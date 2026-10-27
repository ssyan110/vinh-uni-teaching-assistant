"""Synthetic approved version tests; no real teaching approval or output."""
import json
import zipfile
from PIL import Image
import unittest
from pathlib import Path
from unittest.mock import patch

from test_scoped_content_approval import ScopedContentTests
import build_scoped_online_pptx as builder
from workflow_integrity import sha256
from pptx import Presentation
from pptx.util import Inches, Pt
from io import BytesIO


class OnlineBuilderTests(ScopedContentTests):
    def bind_reference(self, record, review, spec):
        reference = self.lesson / '10-design/source-snapshots/reference.pptx'
        reference.parent.mkdir(parents=True, exist_ok=True)
        ppt = Presentation()
        ppt.slide_width, ppt.slide_height = Inches(13.333), Inches(7.5)
        placeholder = BytesIO()
        Image.new('RGB', (16, 16), 'white').save(placeholder, format='PNG')
        for number, item in enumerate(spec['slides'], 1):
            slide = ppt.slides.add_slide(ppt.slide_layouts[6])
            item['reference_slide'] = number
            for e in item['elements']:
                box = [Inches(e[k]) for k in ('x', 'y', 'w', 'h')]
                if e['type'] == 'text':
                    shape = slide.shapes.add_textbox(*box)
                    shape.text = e['text']
                    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(e['font_size'])
                elif e['type'] == 'table':
                    shape = slide.shapes.add_table(len(e['cells']), len(e['cells'][0]), *box)
                    for ri, row in enumerate(e['cells']):
                        for ci, text in enumerate(row):
                            shape.table.cell(ri, ci).text = text
                else:
                    placeholder.seek(0)
                    shape = slide.shapes.add_picture(placeholder, *box)
                e['reference_shape_id'] = shape.shape_id
        ppt.save(reference)
        spec['layout_reference'] = {'path': self.rel(reference), 'sha256': sha256(reference)}
        self.write(review, spec)
        data = json.loads(record.read_text())
        data['artifacts'] = [x for x in data['artifacts'] if x['path'] != self.rel(reference)]
        data['artifacts'].append({'path': self.rel(reference), 'sha256': sha256(reference)})
        next(x for x in data['artifacts'] if x['path'] == self.rel(review))['sha256'] = sha256(review)
        self.write(record, data)

    def make_spec(self):
        record, review = self.make_scoped('2026-fall', 'online', 'online-v1')
        spec = {'lesson_key': self.key, 'offering_id': '2026-fall', 'class_scope': 'group-a', 'artifact_id': 'online-v1',
                'slides': [{'kind': 'cover', 'elements': [{'type': 'text', 'text': '第一课', 'x': 1, 'y': 1, 'w': 7, 'h': 1, 'font_size': 44}]},
                           {'kind': 'learning_flow', 'notes': 'Synthetic', 'elements': [{'type': 'text', 'text': '学习流程', 'x': 1, 'y': 1, 'w': 6, 'h': 1, 'font_size': 34}]}]}
        self.bind_reference(record, review, spec)
        data = json.loads(record.read_text())
        for artifact in data['artifacts']:
            if artifact['path'] == self.rel(review):
                artifact['sha256'] = sha256(review)
        self.write(record, data)
        return record, review, spec

    def test_builds_exact_approved_copy_and_never_overwrites(self):
        record, _, _ = self.make_spec()
        result = builder.build(self.root, self.key, '2026-fall', 'group-a', 'online-v1', self.rel(record))
        self.assertEqual(result['slides'], 2)
        self.assertTrue((self.root / result['pptx']['path']).is_file())
        self.assertEqual(result['status'], 'draft_not_approved')
        with self.assertRaises(FileExistsError):
            builder.build(self.root, self.key, '2026-fall', 'group-a', 'online-v1', self.rel(record))

    def test_free_form_or_changed_reference_cannot_build(self):
        record, review, spec = self.make_spec()
        template = self.root / spec.pop('layout_reference')['path']
        self.write(review, spec)
        data = json.loads(record.read_text())
        next(x for x in data['artifacts'] if x['path'] == self.rel(review))['sha256'] = sha256(review)
        self.write(record, data)
        with self.assertRaisesRegex(ValueError, 'actual approved PPTX layout reference'):
            builder.build(self.root, self.key, '2026-fall', 'group-a', 'online-v1', self.rel(record))
        self.assertFalse((self.lesson / '10-design/pptx-draft/online/2026-fall/online-v1').exists())

    def test_unmapped_reference_text_is_rejected(self):
        record, review, spec = self.make_spec()
        spec['slides'][1]['elements'][0]['reference_shape_id'] = 999
        self.write(review, spec)
        data = json.loads(record.read_text())
        next(x for x in data['artifacts'] if x['path'] == self.rel(review))['sha256'] = sha256(review)
        self.write(record, data)
        with self.assertRaisesRegex(ValueError, 'reference shape bindings'):
            builder.build(self.root, self.key, '2026-fall', 'group-a', 'online-v1', self.rel(record))

    def test_unsupported_reference_object_is_not_silently_dropped(self):
        record, review, spec = self.make_spec()
        path = self.root / spec['layout_reference']['path']
        ppt = Presentation(path)
        ppt.slides[0].shapes.add_group_shape()
        ppt.save(path)
        spec['layout_reference']['sha256'] = sha256(path)
        self.write(review, spec)
        data = json.loads(record.read_text())
        for item in data['artifacts']:
            if item['path'] in (self.rel(path), self.rel(review)):
                item['sha256'] = sha256(self.root / item['path'])
        self.write(record, data)
        with self.assertRaisesRegex(ValueError, 'does not support'):
            builder.build(self.root, self.key, '2026-fall', 'group-a', 'online-v1', self.rel(record))

    def test_native_reference_table_remains_editable_and_complete(self):
        record, review, spec = self.make_spec()
        spec['slides'].append({'kind': 'comparison', 'elements': [
            {'type': 'table', 'cells': [['项目', '时间'], ['方案A', '八点']],
             'x': 1, 'y': 1, 'w': 10, 'h': 4, 'font_size': 25}]})
        self.bind_reference(record, review, spec)
        built = builder.build(self.root, self.key, '2026-fall', 'group-a', 'online-v1', self.rel(record))
        result = Presentation(self.root / built['pptx']['path'])
        table = result.slides[2].shapes[0].table
        self.assertEqual(table.cell(1, 0).text, '方案A')
        self.assertEqual(table.cell(1, 1).text, '八点')

    def test_vocabulary_image_and_plain_audio_are_packaged(self):
        record, review, spec = self.make_spec()
        image = self.lesson / '10-design/assets/example.png'
        image.parent.mkdir(parents=True, exist_ok=True)
        Image.new('RGB', (16, 16), 'white').save(image)
        spec['slides'].append({'kind': 'vocabulary', 'source_refs': ['source/page-1'], 'audio': ['Audio 1-1'],
            'vocabulary': {'word': '学校', 'pinyin': 'xuéxiào', 'part_of_speech': '名词', 'meaning_vi': 'trường học', 'examples': ['我去学校。', '学校很大。']},
            'elements': [{'type': 'text', 'text': '学校 xuéxiào 名词 trường học 我去学校。 学校很大。', 'x': 1, 'y': 1, 'w': 7, 'h': 2, 'font_size': 25},
                         {'type': 'text', 'text': 'Audio 1-1', 'text_role': 'chrome', 'x': 10, 'y': 0.2, 'w': 2, 'h': 0.5, 'font_size': 20},
                         {'type': 'image', 'path': self.rel(image), 'x': 9, 'y': 2, 'w': 2, 'h': 2}]})
        bank = {'lessons': {self.key: {'vocabulary': {'学校': ['我去学校。', '学校很大。']}}}}
        bank_path = self.lesson / '10-design/source-snapshots/example-bank.json'
        self.write(bank_path, bank)
        self.write(self.root / 'course/boya-example-bank.json', bank)
        spec['example_bank_path'] = self.rel(bank_path)
        self.bind_reference(record, review, spec)
        data = json.loads(record.read_text())
        next(x for x in data['artifacts'] if x['path'] == self.rel(review))['sha256'] = sha256(review)
        data['artifacts'].append({'path': self.rel(bank_path), 'sha256': sha256(bank_path)})
        self.write(record, data)
        result = builder.build(self.root, self.key, '2026-fall', 'group-a', 'online-v1', self.rel(record))
        self.assertEqual(result['source_map'][2]['audio'], ['Audio 1-1'])
        with zipfile.ZipFile(self.root / result['pptx']['path']) as package:
            self.assertTrue(any(name.startswith('ppt/media/') for name in package.namelist()))
            self.assertFalse(any(name.endswith('.mp3') for name in package.namelist()))

    def test_corrupt_image_is_rejected_before_writing(self):
        record, review, spec = self.make_spec()
        image = self.lesson / '10-design/assets/bad.png'
        image.parent.mkdir(parents=True, exist_ok=True)
        image.write_bytes(b'\x89PNG\r\n\x1a\n' + b'not an image')
        spec['slides'][1]['elements'].append({'type': 'image', 'path': self.rel(image), 'x': 9, 'y': 2, 'w': 2, 'h': 2})
        self.bind_reference(record, review, spec)
        data = json.loads(record.read_text())
        next(x for x in data['artifacts'] if x['path'] == self.rel(review))['sha256'] = sha256(review)
        self.write(record, data)
        with self.assertRaisesRegex(ValueError, 'cannot be decoded'):
            builder.build(self.root, self.key, '2026-fall', 'group-a', 'online-v1', self.rel(record))
        self.assertFalse((self.lesson / '10-design/pptx-draft/online/2026-fall/online-v1').exists())

    def test_bad_scope_or_changed_copy_never_produces_output(self):
        record, review, _ = self.make_spec()
        with self.assertRaises(ValueError):
            builder.build(self.root, self.key, '2026-fall', 'wrong-class', 'online-v1', self.rel(record))
        review.write_text('{}')
        with self.assertRaises(ValueError):
            builder.build(self.root, self.key, '2026-fall', 'group-a', 'online-v1', self.rel(record))
        self.assertFalse((self.lesson / '10-design/pptx-draft/online/2026-fall/online-v1').exists())

    def test_rejects_unreadable_text_before_output(self):
        record, review, spec = self.make_spec()
        spec['slides'][1]['elements'][0]['font_size'] = 12
        self.bind_reference(record, review, spec)
        data = json.loads(record.read_text())
        next(x for x in data['artifacts'] if x['path'] == self.rel(review))['sha256'] = sha256(review)
        self.write(record, data)
        with self.assertRaises(ValueError):
            builder.build(self.root, self.key, '2026-fall', 'group-a', 'online-v1', self.rel(record))
        self.assertFalse((self.lesson / '10-design/pptx-draft/online/2026-fall/online-v1').exists())

    def test_content_below_25_and_fake_chrome_are_rejected(self):
        record, review, spec = self.make_spec()
        spec['slides'][1]['elements'][0]['font_size'] = 24
        self.bind_reference(record, review, spec)
        with self.assertRaisesRegex(ValueError, 'at least 25 pt'):
            builder.build(self.root, self.key, '2026-fall', 'group-a', 'online-v1', self.rel(record))
        spec['slides'][1]['elements'][0]['text_role'] = 'chrome'
        self.bind_reference(record, review, spec)
        with self.assertRaisesRegex(ValueError, 'chrome is only'):
            builder.build(self.root, self.key, '2026-fall', 'group-a', 'online-v1', self.rel(record))


if __name__ == '__main__':
    unittest.main()
