"""Run: python3 test_gradebook.py. Uses a private disposable copy, never real grades."""
import copy
import hashlib
import io
import json
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
from pathlib import Path
import openpyxl
from server import ROOT, Store, export_workbook, validate_snapshot

def fails(action):
    try:
        action()
    except (ValueError, KeyError, TypeError):
        return
    raise AssertionError('Invalid input was accepted')

def main():
    directory = Path(tempfile.mkdtemp(prefix='qa-', dir=ROOT / 'private'))
    seed = json.loads((ROOT / 'private' / 'seed.json').read_text())
    fixture=copy.deepcopy(seed);fixture['revision']=0
    fixture['classes'][0]['assessments'].append(dict(fixture['classes'][0]['assessments'][0],id='assessment-03'))
    (directory / 'seed.json').write_text(json.dumps(fixture, ensure_ascii=False))
    store = Store(directory)
    assert len(seed['records']) == 152
    assert sum(bool(r['flags']) for r in seed['records'].values()) == 4
    assert {c['id']: sum(s['active'] for s in c['students']) for c in seed['classes']} == {'LT_01':26,'LT_02':25,'LT_03':20,'NNTQ1':28}
    for src in seed['provenance']:
        if 'file' in src:
            path = Path('/Users/ssyan110/Desktop/Work/Teaching/VinhUni/02_評分') / src['file']
            if not path.is_file():
                path = ROOT / 'private' / 'sources' / src['file']
            assert hashlib.sha256(path.read_bytes()).hexdigest() == src['sha256'], 'Original workbook changed'
            wb = openpyxl.load_workbook(path, data_only=False)
            for key,r in seed['records'].items():
                o=r['original']
                if o['file']!=src['file']: continue
                values=[wb[o['sheet']].cell(o['row'],i).value for i in range(4,12)]
                assert values[0]==o['status'] and values[1:6]==r['scores']
                assert values[6]==o['total_cell'] and str(values[7] or '')==r['note']
    c = seed['classes'][0]
    sid = c['students'][0]['code']
    data = {'class_id':c['id'],'assessment_id':'assessment-03','student_code':sid,'revision':0,'status':'done','scores':[2.5,1.5,2,1.5,1],'note':'=Literal feedback'}
    fails(lambda: store.save(dict(data, scores=[None,1,1,1,1])))
    fails(lambda: store.save(dict(data, scores=[3.5,1,1,1,1])))
    fails(lambda: store.save(dict(data, scores=[2.1,1,1,1,1])))
    fails(lambda: store.save(dict(data, student_code='wrong-class')))
    fails(lambda: store.save(dict(data, status='absent')))
    assert store.read()['revision']==0
    saved=store.save(data)
    assert saved['revision']==1
    saved_key='|'.join([c['id'],'assessment-03',sid])
    assert store.read()['records'][saved_key]['total']==8.5
    fails(lambda: store.save(data))
    store.db.close()
    store=Store(directory)
    assert store.read()['records'][saved_key]['note']=='=Literal feedback'
    rkey=next(k for k,r in seed['records'].items() if r['flags'])
    cid,aid,rid=rkey.split('|')
    fails(lambda: store.save({'class_id':cid,'assessment_id':aid,'student_code':rid,'revision':1,'status':'absent','scores':[None]*5,'note':''}))
    partial=dict(data, revision=1,status='supplement',scores=[2,None,None,None,None])
    store.save(partial)
    assert store.read()['records'][saved_key]['total'] is None
    absent=dict(data, revision=2,status='absent',scores=[None]*5)
    store.save(absent)
    assert store.read()['records'][saved_key]['total'] is None
    snapshot=store.read()
    bad=copy.deepcopy(snapshot);bad['records'][saved_key]['scores'][0]=float('nan')
    fails(lambda: store.restore({'revision':3,'snapshot':bad}))
    bad=copy.deepcopy(snapshot);bad['classes'][0]['students'].append(bad['classes'][0]['students'][0])
    fails(lambda: validate_snapshot(bad))
    result=store.restore({'revision':3,'snapshot':fixture})
    assert (directory/'backups'/result['backup']).exists()
    assert len(store.read()['records'])==152 and store.read()['revision']==4
    one=store.add_assessment({'class_id':'NNTQ1','revision':4,'title':'一年级测试','date':'2026-10-02','instructions':'测试记录','source':{'type':'extra','activity':'一年级口语活动'},'rubric':[{'name':'口语总分','max':10,'step':.5,'description':''}]})
    sid1=next(c for c in seed['classes'] if c['id']=='NNTQ1')['students'][0]['code']
    store.save({'class_id':'NNTQ1','assessment_id':one['assessment_id'],'student_code':sid1,'revision':5,'status':'done','scores':[0],'note':'实际零分'})
    assert store.read()['records']['|'.join(['NNTQ1',one['assessment_id'],sid1])]['total']==0
    fails(lambda: store.add_assessment({'class_id':'NNTQ1','revision':6,'title':'bad','rubric':[{'name':'x','max':2.1,'step':.5}]}))
    extra_source={'type':'extra','activity':'课外口语活动'}
    fails(lambda: store.add_assessment({'class_id':'NNTQ1','revision':6,'title':'bad lesson','source':{'type':'textbook','textbook_id':'boya-elementary-i','lesson_keys':['boya-quasi-intermediate-i:lesson-01']},'rubric':[{'name':'总分','max':10,'step':.5}]}))
    for n in range(12):
        store.add_assessment({'class_id':'NNTQ1','revision':store.read()['revision'],'title':f'额外考试 {n+1}','source':extra_source,'rubric':[{'name':'总分','max':10,'step':.5,'description':''}]})
    assert len(next(c for c in store.read()['classes'] if c['id']=='NNTQ1')['assessments'])==13
    exams=next(c for c in store.read()['classes'] if c['id']=='NNTQ1')['assessments']
    first=exams[0]
    edit={'class_id':'NNTQ1','assessment_id':first['id'],'revision':store.read()['revision'],'title':'更正考试名称','source':extra_source,'rubric':first['rubric']}
    store.add_assessment(edit)
    edit['revision']=store.read()['revision'];edit['rubric']=[{'name':'另一标准','max':5,'step':.5}]
    fails(lambda: store.add_assessment(edit))
    dynamic=export_workbook(store.read(),'NNTQ1')
    assert len(openpyxl.load_workbook(io.BytesIO(dynamic)).sheetnames)==14
    export=export_workbook(seed,'LT_02')
    wb=openpyxl.load_workbook(io.BytesIO(export),data_only=False)
    assert len(wb.sheetnames)==3 and wb['Tổng hợp'].max_row==31
    assert sum(row[3].value=='历史' for row in list(wb['Tổng hợp'].rows)[1:])==5
    note_state=copy.deepcopy(seed)
    k=next(iter(note_state['records']));note_state['records'][k]['note']='=HYPERLINK("bad")'
    wb=openpyxl.load_workbook(io.BytesIO(export_workbook(note_state,k.split('|')[0])))
    assert all(cell.data_type!='f' for sheet in wb for row in sheet for cell in row), 'Formula injection'
    (directory/'LT_02-export.xlsx').write_bytes(export)
    store.db.close()
    # The HTTP process uses this same isolated copy. Browser smoke testing may use it later.
    process=subprocess.Popen([sys.executable,str(ROOT/'server.py'),'--port','8768','--data-dir',str(directory)],stdout=subprocess.PIPE,stderr=subprocess.PIPE)
    base='http://127.0.0.1:8768'
    try:
        for _ in range(40):
            try:
                with urllib.request.urlopen(base+'/api/state',timeout=1) as response: body=json.load(response)
                break
            except (OSError,urllib.error.URLError):time.sleep(.1)
        else:raise AssertionError('HTTP startup failed')
        def post(path,value,token=body['token'],origin=base):
            req=urllib.request.Request(base+path,data=json.dumps(value).encode(),headers={'Content-Type':'application/json','X-Gradebook-Token':token,'Origin':origin})
            try:
                with urllib.request.urlopen(req) as response:return response.status,json.load(response)
            except urllib.error.HTTPError as error:return error.code,json.load(error)
        assert post('/api/backup',{},token='bad')[0]==403
        assert post('/api/backup',{},origin='https://untrusted.example')[0]==403
        assert post('/api/backup',{})[0]==200
        req=urllib.request.Request(base+'/api/state',headers={'Host':'untrusted.example:8768'})
        try:urllib.request.urlopen(req);raise AssertionError('DNS rebinding Host accepted')
        except urllib.error.HTTPError as error:assert error.code==403
        with urllib.request.urlopen(base+'/api/export?class=LT_02') as response:
            assert len(openpyxl.load_workbook(io.BytesIO(response.read())).sheetnames)==3
        assert post('/api/restore',{'revision':body['state']['revision'],'snapshot':seed})[0]==200
        report={'passed':True,'checks':['source workbook hashes and every imported raw record','four rosters and historical preservation','score completeness/range/half-point/status validation','stale page conflict prevention','flagged import review gate','restart persistence','partial grades and true zero','invalid restore rejection and pre-restore snapshot','independent first-year assessment','more than twelve exams and dynamic export','textbook identity validation and scored rubric locking','XLSX sheet counts and formula injection prevention','loopback Host/token/Origin protection'],'test_directory':str(directory),'real_database_modified':False}
        (ROOT/'qa'/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
        print(json.dumps(report,ensure_ascii=False,indent=2))
    finally:
        process.terminate();process.wait(timeout=5)

if __name__=='__main__': main()
