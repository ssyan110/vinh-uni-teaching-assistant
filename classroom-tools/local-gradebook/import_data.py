"""One-time, read-only XLSX import. Runtime never contacts Supabase."""
import hashlib
import json
import shutil
from pathlib import Path
from datetime import datetime, timezone
import openpyxl

ROOT = Path(__file__).resolve().parent
SOURCE = Path('/Users/ssyan110/Desktop/Work/Teaching/VinhUni/02_評分').resolve(strict=True)
PRIVATE = ROOT / 'private'
STATUS = {'Đã trình bày': 'done', 'Vắng': 'absent', 'Chưa trình bày': 'pending', 'Chờ bổ sung': 'supplement'}
NAMES = ['内容与任务完成', '组织与连贯', '词语与句式', '发音与易理解程度', '教师问答']

def build():
    seed_path = PRIVATE / 'seed.json'
    if seed_path.exists():
        raise SystemExit('seed.json 已存在；停止，避免覆盖导入资料。')
    live = json.loads((PRIVATE / 'live-roster.json').read_text())['students']
    state = {'format': 'vinh-gradebook-v1', 'classes': [], 'records': {}, 'revision': 0, 'provenance': []}
    for code in ['LT_01', 'LT_02', 'LT_03', 'NNTQ1']:
        current = [s for s in live if s['class_code'] == code]
        students = [{'code': s['student_code'], 'name': s['original_name'], 'seat': s['seat_number'], 'active': True} for s in current]
        cls = {'id': code, 'year': 1 if code == 'NNTQ1' else 2, 'students': students, 'assessments': []}
        state['classes'].append(cls)
        if code == 'NNTQ1':
            continue
        path = SOURCE / (code + '-Bang-diem-bao-cao-ca-nhan.xlsx')
        if path.is_symlink():
            raise ValueError('不接受符号链接来源')
        wb = openpyxl.load_workbook(path, data_only=False)
        cached = openpyxl.load_workbook(path, data_only=True)
        overview = wb['Tổng hợp']
        rubric_rows = [41, 44, 47, 50, 53]
        rubric = [{'name': n, 'max': m, 'step': .5, 'description': str(overview.cell(r, 1).value or '')} for n, m, r in zip(NAMES, [3, 2, 2, 2, 1], rubric_rows)]
        known = {s['code'] for s in students}
        for sheet in wb.worksheets[1:]:
            aid = sheet.title.replace('Lần ', 'assessment-')
            instructions = '\n'.join(str(sheet.cell(r, 1).value or '') for r in [4, 5]).strip()
            exam = {'id': aid, 'title': str(sheet.cell(2, 1).value or sheet.title), 'instructions': instructions, 'rubric': rubric, 'date': '', 'source': {'type': 'textbook', 'textbook_id': 'boya-quasi-intermediate-i', 'lesson_keys': ['boya-quasi-intermediate-i:lesson-' + aid[-2:]], 'activity': ''}}
            has_records = False
            for row in sheet.iter_rows(min_row=8):
                values = [c.value for c in row]
                sid = str(values[2] or '')
                if not sid.isdigit() or len(sid) < 8:
                    continue
                if sid not in known:
                    students.append({'code': sid, 'name': str(values[1]), 'seat': int(values[0]), 'active': False})
                    known.add(sid)
                raw_status = values[3]
                scores = values[4:9]
                note = str(values[10] or '') if len(values) > 10 else ''
                if not raw_status and all(s is None for s in scores) and not note:
                    continue
                has_records = True
                flags = []
                status = STATUS.get(raw_status, 'supplement')
                if raw_status not in STATUS:
                    flags.append('原表状态未识别，请核对')
                if status != 'done' and any(s is not None for s in scores):
                    flags.append('原表非已报告状态仍填有分数；保留原值，核对后再确认')
                valid = all(isinstance(s, (int, float)) and not isinstance(s, bool) and 0 <= s <= c['max'] and s * 2 == int(s * 2) for s,c in zip(scores,rubric))
                if status == 'done' and not valid:
                    flags.append('原表分项分数不完整或超出范围')
                total_cell = row[9]
                source_total = cached[sheet.title].cell(total_cell.row, 10).value
                total = sum(scores) if status == 'done' and valid else None
                if total_cell.data_type != 'f':
                    flags.append('原表总分为手动输入值；保留供核对')
                if isinstance(source_total, (int,float)) and total is not None and source_total != total:
                    flags.append('原表总分与分项合计不一致')
                key = '|'.join([code, aid, sid])
                state['records'][key] = {'status': status, 'scores': scores, 'note': note, 'flags': flags, 'total': total, 'updated_at': '', 'original': {'file': path.name, 'sheet': sheet.title, 'row': total_cell.row, 'status': raw_status, 'scores': scores, 'total_cell': total_cell.value, 'cached_total': source_total, 'note': note}}
            if has_records:
                cls['assessments'].append(exam)
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        backup_dir = PRIVATE / 'sources'
        backup_dir.mkdir(exist_ok=True)
        target = backup_dir / path.name
        if target.exists():
            raise ValueError('来源备份已存在')
        shutil.copy2(path, target)
        state['provenance'].append({'file': path.name, 'sha256': digest})
    state['provenance'].append({'roster_captured_at': '2026-10-02', 'source': 'Current active enrollments; read only; no runtime cloud dependency'})
    seed_path.write_text(json.dumps(state, ensure_ascii=False, indent=2))
    report = {'imported_at': datetime.now(timezone.utc).isoformat(), 'classes': {c['id']: {'active': sum(s['active'] for s in c['students']), 'historical': sum(not s['active'] for s in c['students']), 'assessments': len(c['assessments'])} for c in state['classes']}, 'records': len(state['records']), 'review_records': sum(bool(r['flags']) for r in state['records'].values()), 'source_hashes': state['provenance']}
    (PRIVATE / 'import-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps(report, ensure_ascii=False, indent=2))

if __name__ == '__main__':
    build()
