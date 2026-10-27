"""Single-teacher, authenticated LAN gradebook. SQLite + local JSON recovery snapshots."""
import argparse
import base64
import feedback
import io
import json
import math
import secrets
import socket
import subprocess
import sys
import ipaddress
from http.cookies import SimpleCookie
import sqlite3
import uuid
import qrcode
from datetime import datetime, date
from http.server import HTTPServer, BaseHTTPRequestHandler
from pathlib import Path
from urllib.parse import urlparse, parse_qs
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parent
STATUSES = {'done', 'absent', 'pending', 'supplement'}
MAX_BODY = 8_000_000
CATALOG = json.loads((ROOT / 'catalog.json').read_text())

def lan_address():
    # VPN routes can select a tunnel IP; on this Mac advertise Wi-Fi/Ethernet first.
    if sys.platform == 'darwin':
        for interface in ('en0', 'en1'):
            result = subprocess.run(['/usr/sbin/ipconfig', 'getifaddr', interface], capture_output=True, text=True)
            if result.returncode == 0:
                address = result.stdout.strip()
                if ipaddress.ip_address(address).is_private:
                    return address
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as connection:
            connection.connect(('192.0.2.1', 9))
            address = connection.getsockname()[0]
        return address if ipaddress.ip_address(address).is_private and not ipaddress.ip_address(address).is_loopback else None
    except OSError:
        return None

def now():
    return datetime.now(ZoneInfo('Asia/Ho_Chi_Minh')).isoformat(timespec='seconds')

def number(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool) and math.isfinite(x)

def text(x, limit=10000):
    if not isinstance(x, str) or len(x) > limit:
        raise ValueError('文字格式或长度不正确')
    return x

def validate_rubric(rubric):
    if not isinstance(rubric, list) or not 1 <= len(rubric) <= 10:
        raise ValueError('评分项目必须为 1–10 项')
    for c in rubric:
        text(c['name'], 100)
        text(c.get('description', ''))
        if not c['name'].strip() or not number(c['max']) or not .5 <= c['max'] <= 100 or c['max'] * 2 != int(c['max'] * 2) or c.get('step') != .5:
            raise ValueError('每项满分须为 0.5–100，按半分计')

def validate_source(source):
    if not isinstance(source, dict) or source.get('type') not in ('textbook', 'extra'):
        raise ValueError('请选择教材考试或课外口语活动')
    if source['type'] == 'extra':
        if not text(source.get('activity', ''), 300).strip():
            raise ValueError('请填写课外活动名称')
    else:
        book = next((b for b in CATALOG if b['id'] == source.get('textbook_id')), None)
        keys = source.get('lesson_keys')
        if not book or not isinstance(keys, list) or not keys or len(set(keys)) != len(keys) or any(k not in {l['key'] for l in book['lessons']} for k in keys):
            raise ValueError('请选择对应教材中的课次')

def source_label(assessment):
    source = assessment.get('source', {})
    if source.get('type') == 'extra':
        return '课外口语活动：' + source['activity']
    book = next((b for b in CATALOG if b['id'] == source.get('textbook_id')), None)
    if not book:
        return '原表导入'
    return book['title'] + '；课次：' + '、'.join(str(next(l['number'] for l in book['lessons'] if l['key'] == k)) for k in source['lesson_keys'])

def context(state, class_id, assessment_id, student_code=None):
    cls = next((c for c in state['classes'] if c['id'] == class_id), None)
    if not cls:
        raise ValueError('班级不存在')
    assessment = next((a for a in cls['assessments'] if a['id'] == assessment_id), None)
    if not assessment or (student_code is not None and not any(s['code'] == student_code for s in cls['students'])):
        raise ValueError('评量或学生不存在')
    return cls, assessment

def valid_scores(scores, rubric, complete=False):
    if not isinstance(scores, list) or len(scores) != len(rubric):
        raise ValueError('评分项目数量不正确')
    for score, c in zip(scores, rubric):
        if score is None and not complete:
            continue
        if not number(score) or not 0 <= score <= c['max'] or score * 2 != int(score * 2):
            raise ValueError('请按半分选择范围内的分数；已报告须填齐所有项目')

def validate_snapshot(state):
    if not isinstance(state, dict) or state.get('format') != 'vinh-gradebook-v1':
        raise ValueError('不是本 App 的完整备份')
    if not isinstance(state.get('classes'), list) or not 1 <= len(state['classes']) <= 30 or not isinstance(state.get('records'), dict):
        raise ValueError('备份结构不正确')
    class_ids = set()
    for cls in state['classes']:
        cid = text(cls['id'], 80)
        if not cid or '|' in cid or cid in class_ids or cls['year'] not in (1, 2):
            raise ValueError('班级身份重复或无效')
        class_ids.add(cid)
        ids = set()
        if not isinstance(cls['students'], list) or len(cls['students']) > 1000:
            raise ValueError('名单格式不正确')
        for s in cls['students']:
            code = text(s['code'], 80)
            text(s['name'], 200)
            if not code or '|' in code or code in ids or type(s['active']) is not bool or not isinstance(s['seat'], int):
                raise ValueError('学生身份重复或无效')
            ids.add(code)
        aids = set()
        if not isinstance(cls['assessments'], list):
            raise ValueError('考试格式不正确')
        for a in cls['assessments']:
            aid = text(a['id'], 80)
            if not aid or '|' in aid or aid in aids:
                raise ValueError('评量身份重复或无效')
            aids.add(aid)
            text(a['title'], 300)
            text(a.get('instructions', ''))
            text(a.get('date', ''), 40)
            validate_rubric(a['rubric'])
            if 'source' in a:
                validate_source(a['source'])
    if len(state['records']) > 100000:
        raise ValueError('备份记录过多')
    for key, r in state['records'].items():
        parts = key.split('|')
        if len(parts) != 3:
            raise ValueError('评分身份格式不正确')
        _, a = context(state, *parts)
        if r['status'] not in STATUSES:
            raise ValueError('报告状态不正确')
        text(r.get('note', ''))
        text(r.get('updated_at', ''), 100)
        flags = r.get('flags', [])
        if not isinstance(flags, list) or len(flags) > 20:
            raise ValueError('核对项目格式不正确')
        for flag in flags:
            text(flag, 1000)
        valid_scores(r['scores'], a['rubric'], r['status'] == 'done' and not flags)
        expected = sum(r['scores']) if r['status'] == 'done' and all(number(s) for s in r['scores']) else None
        if r.get('total') != expected:
            raise ValueError('备份总分不符合分项合计')
        if not flags and r['status'] in ('absent', 'pending') and any(s is not None for s in r['scores']):
            raise ValueError('缺席或未报告不能带有分数')
    feedback.validate(state, context, text)
    return state

class Store:
    def __init__(self, directory):
        self.directory = Path(directory).resolve()
        self.directory.mkdir(parents=True, exist_ok=True)
        self.backups = self.directory / 'backups'
        self.backups.mkdir(exist_ok=True)
        self.db = sqlite3.connect(self.directory / 'gradebook.sqlite3')
        self.db.execute('PRAGMA journal_mode=WAL')
        self.db.execute('CREATE TABLE IF NOT EXISTS state (id INTEGER PRIMARY KEY CHECK(id=1), payload TEXT NOT NULL)')
        self.db.execute('CREATE TABLE IF NOT EXISTS history (id INTEGER PRIMARY KEY, timestamp TEXT NOT NULL, record_key TEXT NOT NULL, before_json TEXT, after_json TEXT NOT NULL)')
        if not self.db.execute('SELECT 1 FROM state').fetchone():
            seed = validate_snapshot(json.loads((self.directory / 'seed.json').read_text()))
            with self.db:
                self.db.execute('INSERT INTO state VALUES (1,?)', (json.dumps(seed, ensure_ascii=False),))
            self.backup('initial')

    def read(self):
        return json.loads(self.db.execute('SELECT payload FROM state WHERE id=1').fetchone()[0])

    def backup(self, label='manual', state=None):
        stamp = datetime.now(ZoneInfo('Asia/Ho_Chi_Minh')).strftime('%Y%m%d-%H%M%S-%f')
        name = f'gradebook-{label}-{stamp}.json'
        with (self.backups / name).open('x') as f:
            json.dump(state or self.read(), f, ensure_ascii=False, indent=2)
        return name

    def commit(self, state, key, before, after):
        current = self.read()
        day = now()[:10]
        daily = self.backups / f'gradebook-day-start-{day}.json'
        if not daily.exists():
            with daily.open('x') as f:
                json.dump(current, f, ensure_ascii=False)
        state['revision'] = current['revision'] + 1
        with self.db:
            self.db.execute('UPDATE state SET payload=? WHERE id=1', (json.dumps(state, ensure_ascii=False),))
            self.db.execute('INSERT INTO history(timestamp, record_key, before_json, after_json) VALUES (?,?,?,?)', (now(), key, json.dumps(before, ensure_ascii=False), json.dumps(after, ensure_ascii=False)))
        warning = ''
        try:
            latest = self.backups / f'gradebook-latest-{day}.json'
            temporary = latest.with_suffix('.tmp')
            temporary.write_text(json.dumps(state, ensure_ascii=False))
            temporary.replace(latest)
        except OSError:
            warning = '成绩已保存到数据库，但自动备份失败；请立即手动备份。'
        return {'ok': True, 'revision': state['revision'], 'warning': warning}

    def save(self, data):
        state = self.read()
        if data.get('revision') != state['revision']:
            raise ValueError('资料已在其他页面更新，请先重新载入，避免覆盖')
        cid, aid, sid = data['class_id'], data['assessment_id'], data['student_code']
        _, a = context(state, cid, aid, sid)
        status = data['status']
        if status not in STATUSES:
            raise ValueError('报告状态不正确')
        scores = data['scores']
        valid_scores(scores, a['rubric'], status == 'done')
        if status in ('absent', 'pending') and any(s is not None for s in scores):
            raise ValueError('缺席或未报告不能记作零分，请清除分数或改为待补评')
        note = text(data.get('note', ''))
        key = '|'.join([cid, aid, sid])
        before = state['records'].get(key)
        if before and before.get('flags') and data.get('review_confirmed') is not True:
            raise ValueError('请核对原表异常记录，并勾选已核对')
        r = {'status': status, 'scores': scores, 'note': note, 'total': sum(scores) if status == 'done' else None, 'flags': [], 'updated_at': now()}
        if before and 'original' in before:
            r['original'] = before['original']
        state['records'][key] = r
        return self.commit(state, key, before, r)

    def add_assessment(self, data):
        state = self.read()
        if data.get('revision') != state['revision']:
            raise ValueError('资料已更新，请重新载入')
        cls = next((c for c in state['classes'] if c['id'] == data['class_id']), None)
        if not cls:
            raise ValueError('班级不存在')
        title = text(data['title'], 300).strip()
        if not title:
            raise ValueError('请输入评量名称')
        rubric = data['rubric']
        validate_rubric(rubric)
        source = data['source']
        validate_source(source)
        exam_date = text(data.get('date', ''), 40)
        if exam_date:
            date.fromisoformat(exam_date)
        previous = next((a for a in cls['assessments'] if a['id'] == data.get('assessment_id')), None)
        if data.get('assessment_id') and previous is None:
            raise ValueError('考试不存在')
        if previous and any(k.startswith(cls['id'] + '|' + previous['id'] + '|') for k in state['records']) and rubric != previous['rubric']:
            raise ValueError('已有评分的考试不能更换评分标准；请新增考试')
        a = {'id': previous['id'] if previous else 'assessment-' + uuid.uuid4().hex[:12], 'title': title, 'date': exam_date, 'instructions': text(data.get('instructions', '')), 'rubric': rubric, 'source': source}
        if previous:
            cls['assessments'][cls['assessments'].index(previous)] = a
        else:
            cls['assessments'].append(a)
        result = self.commit(state, 'assessment|' + cls['id'], previous, a)
        result['assessment_id'] = a['id']
        return result

    def restore(self, data):
        state = self.read()
        if data.get('revision') != state['revision']:
            raise ValueError('资料已更新，请重新载入')
        replacement = validate_snapshot(data['snapshot'])
        name = self.backup('before-restore', state)
        result = self.commit(replacement, 'restore', {'revision': state['revision'], 'backup': name}, {'restored': True})
        result['backup'] = name
        return result

def safe_excel(value):
    # Text remains text even if a name/note begins with an Excel formula marker.
    return "'" + value if isinstance(value, str) and value[:1] in ('=', '+', '-', '@', '\t', '\r') else value

def export_workbook(state, class_id):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    cls = next((c for c in state['classes'] if c['id'] == class_id), None)
    if cls is None:
        raise ValueError('班级不存在')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tổng hợp'
    ws.append(['座号', '姓名', '学号', '名单状态'] + [a['title'] for a in cls['assessments']])
    labels = {'done': '已报告', 'absent': '缺席', 'pending': '未报告', 'supplement': '待补评'}
    for s in cls['students']:
        row = [s['seat'], s['name'], s['code'], '当前' if s['active'] else '历史']
        for a in cls['assessments']:
            r = state['records'].get('|'.join([class_id, a['id'], s['code']]))
            row.append('待核对' if r and r['flags'] else r['total'] if r and r['total'] is not None else labels[r['status']] if r else '')
        ws.append([safe_excel(x) for x in row])
    for n,a in enumerate(cls['assessments'], 1):
        sheet = wb.create_sheet(f'Lần {n:02}')
        sheet.append([safe_excel(a['title'])])
        sheet.append([safe_excel(source_label(a) + ('；日期：' + a['date'] if a['date'] else '') + '\n' + a['instructions'])])
        sheet.append(['座号','姓名','学号','名单状态','报告状态'] + [f"{c['name']}/{c['max']}" for c in a['rubric']] + ['总分','备注','待核对事项','修改时间','原表总分单元格','原表缓存总分'])
        for s in cls['students']:
            r = state['records'].get('|'.join([class_id,a['id'],s['code']]))
            original = r.get('original', {}) if r else {}
            row = [s['seat'], s['name'], s['code'], '当前' if s['active'] else '历史', labels[r['status']] if r else '未报告']
            row += r['scores'] if r else [None] * len(a['rubric'])
            row += [r['total'] if r else None, r['note'] if r else '', '；'.join(r['flags']) if r else '', r['updated_at'] if r else '', original.get('total_cell'), original.get('cached_total')]
            sheet.append([safe_excel(x) for x in row])
        sheet.append([])
        for c in a['rubric']:
            sheet.append([safe_excel(c['name']), safe_excel(c['description'])])
    for sheet in wb:
        sheet.freeze_panes = 'D4' if sheet != ws else 'D2'
        sheet.auto_filter.ref = f'A{3 if sheet != ws else 1}:{sheet.cell(len(cls["students"])+(3 if sheet != ws else 1),sheet.max_column).coordinate}'
        for row in sheet:
            for cell in row:
                cell.font = Font(name='Times New Roman', size=12, color='17324D')
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.row == (3 if sheet != ws else 1):
                    cell.fill = PatternFill('solid', fgColor='E7F1EF')
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 22
        sheet.column_dimensions['B'].width = 30
        sheet.column_dimensions['A'].width = 8
        sheet.column_dimensions['C'].width = 22
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        main_columns = sheet.max_column if sheet == ws else 6 + len(cls['assessments'][wb.index(sheet)-1]['rubric'])
        sheet.print_area = f'A1:{sheet.cell(len(cls["students"])+(1 if sheet==ws else 3),main_columns).coordinate}'
        sheet.print_title_rows = '1:1' if sheet == ws else '1:3'
        if sheet != ws:
            for row_number, height in [(1, 36), (2, 76)]:
                sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=main_columns)
                sheet.row_dimensions[row_number].height = height
            for row_number in range(4,len(cls['students'])+4):
                sheet.row_dimensions[row_number].height = 32
            # Keep detailed rubric text in Excel; printed tables use only the grade columns.
            for row_number in range(len(cls['students'])+5,sheet.max_row+1):
                sheet.merge_cells(start_row=row_number,start_column=2,end_row=row_number,end_column=sheet.max_column)
                sheet.row_dimensions[row_number].height = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

class Handler(BaseHTTPRequestHandler):
    def send(self, status, value, mime='application/json; charset=utf-8', filename=None):
        if mime.startswith('application/json'):
            value = json.dumps(value, ensure_ascii=False).encode()
        elif isinstance(value, str):
            value = value.encode()
        self.send_response(status)
        self.send_header('Content-Type', mime)
        self.send_header('Cache-Control', 'no-store')
        self.send_header('X-Content-Type-Options', 'nosniff')
        self.send_header('Referrer-Policy', 'no-referrer')
        self.send_header('Content-Security-Policy', "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'none'; form-action 'self'")
        if filename:
            self.send_header('Content-Disposition', 'attachment; filename="' + filename + '"')
        self.send_header('Content-Length', str(len(value)))
        self.end_headers()
        self.wfile.write(value)

    def allowed(self):
        # Resolve the active interface on each request: DHCP, Wi-Fi changes and
        # waking offline must not leave a startup-only allowlist behind.
        port = self.server.server_port
        hosts = {f'127.0.0.1:{port}', f'localhost:{port}'}
        address = lan_address()
        if address:
            hosts.add(f'{address}:{port}')
        return self.headers.get('Host') in hosts

    def local(self):
        return ipaddress.ip_address(self.client_address[0]).is_loopback

    def authenticated(self):
        if self.local():
            return True
        cookies = SimpleCookie()
        try:
            cookies.load(self.headers.get('Cookie', ''))
            value = cookies.get('gradebook-access')
            return bool(value and secrets.compare_digest(value.value.encode(), self.server.access_key.encode()))
        except Exception:
            return False

    def do_GET(self):
        if not self.allowed():
            return self.send(403, {'error': '本机访问地址不正确'})
        url = urlparse(self.path)
        if url.path == '/connect':
            key = parse_qs(url.query).get('key', [''])[0]
            if not secrets.compare_digest(key.encode(), self.server.access_key.encode()):
                return self.send(403, {'error': '手机入口已失效，请在电脑上重新打开“手机评分”。'})
            self.send_response(303)
            self.send_header('Location', '/')
            self.send_header('Set-Cookie', f'gradebook-access={self.server.access_key}; HttpOnly; SameSite=Strict; Path=/')
            self.send_header('Cache-Control', 'no-store')
            self.send_header('Referrer-Policy', 'no-referrer')
            self.send_header('Content-Length', '0')
            self.end_headers()
            return
        if not self.authenticated():
            return self.send(403, '请先使用电脑上“手机评分”提供的专属链接。', 'text/plain; charset=utf-8')
        try:
            if url.path == '/api/mobile':
                if not self.local():
                    return self.send(403, {'error': '请在电脑上查看手机入口'})
                address = lan_address()
                if not address:
                    raise ValueError('电脑尚未连接局域网。请连接 Wi-Fi 或手机热点后，再点击“手机评分”；无需重启。')
                mobile_url = f'http://{address}:{self.server.server_port}/connect?key={self.server.access_key}'
                qr_image = io.BytesIO()
                qrcode.make(mobile_url).save(qr_image, format='PNG')
                return self.send(200, {'url': mobile_url, 'qr': 'data:image/png;base64,' + base64.b64encode(qr_image.getvalue()).decode('ascii')})
            if url.path == '/api/health':
                return self.send(200, {'ok': True, 'service': 'vinh-local-gradebook'})
            if url.path == '/api/catalog':
                return self.send(200, CATALOG)
            if url.path == '/api/state':
                return self.send(200, {'state': self.server.store.read(), 'token': self.server.token})
            if url.path == '/api/backups':
                names = sorted((p.name for p in self.server.store.backups.glob('*.json')), reverse=True)
                return self.send(200, {'names': names})
            if url.path == '/api/backup-file':
                name = parse_qs(url.query).get('name', [''])[0]
                if not name or Path(name).name != name or not name.endswith('.json'):
                    raise ValueError('备份文件名无效')
                path = self.server.store.backups / name
                if path.is_symlink() or not path.is_file():
                    raise ValueError('备份不存在')
                return self.send(200, path.read_bytes(), 'application/octet-stream', name)
            if url.path == '/api/export':
                cid = parse_qs(url.query).get('class', [''])[0]
                data = export_workbook(self.server.store.read(), cid)
                return self.send(200, data, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', f'{cid}-gradebook.xlsx')
            if url.path == '/api/feedback':
                query = parse_qs(url.query)
                cid, aid = query.get('class', [''])[0], query.get('assessment', [''])[0]
                state = self.server.store.read()
                context(state, cid, aid)
                return self.send(200, {'revision': state['revision'], 'rows': feedback.rows(state, cid, aid)})
            if url.path == '/api/feedback-export':
                query = parse_qs(url.query)
                cid, aid = query.get('class', [''])[0], query.get('assessment', [''])[0]
                state = self.server.store.read()
                context(state, cid, aid)
                return self.send(200, feedback.export_html(state, cid, aid), 'text/html; charset=utf-8', 'individual-feedback.html')
            if url.path == '/api/history':
                key = parse_qs(url.query).get('key', [''])[0]
                rows = self.server.store.db.execute('SELECT timestamp,before_json,after_json FROM history WHERE record_key=? ORDER BY id DESC LIMIT 30', (key,)).fetchall()
                return self.send(200, [{'timestamp': r[0], 'before': json.loads(r[1]), 'after': json.loads(r[2])} for r in rows])
            files = {'/': ('index.html', 'text/html; charset=utf-8'), '/app.js': ('app.js', 'text/javascript; charset=utf-8'), '/queue.js': ('queue.js', 'text/javascript; charset=utf-8'), '/feedback.js': ('feedback.js', 'text/javascript; charset=utf-8'), '/style.css': ('style.css', 'text/css; charset=utf-8')}
            if url.path not in files:
                return self.send(404, {'error': '页面不存在'})
            name, mime = files[url.path]
            self.send(200, (ROOT / 'web' / name).read_bytes(), mime)
        except (ValueError, KeyError, TypeError) as e:
            self.send(400, {'error': str(e)})
        except Exception:
            self.send(500, {'error': '读取失败；本机资料仍保留，请检查终端信息。'})
            raise

    def do_POST(self):
        origin = self.headers.get('Origin')
        if not self.allowed() or not self.authenticated() or self.headers.get('X-Gradebook-Token') != self.server.token or (origin and origin != 'http://' + self.headers.get('Host', '')):
            return self.send(403, {'error': '访问验证失败，请重新载入本机页面'})
        try:
            length = int(self.headers.get('Content-Length', '0'))
            if not 0 < length <= MAX_BODY or not self.headers.get('Content-Type', '').startswith('application/json'):
                raise ValueError('请求格式或大小不正确')
            data = json.loads(self.rfile.read(length))
            path = urlparse(self.path).path
            store = self.server.store
            if path == '/api/save':
                result = store.save(data)
            elif path == '/api/assessment':
                result = store.add_assessment(data)
            elif path == '/api/feedback':
                result = feedback.mutate(store, data, context, text, now)
            elif path == '/api/backup':
                result = {'ok': True, 'name': store.backup()}
            elif path == '/api/restore':
                result = store.restore(data)
            else:
                return self.send(404, {'error': '操作不存在'})
            self.send(200, result)
        except (ValueError, KeyError, TypeError) as e:
            self.send(400, {'error': str(e)})
        except Exception:
            self.send(500, {'error': '保存失败；请保留当前页面并检查磁盘或终端信息。'})
            raise

    def log_message(self, fmt, *args):
        if urlparse(self.path).path != '/connect' and args and str(args[1]) != '200':
            super().log_message(fmt, *args)

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--port', type=int, default=8765)
    parser.add_argument('--data-dir', type=Path, default=ROOT / 'private')
    args = parser.parse_args()
    if not (ROOT == args.data_dir.resolve() or ROOT in args.data_dir.resolve().parents):
        parser.error('资料目录必须在本 App 文件夹内')
    store = Store(args.data_dir)
    server = HTTPServer(('0.0.0.0', args.port), Handler)
    server.store = store
    server.token = secrets.token_urlsafe(32)
    server.access_key = secrets.token_urlsafe(32)
    print(f'教师评分已启动：http://127.0.0.1:{args.port}（Ctrl+C 关闭）', flush=True)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()
        store.db.close()

if __name__ == '__main__':
    main()
