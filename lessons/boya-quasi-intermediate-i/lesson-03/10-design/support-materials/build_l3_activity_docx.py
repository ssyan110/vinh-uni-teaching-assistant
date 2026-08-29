#!/usr/bin/env python3
"""Convert the approved-scope Lesson 3 activity drafts to editable DOCX files.

This is intentionally a small Markdown-to-DOCX converter: it preserves the
existing draft wording and tables and does not invent answers or exercises.
Outputs stay in the lesson's 10-design/support-materials directory.
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
INK = "17324D"
ACCENT = "365F7D"
HEADER = "DCE8F2"
GRID = "B7C6D6"


def font_run(run, size=12, bold=False, color=INK):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.rFonts
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")
    rf.set(qn("w:cs"), "Times New Roman")
    rf.set(qn("w:eastAsia"), "KaiTi")
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")


def style_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.7)
    sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "KaiTi")
    normal_lang = normal._element.rPr.find(qn("w:lang"))
    if normal_lang is None:
        normal_lang = OxmlElement("w:lang")
        normal._element.rPr.append(normal_lang)
    normal_lang.set(qn("w:val"), "zh-CN")
    normal_lang.set(qn("w:eastAsia"), "zh-CN")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2
    for name, size in (("Title", 20), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        st = doc.styles[name]
        st.font.name = "Times New Roman"; st.font.size = Pt(size); st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(ACCENT)
        st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "KaiTi")
        st_lang = st._element.rPr.find(qn("w:lang"))
        if st_lang is None:
            st_lang = OxmlElement("w:lang")
            st._element.rPr.append(st_lang)
        st_lang.set(qn("w:val"), "zh-CN")
        st_lang.set(qn("w:eastAsia"), "zh-CN")


def cell_border(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders"); tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn("w:" + edge))
        if el is None: el = OxmlElement("w:" + edge); borders.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:color"), GRID)


def add_table(doc, rows):
    if not rows: return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for ri, values in enumerate(rows):
        for ci in range(cols):
            cell = table.cell(ri, ci); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_border(cell)
            if ri == 0:
                tcpr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), HEADER); tcpr.append(shd)
            p = cell.paragraphs[0]; p.text = ""; p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(values[ci] if ci < len(values) else "")
            font_run(run, size=10.5, bold=(ri == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def convert(src: Path, dst: Path):
    lines = src.read_text(encoding="utf-8").splitlines()
    doc = Document(); style_doc(doc)
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip(): i += 1; continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells): rows.append(cells)
                i += 1
            add_table(doc, rows); continue
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            level = len(m.group(1)); text = m.group(2)
            p = doc.add_paragraph(style="Title" if level == 1 else f"Heading {level}")
            r = p.add_run(text); font_run(r, size=(20 if level == 1 else 16 if level == 2 else 14), bold=True, color=ACCENT)
            i += 1; continue
        if re.match(r"^[-*]\s+", line):
            text = re.sub(r"^[-*]\s+", "", line)
            p = doc.add_paragraph(style="List Bullet"); r = p.add_run(text); font_run(r); i += 1; continue
        if re.match(r"^\d+[.)]\s+", line):
            text = re.sub(r"^\d+[.)]\s+", "", line)
            p = doc.add_paragraph(style="List Number"); r = p.add_run(text); font_run(r); i += 1; continue
        p = doc.add_paragraph(); r = p.add_run(line); font_run(r); i += 1
    doc.save(dst)


FILES = {
    "lesson-03-活动01-听力证据记录-draft.md": "lesson-03-活动01-听力证据记录-draft.docx",
    "lesson-03-活动02-三段短文信息表-draft.md": "lesson-03-活动02-三段短文信息表-draft.docx",
    "lesson-03-活动03-小组总结与个人迁移-draft.md": "lesson-03-活动03-小组总结与个人迁移-draft.docx",
    "lesson-03-评量表-draft.md": "lesson-03-评量表-draft.docx",
    "lesson-03-课末检查-draft.md": "lesson-03-课末检查-draft.docx",
}

if __name__ == "__main__":
    for src_name, dst_name in FILES.items():
        convert(ROOT / src_name, ROOT / dst_name)
        print(ROOT / dst_name)
