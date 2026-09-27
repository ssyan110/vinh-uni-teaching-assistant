#!/usr/bin/env python3
"""Build the first-class student prep card and teacher evidence sheet.

These are reversible support-draft artifacts. They do not modify the approved
lesson package or the user's icebreaker copy form.
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
INK = "17324D"
BLACK = "000000"
HEADER = "DCE8F2"
PALE = "F5F8FA"
GRID = "B9C9D6"


def set_rfonts(rpr, east_asia="KaiTi"):
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), east_asia)


def set_lang(rpr):
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")


def format_run(run, size=10.5, bold=False, color=INK):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    set_rfonts(rpr)
    set_lang(rpr)


def format_style(style, size=10.5, bold=False, color=INK):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style._element.get_or_add_rPr()
    set_rfonts(rpr)
    set_lang(rpr)


def setup_doc(doc, *, landscape=False, body_size=10.5):
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.top_margin = Inches(0.42)
        section.bottom_margin = Inches(0.42)
        section.left_margin = Inches(0.42)
        section.right_margin = Inches(0.42)
    else:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.38)
        section.bottom_margin = Inches(0.38)
        section.left_margin = Inches(0.48)
        section.right_margin = Inches(0.48)

    format_style(doc.styles["Normal"], body_size, color=INK)
    doc.styles["Normal"].paragraph_format.space_after = Pt(2)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.0
    format_style(doc.styles["Title"], 16, True, BLACK)
    format_style(doc.styles["Heading 1"], 12, True, BLACK)
    format_style(doc.styles["Heading 2"], 11, True, BLACK)


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shading = tcpr.first_child_found_in("w:shd")
    if shading is None:
        shading = OxmlElement("w:shd")
        tcpr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID, size="5"):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn("w:" + edge)
        item = borders.find(tag)
        if item is None:
            item = OxmlElement("w:" + edge)
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), size)
        item.set(qn("w:space"), "0")
        item.set(qn("w:color"), color)


def set_cell_padding(cell, top=55, start=80, bottom=55, end=80):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, *, size=9.4, bold=False, color=INK, align=WD_ALIGN_PARAGRAPH.LEFT, fill=None):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell)
    set_cell_padding(cell)
    if fill:
        shade_cell(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    format_run(run, size=size, bold=bold, color=color)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(3 if level == 1 else 2)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    format_run(run, size=12 if level == 1 else 11, bold=True, color=BLACK)
    return paragraph


def add_para(doc, text, *, size=10.5, bold=False, color=INK, align=WD_ALIGN_PARAGRAPH.LEFT, after=2):
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    format_run(run, size=size, bold=bold, color=color)
    return paragraph


def add_title(doc, title, subtitle):
    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(title)
    format_run(run, size=16, bold=True, color=BLACK)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(subtitle)
    format_run(run, size=8.8, color=INK)


def add_table(doc, headers, rows, widths, *, size=9.2, header_size=9.2, row_height=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, size=header_size, bold=True, color=BLACK, align=WD_ALIGN_PARAGRAPH.CENTER, fill=HEADER)
        table.rows[0].cells[index].width = Inches(widths[index])
    for row in rows:
        cells = table.add_row().cells
        if row_height:
            table.rows[-1].height = Inches(row_height)
            table.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for index, value in enumerate(row):
            set_cell_text(cells[index], value, size=size, fill=PALE if len(table.rows) % 2 == 0 else None)
            cells[index].width = Inches(widths[index])
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)
    return table


def build_student_card(path):
    doc = Document()
    setup_doc(doc, body_size=9.6)
    add_title(doc, "第一课｜学生预习检查卡", "《丽丽是独生女》｜上课开始填写｜这不是正式成绩")

    fields = doc.add_table(rows=1, cols=3)
    fields.alignment = WD_TABLE_ALIGNMENT.CENTER
    fields.autofit = False
    for cell, text, width in zip(fields.rows[0].cells, ["姓名：________________", "班级：____________", "日期：____________"], [2.8, 2.1, 1.8]):
        set_cell_text(cell, text, size=9.5)
        cell.width = Inches(width)
    add_para(doc, "不会写完整句子时，写关键词就可以。请至少写出两项具体内容，不要只写「完成」或「有」。", size=9.3, after=2)

    add_heading(doc, "1｜我带来的预习证据")
    add_table(
        doc,
        ["我做过什么", "我的具体记录"],
        [
            ["教材上一个不懂的地方", "第 ____ 页：____________________________________________"],
            ["听过一个音频", "音频 1-____；我记住的词或信息：________________________"],
            ["看过 P10 信息表或 P11 提纲", "我写下了：______________________________________________"],
            ["准备一个想问同学的问题", "我想问：_______________________________________________"],
        ],
        [2.05, 4.65],
        size=8.9,
        header_size=9.0,
        row_height=0.32,
    )

    add_heading(doc, "2｜我现在能说什么")
    add_table(
        doc,
        ["问题", "我的回答"],
        [
            ["第一课主要谈：", "________________________________________________________"],
            ["我记住的一条人物、关系、数字或事情：", "______________________________________________"],
            ["我还不懂，或者想问老师：", "________________________________________________"],
        ],
        [2.85, 3.85],
        size=8.9,
        header_size=9.0,
        row_height=0.32,
    )

    add_heading(doc, "3｜和同伴说一次")
    add_para(doc, "先说：我预习时注意到 ______________________________。", size=9.4, after=1)
    add_para(doc, "再问：你听到／看到了什么？", size=9.4, after=1)
    add_para(doc, "同伴说到的一个关键词：_______________________________________________", size=9.4, after=2)

    add_heading(doc, "4｜我今天想练习")
    add_para(doc, "□ 我会先听完，再回答。　□ 我听不懂会请求重复。　□ 我会问同伴一个问题。", size=9.4, after=1)
    add_para(doc, "我今天最想练习：______________________________________________________", size=9.4, after=2)
    add_heading(doc, "5｜下课前再写一次")
    add_para(doc, "今天我完成的一件事：________________________________________________", size=9.4, after=1)
    add_para(doc, "我还想问或再练习：__________________________________________________", size=9.4, after=1)
    add_para(doc, "下次上课前我会准备：________________________________________________", size=9.4, after=2)
    add_para(doc, "请把这张卡交给老师，或在老师指定的位置留下。老师会根据你的记录给你合适的支架。", size=8.8, color=INK, after=0)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def build_teacher_sheet(path):
    doc = Document()
    setup_doc(doc, landscape=True, body_size=9.2)
    add_title(doc, "第一课｜预习证据教师速记表", "教师用｜只记原始观察，不计分、不自动加总")
    add_para(doc, "日期：________________　班级：________________　教室：________________　教师：________________", size=9.2, after=2)
    add_para(doc, "标记：✓ 有具体证据　△ 有准备但需要支架　— 尚未看到具体证据　A/B/C 只供教师当天安排任务使用。", size=8.8, after=3)
    rows = [[str(i), "", "", "", "", "", ""] for i in range(1, 21)]
    add_table(
        doc,
        ["序号", "学生", "有书本／音频痕迹", "能说一条具体信息", "能提出一个问题", "能请求重复／确认", "A/B/C｜支架"],
        rows,
        [0.45, 1.35, 1.45, 1.65, 1.35, 1.55, 1.6],
        size=8.4,
        header_size=8.3,
        row_height=0.25,
    )
    add_heading(doc, "3–5 分钟恢复路线记录", level=1)
    add_table(
        doc,
        ["学生", "看题并写两个词", "听一次并写一个信息", "对同伴说一句", "结果"],
        [["", "", "", "", "C→B／继续支架"] for _ in range(5)],
        [1.45, 1.9, 2.2, 1.9, 1.6],
        size=8.5,
        header_size=8.5,
        row_height=0.28,
    )
    add_para(doc, "课末回看：最常见的理解问题：____________________________　最需要的语言支架：____________________________", size=8.8, after=1)
    add_para(doc, "下次优先安排的搭档／互动方式：________________________________　需要个别跟进的学生和原因：________________________", size=8.8, after=0)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


if __name__ == "__main__":
    student = ROOT / "lesson-01-预习检查卡.docx"
    teacher = ROOT / "lesson-01-预习证据-教师速记表.docx"
    build_student_card(student)
    build_teacher_sheet(teacher)
    print(student)
    print(teacher)
