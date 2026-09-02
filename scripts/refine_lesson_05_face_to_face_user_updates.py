from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import json

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path('/Users/ssyan110/Development/vinh-uni-teaching-assistant')
DECK = ROOT / 'lessons/boya-quasi-intermediate-i/lesson-05/10-design/pptx-draft/face-to-face/lesson-05-实体课.pptx'
ONLINE = ROOT / 'lessons/boya-quasi-intermediate-i/lesson-05/10-design/pptx-draft/online/lesson-05-在线预习.pptx'


def shapes_with_text(slide):
    return [shape for shape in slide.shapes if hasattr(shape, 'text')]


def set_text(shape, value: str) -> None:
    runs = [run for para in shape.text_frame.paragraphs for run in para.runs]
    if runs:
        runs[0].text = value
        for run in runs[1:]:
            run.text = ''
    else:
        shape.text = value


def find_exact(slide, value: str):
    for shape in shapes_with_text(slide):
        if shape.text == value:
            return shape
    raise ValueError(f'找不到文字框：{value}')


def find_prefix(slide, prefix: str):
    for shape in shapes_with_text(slide):
        if shape.text.startswith(prefix):
            return shape
    raise ValueError(f'找不到文字框：{prefix}')


def clear_shape(shape) -> None:
    set_text(shape, '')


def replace_slide_shapes(target, source) -> None:
    for shape in list(target.shapes):
        target.shapes._spTree.remove(shape._element)
    for shape in source.shapes:
        target.shapes._spTree.append(deepcopy(shape._element))


def delete_slide(prs, index_zero_based: int) -> None:
    slide_id = prs.slides._sldIdLst[index_zero_based]
    prs.part.drop_rel(slide_id.rId)
    prs.slides._sldIdLst.remove(slide_id)


def replace_route_image(target_slide, online_slide) -> None:
    target_pic = next((s for s in target_slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE), None)
    online_pic = next((s for s in online_slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE), None)
    if target_pic is None or online_pic is None:
        raise ValueError('学习路线图片缺失')
    left, top, width, height = target_pic.left, target_pic.top, target_pic.width, target_pic.height
    target_slide.shapes._spTree.remove(target_pic._element)
    target_slide.shapes.add_picture(BytesIO(online_pic.image.blob), left, top, width, height)


def set_page4(slide):
    set_text(find_exact(slide, '先想一想'), '音乐和我们的生活')
    set_text(find_exact(slide, '你以前来过中国吗？'), '你平时怎么听音乐？')
    set_text(find_prefix(slide, '• 你对中国有什么印象？'), '• 你喜欢听什么音乐？\n• 你常在哪里听音乐？\n• 音乐给你带来什么感受？')
    set_text(find_prefix(slide, '参考句式：'), '参考句式：我喜欢……，因为……')


def set_simple_question(slide, old, new):
    set_text(find_exact(slide, old), new)


def set_source_task(slide, page, title, prompt, words, expressions):
    page_shape = next((s for s in shapes_with_text(slide) if s.text.startswith('教材 ')), None)
    if page_shape is None:
        raise ValueError('教材页码框缺失')
    set_text(page_shape, page)
    set_text(find_prefix(slide, '介绍'), title)
    set_text(find_prefix(slide, '请你用'), prompt)
    set_text(find_prefix(slide, '• 人物'), '• 人物\n• 音乐内容\n• 影响或感受')
    set_text(find_prefix(slide, '参考词语：'), f'参考词语：{words}\n常用表达：{expressions}')


def set_no_page_task(slide, title, prompt):
    for shape in shapes_with_text(slide):
        if shape.text.startswith('教材 '):
            clear_shape(shape)
    title_shape = find_exact(slide, '请你说说')
    prompt_shape = next((s for s in shapes_with_text(slide) if s.text.startswith(('说说', '比较', '请介绍')) and s is not title_shape), None)
    if prompt_shape is None:
        raise ValueError('无教材页任务的问题框缺失')
    set_text(title_shape, title)
    set_text(prompt_shape, prompt)
    set_text(find_prefix(slide, '必须使用'), '请说6—8个句子，不少于60字。')


def main():
    if not DECK.is_file() or not ONLINE.is_file():
        raise FileNotFoundError('线上或实体课 PPTX 缺失')
    prs = Presentation(DECK)
    online = Presentation(ONLINE)
    if len(prs.slides) != 58:
        raise ValueError(f'预期删除前实体课为58页，实际为{len(prs.slides)}页')

    # Explicitly requested cross-deck reuse: route image and complete learning-goals page.
    replace_route_image(prs.slides[1], online.slides[1])
    replace_slide_shapes(prs.slides[2], online.slides[2])

    set_page4(prs.slides[3])
    set_simple_question(prs.slides[8], '你小时候学过什么？', '你喜欢什么音乐？为什么？')
    set_simple_question(prs.slides[9], '你为什么学习中文？', '你常在哪听音乐，为什么？')
    set_simple_question(prs.slides[10], '你觉得中文哪里难？', '你觉得学生应该学习音乐吗？为什么？')
    set_simple_question(prs.slides[11], '你和同学常常怎么互相帮助？', '你想学什么乐器？为什么？')

    set_source_task(
        prs.slides[18], '教材 P45', '介绍他的音乐生活',
        '请你介绍“我”的音乐生活：从小喜欢什么，受了哪些人的影响。说6—8个句子，不少于60字。',
        '积极、参加、成为、音乐迷', '从小就、每……都……、既……又……、可是、受……影响')
    set_no_page_task(prs.slides[27], '说说你的音乐生活。', '请说说你的音乐生活。')
    set_source_task(
        prs.slides[32], '教材 P46–47', '介绍他的音乐课',
        '请你介绍李老师的音乐课：他教“我们”什么，同学们为什么喜欢学吹笛子。说6—8个句子，不少于60字。',
        '欣赏、欧洲、民族、吹、笛子', '越……越……、把、一……就……、比如')
    set_no_page_task(prs.slides[41], '说说你的音乐课。', '请说说你的音乐课。')
    set_source_task(
        prs.slides[46], '教材 P48', '介绍李老师教的音乐',
        '请你介绍李老师教的音乐有什么特色。说6—8个句子，不少于60字。',
        '耐心、清楚、内容、夏天的风、冬天的火、生活、美好', '从……那里、对……感兴趣')
    set_no_page_task(prs.slides[53], '介绍你最喜欢的课。', '请介绍你最喜欢的一门课。')

    slide56 = prs.slides[55]
    set_text(find_prefix(slide56, '请你介绍'), '请介绍朴大宇和音乐老师')
    set_text(find_prefix(slide56, '请填表后'), '请填表后，说一说朴大宇和音乐老师。')

    slide57 = prs.slides[56]
    set_text(find_prefix(slide57, '根据课本'), '根据本课内容，回答问题')
    set_text(find_prefix(slide57, '• 朴大宇以前'), '• “我”从小就喜欢什么？\n• “我”受了哪些人的影响？\n• 李老师教“我们”什么？\n• 同学们为什么越来越喜欢吹笛子？\n• 李老师讲课有什么特点？\n• 音乐对人们的生活有什么影响？')

    slide58 = prs.slides[57]
    set_text(find_prefix(slide58, '请你说说'), '介绍我喜欢的一个人')
    set_text(find_prefix(slide58, '① 我的北京印象'), '你喜欢谁？')
    set_text(find_prefix(slide58, '② 我的中文学习'), '为什么喜欢？')
    set_text(find_prefix(slide58, '③ 我的学习比较'), '介绍一下这个人。')
    set_text(find_prefix(slide58, '必须使用'), '说10到12句，至少100字。')
    # Remove the old three-part prompts and leave one direct, lesson-related question.
    set_text(find_exact(slide58, '____________________________'), '请介绍这个人。')
    for shape in shapes_with_text(slide58):
        if shape.text == '你对北京有什么印象？' or shape.text == '你在哪里学习汉语？' or shape.text == '哪里学习汉语更有效率？':
            clear_shape(shape)

    # Delete the two explicitly requested original slides 13 and 14, from right to left.
    delete_slide(prs, 13)
    delete_slide(prs, 12)
    if len(prs.slides) != 56:
        raise AssertionError(f'删除后应为56页，实际为{len(prs.slides)}页')

    tmp = DECK.with_suffix('.user-update.tmp.pptx')
    prs.save(tmp)
    tmp.replace(DECK)

    manifest = DECK.parent / 'manifest.json'
    if manifest.is_file():
        data = json.loads(manifest.read_text())
        data['sha256'] = sha256(DECK.read_bytes()).hexdigest()
        data['bytes'] = DECK.stat().st_size
        data['status'] = 'draft_user_refinement_2026-09-01'
        manifest.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')
    print(f'updated {DECK}')
    print(f'slides {len(prs.slides)}')
    print(f'sha256 {sha256(DECK.read_bytes()).hexdigest()}')


if __name__ == '__main__':
    main()
