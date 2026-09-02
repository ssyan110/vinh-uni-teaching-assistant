from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from pptx import Presentation

ROOT = Path('/Users/ssyan110/Development/vinh-uni-teaching-assistant')
FACE = ROOT / 'lessons/boya-quasi-intermediate-i/lesson-04/10-design/pptx-draft/face-to-face/lesson-04-实体课.pptx'
ONLINE = ROOT / 'lessons/boya-quasi-intermediate-i/lesson-04/10-design/pptx-draft/online/lesson-04-在线预习.pptx'
L2_ONLINE = ROOT / 'lessons/boya-quasi-intermediate-i/lesson-02/20-approved/pptx/lesson-02-在线预习.pptx'

face = Presentation(FACE)
online = Presentation(ONLINE)

# Slide 3: use the exact current online goals page, including its visual style
# and current wording. The lesson header is already the same lesson identity.
target = face.slides[2]
source = online.slides[2]
for sh in list(target.shapes):
    target.shapes._spTree.remove(sh._element)
for sh in source.shapes:
    target.shapes._spTree.append(deepcopy(sh._element))

# Only the requested activity prompts are changed.
for n in (28, 41, 56):
    slide = face.slides[n - 1]
    for sh in slide.shapes:
        if hasattr(sh, 'text') and '必须使用10个课本中的词语和5个句式。' in sh.text:
            sh.text = sh.text.replace('必须使用10个课本中的词语和5个句式。', '必须使用5个课本中的词语和3个句式。')

# Slide 54 is a divider-like expression page; split it into two focused pages
# while preserving its existing face-to-face layout.
src = face.slides[53]
second = face.slides.add_slide(src.slide_layout)
for sh in src.shapes:
    second.shapes._spTree.append(deepcopy(sh._element))
for sh in src.shapes:
    if hasattr(sh, 'text') and sh.text == '记得住／看得清楚':
        sh.text = '记得住'
        break
for sh in second.shapes:
    if hasattr(sh, 'text') and sh.text == '记得住／看得清楚':
        sh.text = '看得清楚'
        break
# Add the inserted page in sequence and update visible page labels after it.
ids = face.slides._sldIdLst
new_id = ids[-1]
ids.remove(new_id)
src_id = list(ids)[53]
ids.insert(list(ids).index(src_id) + 1, new_id)
for idx in range(54, len(face.slides)):
    for sh in face.slides[idx].shapes:
        if hasattr(sh, 'text') and sh.top < 0.8*914400 and sh.left > 10*914400:
            sh.text = str(idx + 1).zfill(2)
            break

tmp = FACE.with_suffix('.tmp.pptx')
face.save(tmp)
tmp.replace(FACE)

# Replace only the existing route image relationship with Lesson 2 online's
# approved route image; all other manual slide edits and media remain intact.
tmp = FACE.with_suffix('.route.tmp.pptx')
with ZipFile(FACE) as zin, ZipFile(tmp, 'w', ZIP_DEFLATED) as zout, ZipFile(L2_ONLINE) as route_zip:
    route = route_zip.read('ppt/media/image2.png')
    for info in zin.infolist():
        data = route if info.filename == 'ppt/media/image2.png' else zin.read(info.filename)
        zout.writestr(info, data)
tmp.replace(FACE)

with ZipFile(FACE) as z:
    bad = z.testzip()
    if bad: raise ValueError(f'bad zip member: {bad}')
print('slides', len(face.slides))
print('sha256', sha256(FACE.read_bytes()).hexdigest())
