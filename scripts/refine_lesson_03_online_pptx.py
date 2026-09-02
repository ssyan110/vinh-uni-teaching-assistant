from pathlib import Path
import shutil

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "lessons/boya-quasi-intermediate-i/lesson-03/10-design/pptx-draft/online/lesson-03-在线预习.pptx"
OUTPUT = INPUT.with_suffix(".refined.tmp.pptx")


def set_text(shape, text):
    """Replace a text box's text while retaining its first-run styling."""
    tf = shape.text_frame
    sample = None
    for paragraph in tf.paragraphs:
        if paragraph.runs:
            sample = paragraph.runs[0].font
            break
    tf.clear()
    paragraph = tf.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    if sample:
        font = run.font
        font.name = sample.name
        font.size = sample.size
        font.bold = sample.bold
        font.italic = sample.italic
        if sample.color.type is not None:
            font.color.rgb = sample.color.rgb


def shape_by_text(slide, text):
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text == text:
            return shape
    raise ValueError(f"Missing shape text: {text}")


prs = Presentation(str(INPUT))

# Vocabulary slides: remove empty extension labels, then add only explicitly
# supplied extensions. Do not touch the goals slide or any other content.
extensions = {
    5: "扩展：对……有兴趣",
    6: "扩展：又叫做咕咾肉",
    8: "扩展：“寒假”是什么意思？",
    13: "扩展：另外＋名词",
    18: "扩展：重复＋名词／动词：表示同样的内容或动作再次出现或再做一次。",
    20: "扩展：帮＋人＋忙，表示帮助某人做事情、解决问题。例如：帮我忙、帮老师忙。",
    26: "扩展：开＋科目＋课，表示开某种课程。",
    30: "扩展：原因／情况＋，于是＋结果。",
}
for slide_no in range(5, 35):
    slide = prs.slides[slide_no - 1]
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text == "扩展：":
            set_text(shape, extensions.get(slide_no, ""))
            if slide_no in extensions:
                shape.height = int(1.22 * 914400)
            break

set_text(shape_by_text(prs.slides[20 - 1], "词类：—"), "词类：动词")

# Phrase slides: remove the generic situation line and use the requested
# student instruction. Where the user supplied a usage explanation, place it
# in the former situation slot so the established layout remains unchanged.
phrase_updates = {
    40: {"例句：我每周上三节课。": "例句：每次回家，我都会给父母做几道北方菜"},
    41: {
        "情境：谈论中文学习和课堂生活": "扩展：帮＋人＋忙：表示帮助某个人。",
        "例句：周末我帮妈妈做饭。": "例句：做饭的时候，弟弟常常帮妈妈忙。",
        "例句：他帮我拿书。": "例句：我和同学常常互相帮忙",
    },
    42: {"情境：谈论中文学习和课堂生活": "扩展：表示还有其他类似的人或事物，没有全部说出来。"},
    43: {
        "情境：谈论中文学习和课堂生活": "扩展：表示某件事发生的时间。\n用法：动词／句子＋的时候",
    },
    44: {"例句：我记得你的生日。": "例句：记得回家要写作业！"},
    50: {"情境：谈论中文学习和课堂生活": "扩展：可以省略“外”。", "例句：我想买另外一本书。": "例句：我想买另一本书。"},
    51: {"情境：谈论中文学习和课堂生活": "用法：更＋形容词／心理动词，表示在原来的程度上进一步提高或加强，常用于比较。"},
    52: {"情境：谈论中文学习和课堂生活": "用法：动词＋得＋形容词，表示动作做得怎么样。", "例句：她说中文说得很清楚。": "例句：中国人说话说得很快。"},
    54: {"例句：他有兴趣，于是开始学习。": "例句：我对德语很有兴趣，于是我选修德语。"},
    60: {"情境：谈论中文学习和课堂生活": "扩展：互相＋动词：表示双方对对方做同样或相关的动作。", "例句：同学们在课堂上互相学习。": "例句：朋友应该互相关心。"},
}
for slide_no, replacements in phrase_updates.items():
    slide = prs.slides[slide_no - 1]
    for old, new in replacements.items():
        set_text(shape_by_text(slide, old), new)
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text == "完成教材中的练习。":
            set_text(shape, "用这个句式写出三句话")

# All remaining phrase slides use the same requested instruction and no
# generic situation label. Do this by content, not by guessed slide numbers.
for slide in prs.slides:
    for shape in slide.shapes:
        if not hasattr(shape, "text"):
            continue
        if shape.text == "情境：谈论中文学习和课堂生活":
            set_text(shape, "")
        elif shape.text == "完成教材中的练习。":
            set_text(shape, "用这个句式写出三句话")

# Delete the two explicitly requested pages, highest index first. The visible
# slide-number labels are then renumbered to match the resulting deck order.
for index in (61, 59):
    slide_id = prs.slides._sldIdLst[index - 1]
    prs.part.drop_rel(slide_id.rId)
    prs.slides._sldIdLst.remove(slide_id)

for index, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text in {f"{n:02d}" for n in range(1, 100)}:
            if abs(shape.left / 914400 - 12.0) < 0.2 and abs(shape.top / 914400 - 0.25) < 0.2:
                set_text(shape, f"{index:02d}")

prs.save(str(OUTPUT))
shutil.move(str(OUTPUT), str(INPUT))
print(f"saved {INPUT}")
print(f"slides {len(prs.slides)}")
