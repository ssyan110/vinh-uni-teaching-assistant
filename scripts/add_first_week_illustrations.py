from pathlib import Path
from shutil import copy2
from pptx import Presentation
from pptx.util import Inches

root = Path(__file__).resolve().parents[1]
src = root / 'lessons/boya-quasi-intermediate-i/lesson-01/10-design/pptx-draft/first-week-icebreakers/lesson-01-第一周破冰-A-教材延续型-插图版-draft-v2.pptx'
out = src.with_name('lesson-01-第一周破冰-A-教材延续型-插图版-draft-v3.pptx')
asset = root / 'lessons/boya-quasi-intermediate-i/lesson-01/10-design/image-assets-draft/first-week-icebreakers'
copy2(src, out)
prs = Presentation(out)
placements = {
    3: 'node-01-name.png', 5: 'node-03-leaders.png',
    6: 'node-04-listen.png', 7: 'node-05-question.png',
    8: 'node-06-prepare.png', 9: 'node-07-vocab.png',
    10: 'node-08-books.png', 11: 'node-04-listen.png',
}
for slide_no, filename in placements.items():
    slide = prs.slides[slide_no - 1]
    slide.shapes.add_picture(str(asset / filename), Inches(9.5), Inches(3.88), width=Inches(2.81), height=Inches(2.03))
prs.save(out)
print(out)
