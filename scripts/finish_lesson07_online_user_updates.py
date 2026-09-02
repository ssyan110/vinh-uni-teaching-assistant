from hashlib import sha256
from pathlib import Path
import json
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Pt

ROOT = Path('/Users/ssyan110/Development/vinh-uni-teaching-assistant')
DECK = ROOT / 'lessons/boya-quasi-intermediate-i/lesson-07/10-design/pptx-draft/online/lesson-07-在线预习.pptx'
MANIFEST = DECK.parent / 'manifest.json'

def set_text(shape, text, size=None):
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(3); tf.margin_right = Pt(3); tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text; r.font.name = 'KaiTi'; r.font.color.rgb = RGBColor(56,69,86)
    if size is not None: r.font.size = Pt(size)

def main():
    prs = Presentation(DECK)
    # Existing slide numbers are intentionally used here; these pages were
    # retained in place so the user's manual edits remain on the same slides.
    s47 = prs.slides[46]
    for sh in s47.shapes:
        if not hasattr(sh, 'text'): continue
        if sh.text == '我学习中文的经历': set_text(sh, '介绍你喜欢的课外活动，并说说为什么。', 30)
        elif sh.text == '请用三到五句介绍你学习中文的经历。': set_text(sh, '用自己的话介绍一项你喜欢的课外活动。', 22)

    s49 = prs.slides[48]
    for sh in s49.shapes:
        if hasattr(sh, 'text') and sh.text == '请填写课本第30页的表格。': set_text(sh, '请填写课本第66页的表格。', 22)

    s50 = prs.slides[49]
    replacements = {
        '我的个人介绍': ('我的个人介绍', 30),
        '按照三段准备自己的信息：': ('回答下面的问题：', 22),
        '北京印象': ('喜欢的运动', 22),
        '你对北京有什么印象？': ('你喜欢什么运动？', 20),
        '中文学习': ('喜欢的原因', 22),
        '你在哪里学习汉语？': ('你为什么喜欢？', 20),
        '学习比较': ('运动的好处和坏处', 20),
        '哪里学习汉语更有效率？': ('运动有什么好处和坏处？', 20),
        '至少使用课本中的10个词语和5个句式。': ('说10—12句，至少100字。', 24),
    }
    for sh in s50.shapes:
        if hasattr(sh, 'text') and sh.text in replacements:
            text, size = replacements[sh.text]; set_text(sh, text, size)

    # The 12 inserted expression pages move the original final three pages
    # from 51–53 to 63–65.
    for slide, old, new in ((prs.slides[62], '64', '63'), (prs.slides[63], '65', '64'), (prs.slides[64], '66', '65')):
        for sh in slide.shapes:
            if hasattr(sh, 'text') and sh.text == old:
                set_text(sh, new, 12)

    tmp = DECK.with_suffix('.tmp.pptx'); prs.save(tmp); tmp.replace(DECK)
    digest = sha256(DECK.read_bytes()).hexdigest()
    manifest = json.loads(MANIFEST.read_text())
    manifest['status'] = 'draft_user_updates_applied_2026-09-01'
    manifest['output']['sha256'] = digest; manifest['output']['bytes'] = DECK.stat().st_size
    manifest['sha256'] = digest; manifest['bytes'] = DECK.stat().st_size
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'sha256': digest, 'bytes': DECK.stat().st_size, 'slides': len(prs.slides)}, ensure_ascii=False))

if __name__ == '__main__': main()
