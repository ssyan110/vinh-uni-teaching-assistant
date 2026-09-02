from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from pptx import Presentation

ROOT = Path('/Users/ssyan110/Development/vinh-uni-teaching-assistant')
DECK = ROOT / 'lessons/boya-quasi-intermediate-i/lesson-06/10-design/pptx-draft/face-to-face/lesson-06-实体课.pptx'
ONLINE = ROOT / 'lessons/boya-quasi-intermediate-i/lesson-06/10-design/pptx-draft/online/lesson-06-在线预习.pptx'

prs = Presentation(DECK)
online = Presentation(ONLINE)

def texts(slide):
    return [sh.text for sh in slide.shapes if hasattr(sh, 'text') and sh.text.strip()]

def set_exact(slide, old, new):
    hits = [sh for sh in slide.shapes if hasattr(sh, 'text') and sh.text == old]
    if len(hits) != 1:
        raise ValueError(f'{old!r}: expected one shape, got {len(hits)}')
    hits[0].text = new

def set_first(slide, predicate, new):
    hits = [sh for sh in slide.shapes if hasattr(sh, 'text') and predicate(sh.text)]
    if not hits:
        raise ValueError('target text shape not found')
    hits[0].text = new

def copy_content(dst, src, keep_number=None):
    for sh in list(dst.shapes):
        dst.shapes._spTree.remove(sh._element)
    for sh in src.shapes:
        dst.shapes._spTree.append(deepcopy(sh._element))
    if keep_number is not None:
        nums = [sh for sh in dst.shapes if hasattr(sh, 'text') and sh.text in {f'{n:02d}' for n in range(1, 100)}]
        if nums:
            nums[0].text = f'{keep_number:02d}'

def sid_of(slide):
    return next(sid for sid in prs.slides._sldIdLst if prs.part.related_part(sid.rId) is slide.part)

def move_before(slide, anchor):
    ids = prs.slides._sldIdLst
    sid = sid_of(slide)
    ids.remove(sid)
    ids.insert(list(ids).index(sid_of(anchor)), sid)

# Directly reuse the online goals page; retain its current wording.
goals_before = texts(prs.slides[2])
copy_content(prs.slides[2], online.slides[2], keep_number=3)
assert texts(prs.slides[2]) == texts(online.slides[2]), 'goals page is not identical to online page'

# Replace only the route image and retain the face-to-face route-page text.
face_route = prs.slides[1]
online_route = online.slides[1]
for sh in list(face_route.shapes):
    if getattr(sh, 'shape_type', None) == 13:
        face_route.shapes._spTree.remove(sh._element)
route_pic = next(sh for sh in online_route.shapes if getattr(sh, 'shape_type', None) == 13)
face_route.shapes.add_picture(BytesIO(route_pic.image.blob), route_pic.left, route_pic.top, route_pic.width, route_pic.height)

# Lesson-specific warm-up.
s4 = prs.slides[3]
set_exact(s4, '先想一想', '音乐和课外活动')
set_exact(s4, '你以前来过中国吗？', '你平时参加什么课外活动？')
set_exact(s4, '• 你对中国有什么印象？\n• 你想在中国学习什么？\n• 你觉得在哪里学中文更有效率？', '• 你有什么业余爱好？\n• 你喜欢参加什么活动？\n• 你觉得参加活动有什么收获？')
set_exact(s4, '参考句式：我觉得……，因为……', '参考句式：我喜欢……，因为……')

for no, question in {
    9: '你的业余爱好是什么？为什么？',
    10: '你的性格怎么样？',
    11: '你喜欢什么课？不喜欢什么课？为什么？',
    12: '你的假期生活充实吗？请你介绍。',
}.items():
    set_first(prs.slides[no - 1], lambda t: t.startswith('你'), question)

# User-requested speaking tasks, using existing task layouts.
copy_content(prs.slides[18], prs.slides[19], keep_number=19)
set_exact(prs.slides[18], '介绍朴大宇', '介绍大岛的爱好')
set_first(prs.slides[18], lambda t: t.startswith('请你用6—8句话'), '请你介绍大岛的爱好。说6—8个句子，不少于60字。')
set_exact(prs.slides[18], '参考词语：钢琴\u3000合唱节\u3000话剧\u3000交响乐\u3000住院\u3000弹\n常用表达：', '参考词语：专业、硕士生、弹、参观、话剧、语速、正式\n常用表达：相当、和……相比')

set_exact(prs.slides[20], '你说的跟短文（一）哪里不一样？', '介绍你朋友的爱好。')

copy_content(prs.slides[25], prs.slides[27], keep_number=26)
set_exact(prs.slides[25], '介绍短文', '介绍大岛参加合唱团的经过')
set_first(prs.slides[25], lambda t: t.startswith('请你用6—8句话'), '请你介绍大岛参加合唱团的经过。说6—8个句子，不少于60字。')
set_exact(prs.slides[25], '参考词语：钢琴\u3000合唱节\u3000话剧\u3000交响乐\u3000住院\u3000弹\n常用表达：', '参考词语：弹钢琴、交响乐、感兴趣、合唱比赛、住院、代替、答应、第三名、邀请、同意\n常用表达：上……的时候')

copy_content(prs.slides[32], prs.slides[35], keep_number=33)
set_exact(prs.slides[32], '介绍短文', '介绍他参加合唱团以后的情况')
set_first(prs.slides[32], lambda t: t.startswith('请你用6—8句话'), '请你介绍大岛参加合唱团以后的情况。说6—8个句子，不少于60字。')
set_exact(prs.slides[32], '参考词语：钢琴\u3000合唱节\u3000话剧\u3000交响乐\u3000住院\u3000弹\n常用表达：', '参考词语：充实、忙碌、排练、中文、交流、有得必有失\n常用表达：没有那么多时间、有得必有失')

copy_content(prs.slides[34], prs.slides[37], keep_number=35)
set_exact(prs.slides[34], '请你说说', '请说说你的爱好对你的影响。')
set_first(prs.slides[34], lambda t: t.startswith('比较两地学习汉语'), '请说说你的爱好对你的影响。')

# Comprehensive practice based on textbook page 57 and this lesson.
copy_content(prs.slides[49], prs.slides[52], keep_number=50)
set_exact(prs.slides[49], '请你介绍朴大宇', '介绍大岛的爱好、参加合唱团的经历和影响')
set_exact(prs.slides[49], '请填表后，说一说朴大宇在北京的生活、学习和课堂经验。', '请填表后，说一说大岛的爱好、参加合唱团的经历和参加以后生活的变化。')

# Preserve the existing “没有那么多” sentence page before repurposing
# slide 51 as the requested lesson-specific question page.
no_more_source = prs.slides[50]
no_more_copy = prs.slides.add_slide(no_more_source.slide_layout)
for sh in no_more_source.shapes:
    no_more_copy.shapes._spTree.append(deepcopy(sh._element))

copy_content(prs.slides[50], prs.slides[53], keep_number=51)
set_exact(prs.slides[50], '根据课本，回答问题', '根据本课内容，回答问题')
set_exact(prs.slides[50], '• 朴大宇以前来过北京吗？\n• 他这次在北京发现了什么？\n• 他在北京上什么课？\n• 他晚上和谁一起学习？\n• 在韩国和中国学习汉语有什么不同？\n• 为什么说在中国学习效率更高？', '• 大岛的业余爱好是什么？\n• 她为什么参加合唱比赛？\n• 她怎样进入学校合唱团？\n• 参加合唱团以后，她的生活有什么变化？\n• 她参加合唱团后得到什么，又失去了什么？\n• 你觉得参加业余活动有什么影响？')

copy_content(prs.slides[51], prs.slides[54], keep_number=52)
set_exact(prs.slides[51], '请你说说', '介绍自己的爱好')
set_exact(prs.slides[51], '① 我的北京印象', '你喜欢什么？为什么喜欢？')
set_exact(prs.slides[51], '② 我的中文学习', '有什么有趣的事情？')
set_exact(prs.slides[51], '③ 我的学习比较', '说10到12句，至少100字。')
set_first(prs.slides[51], lambda t: t.startswith('必须使用10个课本'), '使用本课学过的词语和常用表达。')
for sh in prs.slides[51].shapes:
    if hasattr(sh, 'text') and sh.text == '____________________________':
        sh.text = ''

# The existing file has no slide 57; make its final personal task exact.
final = prs.slides[54]
set_exact(final, '请你说说', '介绍自己的爱好')
set_exact(final, '① 我的北京印象', '你喜欢什么？为什么喜欢？')
set_exact(final, '② 我的中文学习', '有什么有趣的事情？')
set_exact(final, '③ 我的学习比较', '说10到12句，至少100字。')
set_first(final, lambda t: t.startswith('必须使用10个课本'), '使用本课学过的词语和常用表达。')
for sh in final.shapes:
    if hasattr(sh, 'text') and sh.text == '____________________________':
        sh.text = ''

# Delete the explicitly requested original pages.
for index in (14, 13):
    sid = prs.slides._sldIdLst[index - 1]
    prs.part.drop_rel(sid.rId)
    prs.slides._sldIdLst.remove(sid)

def find_text(text):
    for s in prs.slides:
        if any(hasattr(sh, 'text') and sh.text == text for sh in s.shapes):
            return s
    raise ValueError(text)

# Keep only each short text's own sentence patterns in its section.
short2_divider = find_text('短文（二）')
short3_divider = find_text('短文（三）')
reflection = find_text('介绍大岛的爱好、参加合唱团的经历和影响')
move_before(find_text('相当'), short2_divider)
move_before(find_text('和……相比'), short2_divider)
move_before(find_text('上……的时候'), short3_divider)
move_before(find_text('没有那么多'), reflection)

# Add the missing textbook expression to short text three.
template = find_text('没有那么多')
extra = prs.slides.add_slide(template.slide_layout)
for sh in template.shapes:
    extra.shapes._spTree.append(deepcopy(sh._element))
set_exact(extra, '没有那么多', '有得必有失')
for sh in extra.shapes:
    if hasattr(sh, 'text') and sh.text.startswith('教材 P'):
        sh.text = '教材 P56–57'
move_before(extra, reflection)

assert texts(prs.slides[2]) == texts(online.slides[2]), 'goals changed after edits'

tmp = DECK.with_suffix('.tmp.pptx')
prs.save(tmp)
tmp.replace(DECK)
print(f'slides={len(prs.slides)}')
print(f'sha256={sha256(DECK.read_bytes()).hexdigest()}')
