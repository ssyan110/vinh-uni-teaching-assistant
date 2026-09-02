from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from pptx import Presentation

ROOT = Path('/Users/ssyan110/Development/vinh-uni-teaching-assistant')
DECK = ROOT / 'lessons/boya-quasi-intermediate-i/lesson-06/10-design/pptx-draft/face-to-face/lesson-06-实体课.pptx'
prs = Presentation(DECK)

def find_slide(text):
    for s in prs.slides:
        if any(hasattr(sh, 'text') and sh.text == text for sh in s.shapes):
            return s
    raise ValueError(text)

def sid_of(slide):
    return next(sid for sid in prs.slides._sldIdLst if prs.part.related_part(sid.rId) is slide.part)

src = find_slide('有得必有失')
copy = prs.slides.add_slide(src.slide_layout)
for sh in src.shapes:
    copy.shapes._spTree.append(deepcopy(sh._element))
for sh in copy.shapes:
    if hasattr(sh, 'text') and sh.text == '有得必有失':
        sh.text = '没有那么多'
    elif hasattr(sh, 'text') and sh.text.startswith('教材 P'):
        sh.text = '教材 P56–57'
    elif hasattr(sh, 'text') and sh.text.strip().isdigit() and abs(sh.left/914400-12) < .2:
        sh.text = '56'

anchor = find_slide('介绍大岛的爱好、参加合唱团的经历和影响')
ids = prs.slides._sldIdLst
sid = sid_of(copy)
ids.remove(sid)
ids.insert(list(ids).index(sid_of(anchor)), sid)

tmp = DECK.with_suffix('.tmp.pptx')
prs.save(tmp)
tmp.replace(DECK)
print(f'slides={len(prs.slides)}')
print(f'sha256={sha256(DECK.read_bytes()).hexdigest()}')
