from hashlib import sha256
from pathlib import Path
from pptx import Presentation

ROOT = Path('/Users/ssyan110/Development/vinh-uni-teaching-assistant')
DECK = ROOT / 'lessons/boya-quasi-intermediate-i/lesson-06/10-design/pptx-draft/online/lesson-06-在线预习.pptx'
prs = Presentation(DECK)

def find_exact(text):
    for s in prs.slides:
        if any(hasattr(sh, 'text') and sh.text == text for sh in s.shapes): return s
    raise ValueError(text)

def find_divider(number):
    for s in prs.slides:
        texts = [sh.text for sh in s.shapes if hasattr(sh, 'text')]
        if '句式练习' in texts and f'{number:02d}' in texts and not any(t.startswith('教材 P') for t in texts): return s
    raise ValueError(f'divider {number}')

def sid_of(slide):
    return next(sid for sid in prs.slides._sldIdLst if prs.part.related_part(sid.rId) is slide.part)

def move_before(slide, anchor):
    ids = prs.slides._sldIdLst; sid = sid_of(slide); ids.remove(sid); ids.insert(list(ids).index(sid_of(anchor)), sid)

short3 = find_exact('短文（三）')
short2 = find_exact('短文（二）')
reflection = find_exact('我觉得很难的地方')
move_before(find_exact('相当＋形容词'), short2)
move_before(find_exact('和……相比'), short2)
move_before(find_divider(58), short3)
move_before(find_exact('上……的时候'), short3)
move_before(find_divider(59), reflection)
move_before(find_exact('没有那么多'), reflection)
move_before(find_exact('有得必有失'), reflection)

tmp = DECK.with_suffix('.tmp.pptx'); prs.save(tmp); tmp.replace(DECK)
print(f'slides={len(prs.slides)}')
print(f'sha256={sha256(DECK.read_bytes()).hexdigest()}')
