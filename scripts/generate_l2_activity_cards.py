from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('lessons/boya-quasi-intermediate-i/lesson-02/10-design/support-materials')

def set_fonts(run):
    run.font.name = 'Times New Roman'
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:cs'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'KaiTi')
    r_pr = run._element.get_or_add_rPr()
    lang = r_pr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang'); r_pr.append(lang)
    lang.set(qn('w:val'), 'zh-CN'); lang.set(qn('w:eastAsia'), 'zh-CN')

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_text(cell, text, bold=False):
    cell.text = ''
    p = cell.paragraphs[0]; r = p.add_run(text); r.bold = bold; set_fonts(r); r.font.size = Pt(11)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def style_doc(doc):
    sec = doc.sections[0]; sec.top_margin = Inches(.6); sec.bottom_margin = Inches(.6); sec.left_margin = Inches(.65); sec.right_margin = Inches(.65)
    st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman'); st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman'); st._element.rPr.rFonts.set(qn('w:cs'), 'Times New Roman'); st._element.rPr.rFonts.set(qn('w:eastAsia'), 'KaiTi'); st.font.size = Pt(11)
    lang = st._element.rPr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang'); st._element.rPr.append(lang)
    lang.set(qn('w:val'), 'zh-CN'); lang.set(qn('w:eastAsia'), 'zh-CN')
    for style_name in ['Title','Heading 1','Heading 2']:
        st = doc.styles[style_name]; st.font.name='Times New Roman'; st._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman'); st._element.rPr.rFonts.set(qn('w:cs'),'Times New Roman'); st._element.rPr.rFonts.set(qn('w:eastAsia'),'KaiTi')
        lang = st._element.rPr.find(qn('w:lang'))
        if lang is None:
            lang = OxmlElement('w:lang'); st._element.rPr.append(lang)
        lang.set(qn('w:val'), 'zh-CN'); lang.set(qn('w:eastAsia'), 'zh-CN')

def add_title(doc, title, pages):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(title); r.bold=True; set_fonts(r); r.font.size=Pt(18); r.font.color.rgb=RGBColor(31,60,86)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(f'第二课《王红的一天》｜教材 {pages}'); set_fonts(r); r.font.size=Pt(10)

def add_line_table(doc, headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers): set_cell_text(t.rows[0].cells[i],h,True); set_cell_shading(t.rows[0].cells[i],'D9E7F2')
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row): set_cell_text(cells[i],val)
    return t

def add_prompt(doc, text):
    p=doc.add_paragraph(); r=p.add_run(text); set_fonts(r); r.font.size=Pt(11)

def save(name, builder):
    d=Document(); style_doc(d); builder(d); d.save(OUT/name)

def keywords(d):
    add_title(d,'听力关键词记录表','P15–P19')
    add_prompt(d,'使用方法：听前看题，先写下你猜到的时间、地点、人物和活动；听后补充关键词。记录不要求完整句子，最后和同伴互相确认。')
    add_line_table(d,['听力','时间／地点','人物','活动／事情','我听到的关键词'],[['2-5 王红喜欢上课','________________','________________','________________','________________________________'],['2-6 生日午餐','________________','________________','________________','________________________________'],['2-7 课外活动','________________','________________','________________','________________________________']])
    add_prompt(d,'和同伴确认：你听到的是 ________________________________ 吗？    我需要再听／确认：________________________')
    add_prompt(d,'开放记录没有唯一答案，请根据音频和教材内容核对。')

def dayplan(d):
    add_title(d,'王红一天安排表','P15–P20')
    add_prompt(d,'根据三段短文，写出王红一天的安排。可以先写关键词，再和同伴互相补充。')
    add_line_table(d,['时间','地点','王红做什么','我听到的依据／关键词'],[['早上','________________','____________________________','____________________________'],['上午','________________','____________________________','____________________________'],['中午','________________','____________________________','____________________________'],['下午','________________','____________________________','____________________________']])
    add_prompt(d,'和同伴核对：我们有哪一项不一样？______________________________________________')
    add_prompt(d,'开放记录：表格中的说法可以不同，但要能回到三段短文的内容。')

def school(d):
    add_title(d,'学校生活口语任务卡','P16、P20')
    add_prompt(d,'任务：介绍王红的大学生活，重点说说她上课的情况。先一个人说，再由同伴补充，最后一人总结。')
    add_prompt(d,'第一位同学：说6～8个句子，不少于60字。可参考：开始、每次、收获、周围、环境；对……有了了解、对……熟悉了、尤其、虽然……可是……。')
    add_line_table(d,['先说的人记录关键词','同伴补充的关键词','总结时要说的重点'],[['____________________________','____________________________','____________________________'],['____________________________','____________________________','____________________________'],['____________________________','____________________________','____________________________']])
    add_prompt(d,'同伴可以问：你说的是数学课吗？／王红什么时候开始上大学？／她觉得大学生活怎么样？')
    add_prompt(d,'最后总结：说6～8个句子，不少于60字。开放口语，没有唯一答案；请依据教材内容完成。')

def birthday(d):
    add_title(d,'生日午餐口语任务卡','P17–P18')
    add_prompt(d,'任务：根据《生日午餐》，一起说清楚王红的生日午餐。一个人先说，其他同学补充，最后一人总结。')
    add_prompt(d,'先说的人：说4～6个句子，不少于40字。可参考：摆、插、生日歌、愉快；给……开生日会、顿、开开心心。')
    add_line_table(d,['先说的人记录','其他同学补充','最后总结的重点'],[['________________________','________________________','________________________'],['________________________','________________________','________________________'],['________________________','________________________','________________________']])
    add_prompt(d,'同伴可以问：她们在哪里吃饭？／桌子上摆着什么？／王红感觉怎么样？')
    add_prompt(d,'最后总结：说6～8个句子，不少于60字。开放口语，没有唯一答案；请依据教材内容完成。')

def extracurricular(d):
    add_title(d,'课外活动口语任务卡','P18–P19')
    add_prompt(d,'任务：介绍王红的课外活动。一个人先说，其他同学补充，最后一人总结。')
    add_prompt(d,'先说的人：说4～6个句子，不少于40字。可参考：志愿者、翻译、布置、回答、讲解员、增长、服务；自从……以来、每……都……、不但……同时……。')
    add_line_table(d,['先说的人记录','其他同学补充','最后总结的重点'],[['________________________','________________________','________________________'],['________________________','________________________','________________________'],['________________________','________________________','________________________']])
    add_prompt(d,'同伴可以问：王红每星期二下午去哪儿？／她今天做什么？／这个生日她觉得怎么样？')
    add_prompt(d,'最后总结：说6～8个句子，不少于60字。开放口语，没有唯一答案；请依据教材内容完成。')

OUT.mkdir(parents=True, exist_ok=True)
save('lesson-02-听力关键词记录表-draft.docx', keywords)
save('lesson-02-王红一天安排表-draft.docx', dayplan)
save('lesson-02-学校生活口语任务卡-draft.docx', school)
save('lesson-02-生日午餐口语任务卡-draft.docx', birthday)
save('lesson-02-课外活动口语任务卡-draft.docx', extracurricular)
