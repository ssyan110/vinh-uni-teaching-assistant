from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
import hashlib
import os
import shutil


ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "lessons/boya-quasi-intermediate-i/lesson-03/10-design/pptx-draft/online/lesson-03-在线预习.pptx"
L1_ROUTE = ROOT / "lessons/boya-quasi-intermediate-i/lesson-01/10-design/image-assets-draft/learning-route-user-supplied-transparent.png"
ROUTE = ROOT / "lessons/boya-quasi-intermediate-i/lesson-03/10-design/assets/lesson-03-learning-route.png"
TMP = DECK.with_suffix(".route-goals.tmp.pptx")


def make_route():
    image = Image.open(L1_ROUTE).convert("RGB")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("/Library/Fonts/Microsoft/Kaiti.ttf", 46)
    # Retain the approved Lesson 1 visual family and change only the steps
    # that differ in Lesson 3's online-preparation sequence.
    cards = [
        (265, 585, 465, 700, image.getpixel((365, 720)), (0, 128, 128), "练习听力"),
        (735, 585, 935, 700, image.getpixel((835, 720)), (195, 125, 0), "读懂短文"),
    ]
    for x1, y1, x2, y2, fill, color, label in cards:
        draw.rectangle((x1, y1, x2, y2), fill=fill)
        box = draw.textbbox((0, 0), label, font=font)
        tw = box[2] - box[0]
        th = box[3] - box[1]
        draw.text(((x1 + x2 - tw) / 2, 618 - box[1]), label, font=font, fill=color, stroke_width=1, stroke_fill=color)
    ROUTE.parent.mkdir(parents=True, exist_ok=True)
    image.save(ROUTE, format="PNG")


def style_goals():
    prs = Presentation(str(DECK))
    slide = prs.slides[2]
    # Copy the actual Lesson 1 body shapes, including the colored number
    # circles and divider lines, then restore the current Lesson 3 wording.
    l1 = Presentation(str(ROOT / "lessons/boya-quasi-intermediate-i/lesson-01/20-approved/pptx/lesson-01-在线预习.pptx"))
    l1_body = list(l1.slides[2].shapes)[3:]
    for shape in list(slide.shapes)[3:]:
        slide.shapes._spTree.remove(shape._element)
    for shape in l1_body:
        slide.shapes._spTree.append(deepcopy(shape._element))
    wording = {
        "听懂关于家庭、工作和爱好的主要信息。": "听懂李大为学习中文经历中的主要信息。",
        "用本课词语介绍自己的家庭、学习／工作和爱好。": "说明自己为什么学习或选修中文。",
        "回答和讨论跟课本主题有关的问题。": "回答并讨论家庭、学习和中文课堂问题。",
        "根据要求写出并准备一段个人介绍。": "能够介绍个人的生活和经历。",
    }
    for shape in slide.shapes:
        if not hasattr(shape, "text"):
            continue
        for old, new in wording.items():
            if shape.text == old:
                shape.text_frame.paragraphs[0].runs[0].text = new
                break
    prs.save(str(TMP))
    os.replace(TMP, DECK)


def replace_route_media():
    tmp = DECK.with_suffix(".media.tmp.pptx")
    with ZipFile(DECK) as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    route = ROUTE.read_bytes()
    replaced = False
    with ZipFile(tmp, "w", ZIP_DEFLATED) as target:
        for info, data in entries:
            if info.filename == "ppt/media/image2.png":
                data = route
                replaced = True
            target.writestr(info, data)
    if not replaced:
        raise RuntimeError("ppt/media/image2.png not found")
    os.replace(tmp, DECK)


make_route()
style_goals()
replace_route_media()
with ZipFile(DECK) as z:
    assert z.testzip() is None
    assert hashlib.sha256(z.read("ppt/media/image2.png")).hexdigest() == hashlib.sha256(ROUTE.read_bytes()).hexdigest()
print(f"saved {DECK}")
print(f"route {ROUTE} sha256={hashlib.sha256(ROUTE.read_bytes()).hexdigest()}")
print(f"deck sha256={hashlib.sha256(DECK.read_bytes()).hexdigest()}")
