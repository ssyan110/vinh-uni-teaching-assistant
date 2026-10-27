#!/usr/bin/env python3
"""Build one online draft from its explicitly approved, structured slide copy."""
from __future__ import annotations

import argparse
import json
import math
import os
import re
import zipfile
import xml.etree.ElementTree as ET

from PIL import Image
from pathlib import Path

import production_gate
from lesson_context import resolve_lesson_context
from workflow_integrity import resolve_relative_path, sha256
from verify_pptx_audio_policy import verify as verify_audio_policy

ROOT = Path(__file__).resolve().parents[1]


def nonempty(value):
    return isinstance(value, str) and bool(value.strip())


def build(root, lesson_key, offering, class_scope, artifact_id, approval_record):
    root = Path(root).resolve()
    context = resolve_lesson_context(root, lesson_key, offering)
    output = context.design_root / 'pptx-draft' / 'online' / offering / artifact_id
    check = production_gate.check_lesson_ppt_draft(lesson_key, str(output), offering, approval_record, artifact_id, class_scope)
    if check['status'] != 'ready':
        raise ValueError('; '.join(check['blockers']))
    record_path, _ = resolve_relative_path(root, approval_record, required_root=context.design_root / 'approvals')
    record = json.loads(record_path.read_text(encoding='utf-8'))
    review_path, review_rel = resolve_relative_path(root, record['review_path'], required_root=context.design_root)
    spec = json.loads(review_path.read_text(encoding='utf-8'))
    identity = {'lesson_key': lesson_key, 'offering_id': offering, 'class_scope': class_scope, 'artifact_id': artifact_id}
    if not isinstance(spec, dict) or any(spec.get(k) != v for k, v in identity.items()):
        raise ValueError('approved slide specification identity does not match the requested version')
    slides = spec.get('slides')
    if not isinstance(slides, list) or not slides:
        raise ValueError('approved slide specification must contain slides')
    if len(slides) < 2 or not all(isinstance(x, dict) for x in slides) or slides[0].get('kind') != 'cover' or slides[1].get('kind') != 'learning_flow':
        raise ValueError('online presentation requires a cover and learning flow')
    reference = spec.get('layout_reference')
    if not isinstance(reference, dict):
        raise ValueError('online production requires an actual approved PPTX layout reference; free-form coordinates are not a template')
    reference_path, reference_rel = resolve_relative_path(root, reference.get('path'), required_root=context.design_root)
    if not reference_path.is_file() or sha256(reference_path) != reference.get('sha256'):
        raise ValueError('online PPTX layout reference is missing or changed')
    if not any(x.get('path') == reference_rel and x.get('sha256') == reference['sha256'] for x in record['artifacts']):
        raise ValueError('online layout reference must be bound to the scoped approval inputs')
    from pptx import Presentation
    from render_reference_online_pptx import render
    template = Presentation(reference_path)
    reference['path'] = str(reference_path)
    resources = {}
    example_bank = None
    chrome_minimum_pt = 23 if offering == '2026-fall' and context.textbook_id == 'boya-elementary-i' else 20
    for index, slide in enumerate(slides, 1):
        if not isinstance(slide, dict) or not nonempty(slide.get('kind')) or not isinstance(slide.get('elements'), list) or not slide['elements']:
            raise ValueError(f'slide {index} requires kind and visible elements')
        if slide.get('notes') is not None and not isinstance(slide['notes'], str):
            raise ValueError(f'slide {index} notes must be text')
        if any(not isinstance(e, dict) or e.get('type') not in ('text', 'image', 'shape', 'table') for e in slide['elements']):
            raise ValueError(f'slide {index} contains an unsupported element')
        page_number = slide.get('reference_slide')
        if isinstance(page_number, bool) or not isinstance(page_number, int) or not 1 <= page_number <= len(template.slides):
            raise ValueError(f'slide {index} must map to an actual reference slide')
        native = template.slides[page_number - 1]
        if any(s.shape_type not in (1, 5, 9, 13, 17, 19) for s in native.shapes):
            raise ValueError(f'slide {index} reference has a grouped, chart or media object that this native copier does not support')
        slots = {str(s.shape_id): s for s in native.shapes}
        mapped = [str(e.get('reference_shape_id')) for e in slide['elements']]
        removed = set(map(str, slide.get('remove_shape_ids', [])))
        if len(mapped) != len(set(mapped)) or set(mapped) - set(slots) or removed - set(slots) or removed.intersection(mapped):
            raise ValueError(f'slide {index} has invalid or repeated reference shape bindings')
        required = {sid for sid, s in slots.items() if (s.has_text_frame and s.text.strip()) or s.has_table} - removed
        if required - set(mapped):
            raise ValueError(f'slide {index} leaves reference text unreplaced')
        for e in slide['elements']:
            shape = slots[str(e['reference_shape_id'])]
            if (e['type'] == 'text' and not shape.has_text_frame) or (e['type'] == 'image' and shape.shape_type != 13):
                raise ValueError(f'slide {index} reference slot has the wrong type')
            if e['type'] == 'table':
                cells = e.get('cells')
                if (not shape.has_table or not isinstance(cells, list) or len(cells) != len(shape.table.rows)
                        or any(not isinstance(row, list) or len(row) != len(shape.table.columns)
                               or any(not isinstance(text, str) for text in row) for row in cells)):
                    raise ValueError(f'slide {index} must replace every native table cell')
        refs = slide.get('source_refs', [])
        if not isinstance(refs, list) or any(not nonempty(ref) for ref in refs):
            raise ValueError(f'slide {index} has invalid source references')
        audio = slide.get('audio', [])
        if not isinstance(audio, list) or any(not isinstance(n, str) or not re.fullmatch(r'Audio [0-9]+-[0-9]+', n) for n in audio):
            raise ValueError(f'slide {index} has invalid audio labels')
        visible_audio = re.findall(r'Audio [0-9]+-[0-9]+', ' '.join(e.get('text', '') for e in slide['elements'] if isinstance(e, dict) and isinstance(e.get('text'), str)))
        if sorted(visible_audio) != sorted(audio):
            raise ValueError(f'slide {index} visible audio labels do not match source mapping')
        for label in audio:
            if not any(e.get('type') == 'text' and e.get('text') == label and e.get('x', 0) >= 9 and e.get('y', 8) <= 1 for e in slide['elements'] if isinstance(e, dict)):
                raise ValueError(f'slide {index} audio must be plain text in the upper right')
        if slide.get('kind') == 'vocabulary':
            vocab = slide.get('vocabulary')
            if not isinstance(vocab, dict) or any(not nonempty(vocab.get(key)) for key in ('word', 'pinyin', 'part_of_speech', 'meaning_vi')):
                raise ValueError(f'slide {index} vocabulary requires word, pinyin, part of speech and Vietnamese meaning')
            if example_bank is None:
                bank_path, bank_rel = resolve_relative_path(root, spec.get('example_bank_path'), required_root=context.design_root)
                if bank_rel not in {item['path'] for item in record['artifacts']}:
                    raise ValueError('versioned example bank must be included in content approval')
                example_bank = json.loads(bank_path.read_text(encoding='utf-8'))
                global_bank_path, _ = resolve_relative_path(root, 'course/boya-example-bank.json', required_root=root / 'course')
                global_bank = json.loads(global_bank_path.read_text(encoding='utf-8'))
            examples = vocab.get('examples')
            if not isinstance(examples, list) or len(examples) != 2 or examples[0] == examples[1] or any(not nonempty(x) or len(re.findall(r'[\u3400-\u9fff]', x)) > 24 for x in examples):
                raise ValueError(f'slide {index} vocabulary requires two short examples')
            if example_bank.get('lessons', {}).get(lesson_key, {}).get('vocabulary', {}).get(vocab['word']) != examples or global_bank.get('lessons', {}).get(lesson_key, {}).get('vocabulary', {}).get(vocab['word']) != examples:
                raise ValueError(f'slide {index} vocabulary examples do not match the version snapshot and course example bank')
            visible = ' '.join(str(e.get('text', '')) for e in slide['elements'] if isinstance(e, dict))
            if any(vocab[key] not in visible for key in ('word', 'pinyin', 'part_of_speech', 'meaning_vi')) or any(x not in visible for x in examples) or not any(e.get('type') == 'image' for e in slide['elements'] if isinstance(e, dict)):
                raise ValueError(f'slide {index} vocabulary content and image must be visible')
        for element in slide['elements']:
            if not isinstance(element, dict) or element.get('type') not in ('text', 'image', 'shape', 'table'):
                raise ValueError(f'slide {index} contains an unsupported element')
            for field in ('x', 'y', 'w', 'h'):
                v = element.get(field)
                if isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v):
                    raise ValueError(f'slide {index} element has invalid {field}')
            if element['x'] < 0 or element['y'] < 0 or element['w'] <= 0 or element['h'] <= 0 or element['x'] + element['w'] > 13.333 or element['y'] + element['h'] > 7.5:
                raise ValueError(f'slide {index} element is outside the slide')
            for field in ('color', 'line_color', 'fill_color'):
                if field in element and (not isinstance(element[field], str) or not re.fullmatch(r'[0-9A-Fa-f]{6}', element[field])):
                    raise ValueError(f'slide {index} contains an invalid color')
            if element['type'] in ('text', 'table'):
                size = element.get('font_size')
                role = element.get('text_role', 'content')
                if role not in ('content', 'chrome'):
                    raise ValueError(f'slide {index} has an invalid text role')
                if role == 'chrome' and (element['type'] != 'text' or not (
                    (element['y'] < 0.7 and (re.fullmatch(r'第\d+(?:[—–-]\d+)?课[｜|].+', element['text']) or re.fullmatch(r'\d{1,3}', element['text'])))
                    or (element['x'] >= 9 and element['y'] < 1.3 and re.fullmatch(r'Audio \d+-\d+', element['text']))
                    or (element['y'] >= 6.5 and re.fullmatch(r'教材 P\d+(?:[–—-]\d+)?', element['text']))
                )):
                    raise ValueError(f'slide {index} chrome is only for the page header, number, audio or textbook marker')
                minimum_pt = chrome_minimum_pt if role == 'chrome' else 25
                if ((element['type'] == 'text' and not nonempty(element.get('text')))
                        or isinstance(size, bool) or not isinstance(size, (int, float)) or not math.isfinite(size) or size < minimum_pt or size > 96):
                    raise ValueError(f'slide {index} requires visible text at least {minimum_pt} pt')
                if element.get('font', 'cjk') not in ('cjk', 'latin'):
                    raise ValueError('font must be cjk or latin')
            if element['type'] == 'image':
                img, rel = resolve_relative_path(root, element.get('path'), required_root=context.design_root)
                if not img.is_file() or img.suffix.lower() not in ('.png', '.jpg', '.jpeg'):
                    raise ValueError(f'slide {index} image must be a local PNG or JPEG')
                data = img.read_bytes()
                if not (data.startswith(b'\x89PNG\r\n\x1a\n') or data.startswith(b'\xff\xd8\xff')):
                    raise ValueError(f'slide {index} image is invalid')
                try:
                    with Image.open(img) as decoded:
                        decoded.load()
                        if decoded.format not in ('PNG', 'JPEG'):
                            raise ValueError('unsupported image format')
                except (OSError, Image.DecompressionBombError) as exc:
                    raise ValueError(f'slide {index} image cannot be decoded: {rel}') from exc
                resources[rel] = sha256(img)
                element['path'] = str(img)
            if element['type'] == 'shape' and element.get('shape', 'rect') not in ('rect', 'roundRect', 'ellipse', 'line'):
                raise ValueError('unsupported shape')
    if output.exists():
        raise FileExistsError('this version output already exists; choose a new artifact ID')
    # Recheck the exact inputs immediately before creating the version directory.
    check = production_gate.check_lesson_ppt_draft(lesson_key, str(output), offering, approval_record, artifact_id, class_scope)
    if check['status'] != 'ready':
        raise ValueError('; '.join(check['blockers']))
    output.mkdir(parents=True, exist_ok=False)
    pptx_path = output / f'{context.lesson_id}-在线预习.pptx'
    render(spec, pptx_path)
    if not pptx_path.is_file() or not pptx_path.stat().st_size:
        raise ValueError('renderer did not produce a PPTX')
    # PptxGenJS applies one font to every script; set the project font mapping per run.
    rewritten = output / 'font-mapped.pptx.tmp'
    with zipfile.ZipFile(pptx_path) as original, zipfile.ZipFile(rewritten, 'x') as result:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename.startswith('ppt/slides/slide') and info.filename.endswith('.xml'):
                xml = data.decode('utf-8')
                for name, font in (('latin', 'Times New Roman'), ('ea', 'KaiTi'), ('cs', 'Times New Roman')):
                    xml = re.sub(r'(<a:' + name + r'\b[^>]*typeface=")[^"]*(")', lambda m: m.group(1) + font + m.group(2), xml)
                data = xml.encode('utf-8')
            result.writestr(info, data)
    os.replace(rewritten, pptx_path)
    if verify_audio_policy(pptx_path):
        raise ValueError('PPTX violates plain audio reference policy')
    # Verify the actual editable slide texts and forbid external or audio links.
    with zipfile.ZipFile(pptx_path) as package:
        if package.testzip():
            raise ValueError('PPTX ZIP integrity failed')
        names = set(package.namelist())
        ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        for index, item in enumerate(slides, 1):
            name = f'ppt/slides/slide{index}.xml'
            if name not in names:
                raise ValueError('PPTX slide order/count does not match approved copy')
            xml = ET.fromstring(package.read(name))
            actual = ''.join((node.text or '') for node in xml.findall('.//a:t', ns)).replace('\n', '')
            expected = ''.join((e['text'] if e['type'] == 'text' else ''.join(x for row in e['cells'] for x in row)).replace('\n', '')
                               for e in item['elements'] if e['type'] in ('text', 'table'))
            if actual != expected:
                raise ValueError(f'PPTX slide {index} text differs from approved copy')
        if f'ppt/slides/slide{len(slides) + 1}.xml' in names:
            raise ValueError('PPTX has unexpected extra slides')
        for name in names:
            if name.endswith('.rels'):
                rels = ET.fromstring(package.read(name))
                for rel in rels:
                    if rel.get('TargetMode') == 'External' or '/audio' in (rel.get('Type') or ''):
                        raise ValueError('PPTX contains an external or audio relationship')
    # Inputs changed during generation require a fresh run; no manifest is written.
    check = production_gate.check_lesson_ppt_draft(lesson_key, str(output), offering, approval_record, artifact_id, class_scope)
    if check['status'] != 'ready' or any(sha256(root / p) != h for p, h in resources.items()):
        raise ValueError('approved inputs or images changed while generating')
    manifest = {**identity, 'approval_record': approval_record, 'review_path': review_rel,
                'review_sha256': sha256(review_path), 'slides': len(slides),
                'layout_reference': {'path': reference_rel, 'sha256': sha256(reference_path)},
                'layout_map': [{'slide': i, 'reference_slide': s['reference_slide']} for i, s in enumerate(slides, 1)],
                'pptx': {'path': pptx_path.relative_to(root).as_posix(), 'sha256': sha256(pptx_path)},
                'resources': [{'path': p, 'sha256': h} for p, h in sorted(resources.items())],
                'source_map': [{'slide': i, 'source_refs': item.get('source_refs', []), 'audio': item.get('audio', [])} for i, item in enumerate(slides, 1)],
                'status': 'draft_not_approved', 'visual_qa': 'pending', 'playback_qa': 'pending'}
    (output / 'build-manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return manifest


def main():
    p = argparse.ArgumentParser(description=__doc__)
    for name in ('lesson-key', 'offering-id', 'class-scope', 'artifact-id', 'approval-record'):
        p.add_argument('--' + name, required=True)
    a = p.parse_args()
    try:
        print(json.dumps(build(ROOT, a.lesson_key, a.offering_id, a.class_scope, a.artifact_id, a.approval_record), ensure_ascii=False, indent=2))
    except (ValueError, OSError, zipfile.BadZipFile) as exc:
        p.exit(1, f'Blocked: {exc}\n')


if __name__ == '__main__':
    main()
