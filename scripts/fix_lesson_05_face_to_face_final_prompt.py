from hashlib import sha256
from pathlib import Path
import json

from pptx import Presentation


ROOT = Path('/Users/ssyan110/Development/vinh-uni-teaching-assistant')
DECK = ROOT / 'lessons/boya-quasi-intermediate-i/lesson-05/10-design/pptx-draft/face-to-face/lesson-05-实体课.pptx'


def set_text(shape, value):
    runs = [run for para in shape.text_frame.paragraphs for run in para.runs]
    if runs:
        runs[0].text = value
        for run in runs[1:]:
            run.text = ''
    else:
        shape.text = value


def main():
    prs = Presentation(DECK)
    if len(prs.slides) != 56:
        raise ValueError(f'预期实体课为56页，实际为{len(prs.slides)}页')
    slide = prs.slides[55]
    text_shapes = [s for s in slide.shapes if hasattr(s, 'text')]
    title = next(s for s in text_shapes if s.text == '介绍我喜欢的一个人')
    set_text(title, '介绍我喜欢的一个人')
    prompt = next(s for s in text_shapes if s.text == '你喜欢谁？')
    set_text(prompt, '你喜欢谁？为什么喜欢？介绍一下这个人。\n说10到12句，至少100字。')
    for shape in text_shapes:
        if shape is title or shape is prompt:
            continue
        if shape.text.startswith('教材 ') or shape.text == '第五课｜我的音乐老师' or shape.text == '61' or not shape.text.strip():
            continue
        shape.text_frame.clear()
    tmp = DECK.with_suffix('.prompt.tmp.pptx')
    prs.save(tmp)
    tmp.replace(DECK)
    manifest = DECK.parent / 'manifest.json'
    if manifest.is_file():
        data = json.loads(manifest.read_text())
        data['sha256'] = sha256(DECK.read_bytes()).hexdigest()
        data['bytes'] = DECK.stat().st_size
        data['status'] = 'draft_user_refinement_2026-09-01'
        manifest.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')
    print(f'updated {DECK}')
    print(f'sha256 {sha256(DECK.read_bytes()).hexdigest()}')


if __name__ == '__main__':
    main()
