#!/usr/bin/env python3
"""Remove production metadata from L2/L3 student-facing DOCX drafts.

The lesson Markdown files are the wording source for cards generated from
Markdown.  This small pass only removes internal identity/status notes from
the two pre-study cards that were created by an earlier hand-built draft; it
does not alter textbook wording, prompts, answers, page references, or task
actions.  Outputs remain in each lesson's ``10-design/support-materials``
directory.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]


def all_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def set_text(paragraph, text: str) -> None:
    # These paragraphs are simple body text.  Replacing the text in place
    # keeps the existing paragraph style and document-level KaiTi/TNR rules.
    paragraph.text = text


def drop_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def clean_l2_prestudy(path: Path) -> None:
    document = Document(path)
    for paragraph in list(all_paragraphs(document)):
        text = paragraph.text.strip()
        if not text:
            continue
        if text == "课次身份：boya-quasi-intermediate-i:lesson-02　教材页码：P12–P21":
            set_text(paragraph, "教材页码：P12–P21")
        elif text.startswith("草案状态："):
            drop_paragraph(paragraph)
        elif "（草案）" in text:
            set_text(paragraph, text.replace("（草案）", ""))
    document.save(path)


def clean_l3_prestudy(path: Path) -> None:
    document = Document(path)
    for paragraph in list(all_paragraphs(document)):
        text = paragraph.text.strip()
        if not text:
            continue
        if text == "lesson_key：`boya-quasi-intermediate-i:lesson-03`":
            drop_paragraph(paragraph)
        elif text == "使用说明":
            drop_paragraph(paragraph)
        elif text.startswith("本卡依据第三课来源包编写"):
            set_text(paragraph, "课前完成以下阅读、听力和口语准备。")
        elif text.startswith("状态：`draft`"):
            drop_paragraph(paragraph)
        elif text.startswith("教师检查：") or text.startswith("**教师检查**："):
            drop_paragraph(paragraph)
        elif "（草案）" in text:
            set_text(paragraph, text.replace("（草案）", ""))
    document.save(path)


def main() -> None:
    l2 = ROOT / "lessons/boya-quasi-intermediate-i/lesson-02/10-design/support-materials"
    l3 = ROOT / "lessons/boya-quasi-intermediate-i/lesson-03/10-design/support-materials"
    clean_l2_prestudy(l2 / "lesson-02-预习卡-draft.docx")
    clean_l3_prestudy(l3 / "lesson-03-预习卡-draft.docx")
    for path in (
        l2 / "lesson-02-课末检查-draft.docx",
        l2 / "lesson-02-综合口语评量表-draft.docx",
        l3 / "lesson-03-课末检查-draft.docx",
        l3 / "lesson-03-评量表-draft.docx",
        l3 / "lesson-03-活动01-听力证据记录-draft.docx",
        l3 / "lesson-03-活动02-三段短文信息表-draft.docx",
        l3 / "lesson-03-活动03-小组总结与个人迁移-draft.docx",
    ):
        if not path.is_file():
            raise FileNotFoundError(path)
    print("cleaned L2/L3 pre-study cards")


if __name__ == "__main__":
    main()
