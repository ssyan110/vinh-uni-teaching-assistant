#!/usr/bin/env python3
"""Build the 2026-fall 11-week semester teacher manual."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT = PROJECT_ROOT / "course/offerings/2026-fall/semester-teacher-manual.docx"

LATIN_FONT = "Times New Roman"
CJK_FONT = "KaiTi"
INK = "17324D"
MUTED = "64748B"
BLUE = "2E6F95"
CORAL = "C66A56"
MINT = "DCEFE9"
GOLD = "FBF1D9"
LIGHT = "F2F6F6"
LINE = "DCE3E6"


LESSONS = [
    ("第1课", "丽丽是独生女", "P1–P11", "1-1–1-6", "第1周", "家庭信息交换与家庭说明"),
    ("第2课", "王红的一天", "P12–P21", "2-1–2-6", "第2周", "日程交换与共同安排"),
    ("第3课", "我对学中文越来越有兴趣", "P22–P31", "3-1–3-6", "第3周", "学习兴趣访谈与学习安排"),
    ("第4课", "在中国学汉语", "P32–P41", "4-1–4-6", "第4周", "学习咨询与建议说明"),
    ("第5课", "我的音乐老师", "P42–P50", "5-1–5-6", "第5周", "人物／兴趣推荐"),
    ("第6课", "大圣参加了学校的合唱团", "P51–P58", "6-1–6-6", "第6周", "团队活动方案与阶段检核"),
    ("第7课", "小张热爱登山", "P59–P67", "7-1–7-6", "第7周", "户外活动与安全建议"),
    ("第8课", "孙子和《孙子兵法》", "P68–P76", "8-1–8-6", "第8周", "文化材料概括与策略案例"),
    ("第9课", "北方菜和南方菜", "P77–P86", "9-1–9-6", "第9周", "比较、推荐与理由"),
    ("第10课", "中国人喜欢聚餐", "P87–P95", "10-1–10-6", "第9周", "聚餐共同计划与变化回应"),
    ("第11课", "原来他们是关心我", "P96–P104", "11-1–11-6", "第10周", "澄清误解与回应关心"),
    ("第12课", "散步", "P105–P113", "12-1–12-6", "第10周", "生活／健康计划与建议"),
]


WEEKS = [
    {
        "week": "第1周", "dates": "09/07–09/13", "title": "第1课《丽丽是独生女》",
        "function": "从家庭与身份信息开始，练习听懂人物关系、交换个人信息和澄清理解。",
        "prep": "打开教材P1–P11；听1-1至1-6；记录人物、关系、数量或事件等关键词；准备3–5句家庭介绍。",
        "rows": [
            ("第1节", "课程说明；熟悉主题口语诊断；回收预习；第一次听力", "先让学生完成熟悉主题口语，再进入教材听力；未预习学生走3–5分钟恢复路线", "主旨、关键词、一个证据；教师初始观察记录"),
            ("第2节", "核对听力信息；比较答案；人物关系口头改述", "以教材题目为中心核对；只修补影响理解的词语或句子", "个人听力表；人物关系与关键细节"),
            ("第3节", "家庭信息交换与澄清", "一人说明、一人提问；要求至少一次确认或澄清", "访谈记录；追问和确认"),
            ("第4节", "60秒家庭说明；反馈与重做；出口卡", "反馈聚焦一项高影响问题；安排学生重新说一遍", "说明初稿、重做版、出口卡"),
        ],
    },
    {
        "week": "第2周", "dates": "09/14–09/20", "title": "第2课《王红的一天》",
        "function": "根据听到的生活安排整理顺序、交换日程，并协商一个共同计划。",
        "prep": "打开教材P12–P21；听2-1至2-6；记录时间、活动和先后关系；准备自己一天中的三个活动。",
        "rows": [
            ("第1节", "个人日程回收；听力主旨和细节", "先看题目、抓关键词，再听；让同伴互相补充信息", "日程关键词表；主旨和两个细节"),
            ("第2节", "时间与活动顺序口头复述", "按教材题目核对；把答案改说成完整句", "日程复述；时间追问回应"),
            ("第3节", "带限制安排一天", "两人交换不同日程限制；要求说明一个取舍", "共同日程方案；提问、确认和协商"),
            ("第4节", "介绍方案；反馈与重做", "引导听者追问时间安排；保留一次改说", "60–90秒说明；修订表达"),
        ],
    },
    {
        "week": "第3周", "dates": "09/21–09/27", "title": "第3课《我对学中文越来越有兴趣》",
        "function": "听懂兴趣和变化，说明个人学习经历与原因，并提出一项可执行的学习安排。",
        "prep": "打开教材P22–P31；听3-1至3-6；标记兴趣、变化或原因信息；准备自己的学习经历。",
        "rows": [
            ("第1节", "主旨、变化和原因听辨", "比较课前预测；要求学生用关键词支持判断", "听力证据表；变化与支持信息"),
            ("第2节", "教材练习口头回答、复述和改述", "语言修补只处理影响表达的问题", "个人口头回答；原因或变化说明"),
            ("第3节", "学习兴趣访谈", "询问开始学习、困难和现在的变化；听者追问", "访谈表；两个问题和一个追问"),
            ("第4节", "说明一项学习安排", "同伴提出可执行性问题；学生根据反馈重做", "60–90秒安排；反馈和修订版"),
        ],
    },
    {
        "week": "第4周", "dates": "09/28–10/04", "title": "第4课《在中国学汉语》",
        "function": "理解学习经历和学习选择，向同伴说明情况并提出有依据的建议。",
        "prep": "打开教材P32–P41；听4-1至4-6；记录地点、经历、困难或选择；准备一个想咨询的学习问题。",
        "rows": [
            ("第1节", "听前预测；地点、人物和问题", "让小组比较听力证据，不先讲解全文", "主旨、关键词和细节证据"),
            ("第2节", "教材理解题转为同伴问答", "用短时间修补高影响表达；要求补充信息", "问答记录；完整句回答"),
            ("第3节", "学习者咨询任务", "根据不同需求提出建议；听者说明是否接受", "建议卡；原因、条件和回应"),
            ("第4节", "小组分享与个人选择说明", "安排反馈和一次根据听者问题的改说", "60秒建议说明；改说证据"),
        ],
    },
    {
        "week": "第5周", "dates": "10/05–10/11", "title": "第5课《我的音乐老师》",
        "function": "整理人物信息、描述重要人物或兴趣，并面向听众完成简短推荐。",
        "prep": "打开教材P42–P50；听5-1至5-6；记录人物特点、事件和说话者态度；准备一位重要人物或一项兴趣。",
        "rows": [
            ("第1节", "人物、事件和态度听辨", "根据教材图片和题目预测；听后回到证据", "人物信息表；主旨和两个细节"),
            ("第2节", "教材练习核对；人物复述", "两人互相复述并指出遗漏；修补影响听说的语言", "人物复述；“为什么重要”追问回应"),
            ("第3节", "人物／兴趣访谈", "听者选择一个信息继续追问；交换角色", "访谈记录；提问、追问和确认"),
            ("第4节", "人物或兴趣推荐", "同伴指出一项清楚处和一项可改进处；重做", "60–90秒推荐初稿与重做版"),
        ],
    },
    {
        "week": "第6周", "dates": "10/12–10/18", "title": "第6课《大圣参加了学校的合唱团》＋阶段检核",
        "function": "听懂校园团队活动信息，协商共同安排，并用阶段性表现检验前半学期听说能力。",
        "prep": "打开教材P51–P58；听6-1至6-6；记录参加团队、时间、人物和安排；准备一次集体活动经历或愿望。",
        "rows": [
            ("第1节", "参加活动、人物和安排听辨", "先抓关键词，再完成教材听力理解和同伴核对", "主旨、细节和一个依据"),
            ("第2节", "团队活动介绍", "把教材信息改说成介绍；修补时间、顺序和参加表达", "口头复述；澄清问题回应"),
            ("第3节", "团队活动方案", "根据成员、时间或任务限制安排排练／活动", "小组共同方案；协商和确认"),
            ("第4节", "阶段听说检核与重做", "新短听力＋双人互动＋60–90秒说明；反馈后重做", "阶段表现记录；个人改进目标"),
        ],
    },
    {
        "week": "第7周", "dates": "10/19–10/25", "title": "第7课《小张热爱登山》",
        "function": "理解户外活动经历和条件，说明计划并提出安全或准备建议。",
        "prep": "打开教材P59–P67；听7-1至7-6；记录人物、地点、时间、条件和结果；准备一个户外活动问题。",
        "rows": [
            ("第1节", "事件顺序和关键条件", "听前预测；要求学生比较不同听力证据", "事件顺序表；主旨、细节和结果"),
            ("第2节", "条件与结果口头说明", "完成教材理解练习并口头核对", "口头复述；条件与结果关系"),
            ("第3节", "户外活动顾问", "根据时间、天气或能力限制决定路线和准备物品", "路线／准备方案；询问和确认"),
            ("第4节", "计划或建议说明", "同伴提出一个风险问题；反馈后重做", "60–90秒说明；风险回应"),
        ],
    },
    {
        "week": "第8周", "dates": "10/26–11/01", "title": "第8课《孙子和〈孙子兵法〉》",
        "function": "理解文化材料的主旨和观点，用自己的话概括，并把一个可核对的观点应用到学习或团队问题。",
        "prep": "打开教材P68–P76；听8-1至8-6；记录人物、事件、观点和例子；标记一个想讨论的问题。",
        "rows": [
            ("第1节", "人物、事件和观点听辨", "先看题目和标题；小组比较支持判断的依据", "文化材料主旨与细节"),
            ("第2节", "文化材料概括和问答", "区分教材信息与学生自己的推论", "一段口头概括；信息来源说明"),
            ("第3节", "策略案例讨论", "把材料中的一个观点用于学习、团队或日常问题", "案例卡；问题、做法和理由"),
            ("第4节", "小组短报告和追问", "听众提出一个追问；报告者澄清或修订一次", "1–2分钟报告；追问回应和重做"),
        ],
    },
    {
        "week": "第9周", "dates": "11/02–11/08", "title": "第9课《北方菜和南方菜》＋第10课《中国人喜欢聚餐》",
        "function": "在同一周内完成饮食比较和聚餐安排，练习根据对象、口味或条件做推荐并处理变化。",
        "prep": "打开教材P77–P95；听9-1至9-6、10-1至10-6；各记录一组比较信息和一项聚餐习惯；准备一个推荐理由。",
        "rows": [
            ("第1节", "第9课比较听力与教材题目", "引导学生找出比较信息并回到听力证据", "比较表；两项差异和一个依据"),
            ("第2节", "第9课菜单顾问", "按对象、口味或预算推荐；处理不同意见", "推荐说明；比较、理由和回应"),
            ("第3节", "第10课聚餐信息理解", "记录人物、场合和安排；核对教材理解题", "聚餐信息表；主旨、细节和回应"),
            ("第4节", "聚餐主办任务", "共同安排聚餐并处理一个临时变化；快速反馈", "共同计划；变化回应；90秒小组说明"),
        ],
    },
    {
        "week": "第10周", "dates": "11/09–11/15", "title": "第11课《原来他们是关心我》＋第12课《散步》",
        "function": "先处理关心与误解，再把生活和健康主题迁移到共同计划，练习澄清、回应和建议。",
        "prep": "打开教材P96–P113；听11-1至11-6、12-1至12-6；记录误解、关心、活动和安排信息；准备一个生活情境。",
        "rows": [
            ("第1节", "第11课关系、误解和关心", "让学生记录依据，再两人核对并改述", "关系信息图；误解说明"),
            ("第2节", "关心对话诊所", "根据角色信息重建对话；安排澄清、回应和重说", "澄清对话；确认、改述和修补"),
            ("第3节", "第12课生活或健康安排", "记录散步、生活或健康信息；口头核对教材题目", "计划信息表；主旨、细节和建议"),
            ("第4节", "生活计划任务与期末提纲", "根据对象需求安排计划；个人说明并回答追问", "60–90秒说明；期末提纲初稿"),
        ],
    },
    {
        "week": "第11周", "dates": "11/16–11/22", "title": "期末整合表现、学习档案与补做",
        "function": "在新听力材料和新任务限制下综合使用本学期的理解、互动、表达和修补策略。",
        "prep": "整理全学期听力记录、口语样本、反馈和重做；准备期末个人口语提纲。",
        "rows": [
            ("第1节", "新准中级听力理解", "先看题、抓关键词；听后回答主旨、细节、态度和依据", "Interpretive记录；答案和听力证据"),
            ("第2节", "综合互动任务", "给出新情境和限制；观察提问、确认、不同意见回应", "互动观察记录；共同方案"),
            ("第3节", "个人成段表现", "每人完成2–3分钟表达并回答至少一个追问", "Presentational口语；任务、可理解度、组织和回应"),
            ("第4节", "反馈、重做、档案和补做", "按rubric反馈；完成重做、反思和缺席补做", "期末重做版；学习档案；下阶段建议"),
        ],
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=LINE, size="6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, size=12, bold=False, color=INK, italic=False) -> None:
    run.font.name = CJK_FONT
    run.font.size = Pt(size)
    # KaiTi has no separate bold face in the LibreOffice preview runtime;
    # size, color and table shading provide hierarchy without losing CJK glyphs.
    run.font.bold = False
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:eastAsia"), CJK_FONT)
    r_fonts.set(qn("w:ascii"), LATIN_FONT)
    r_fonts.set(qn("w:hAnsi"), LATIN_FONT)
    r_fonts.set(qn("w:cs"), LATIN_FONT)
    r_fonts.set(qn("w:hint"), "eastAsia")


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.18) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_text(paragraph, text, size=12, bold=False, color=INK, italic=False):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return run


def add_para(doc, text="", size=12, bold=False, color=INK, align=None, before=0, after=6, line=1.18):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    set_paragraph_spacing(p, before=before, after=after, line=line)
    if text:
        add_text(p, text, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=14 if level == 1 else 10, after=6, line=1.1)
    add_text(p, text, size=16 if level == 1 else 14, bold=True, color=INK)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_paragraph_spacing(p, after=3, line=1.12)
    add_text(p, text, size=11.5)
    return p


def set_table_widths(table, widths):
    for row in table.rows:
        for index, width in enumerate(widths):
            if index < len(row.cells):
                row.cells[index].width = Inches(width)


def format_table(table, header=True, font_size=10.5, header_fill=LIGHT):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        if header and row_index == 0 and tr_pr.find(qn("w:tblHeader")) is None:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            set_cell_border(cell)
            if header and row_index == 0:
                set_cell_shading(cell, header_fill)
            for paragraph in cell.paragraphs:
                set_paragraph_spacing(paragraph, after=2, line=1.08)
                for run in paragraph.runs:
                    set_run_font(run, size=font_size, bold=(header and row_index == 0), color=INK)


def add_table(doc, headers, rows, widths=None, font_size=10.5, header_fill=LIGHT):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cells[index].text = str(value)
    format_table(table, font_size=font_size, header_fill=header_fill)
    if widths:
        set_table_widths(table, widths)
    return table


def add_callout(doc, label, text, fill=GOLD, accent=CORAL):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=accent, size="12")
    set_cell_margins(cell, top=140, start=160, bottom=140, end=160)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=0, line=1.15)
    add_text(p, label, size=11.5, bold=True, color=accent)
    add_text(p, text, size=11.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def set_document_defaults(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = CJK_FONT
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
    for style_name, size, color in (("Title", 22, INK), ("Heading 1", 16, INK), ("Heading 2", 14, INK), ("Heading 3", 12.5, BLUE)):
        style = styles[style_name]
        style.font.name = CJK_FONT
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
        style._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_header_footer(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(header, "荣市大学华语听说中级课程 · 2026年秋季学期", size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "《博雅汉语听说：准中级加速篇 I》· 整学期教学计划与教案 · 第", size=9, color=MUTED)
    add_page_field(footer)
    add_text(footer, "页", size=9, color=MUTED)


def build():
    doc = Document()
    set_document_defaults(doc)
    add_header_footer(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title, before=25, after=6, line=1.05)
    add_text(title, "荣市大学", size=22, bold=True, color=INK)
    title2 = doc.add_paragraph()
    title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title2, after=7, line=1.05)
    add_text(title2, "华语听说中级课程", size=21, bold=True, color=INK)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(subtitle, after=9, line=1.05)
    add_text(subtitle, "2026年秋季学期整学期教学计划与教案", size=19, bold=True, color=CORAL)
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(intro, after=22, line=1.2)
    add_text(intro, "教材：《博雅汉语听说：准中级加速篇 I》｜11周｜44课时", size=13, color=MUTED)

    add_callout(doc, "排课基准：", "2026/09/07–2026/11/22，共11周；每周一次、每次4节、每节50分钟，课堂教学共44课时。线上预习为课外自主学习，不抵扣课堂课时。", fill=MINT, accent=BLUE)
    add_para(doc, "", after=2)
    add_table(doc, ["项目", "内容"], [
        ("开课学校", "荣市大学"),
        ("开课实例", "2026-fall"),
        ("课程名称", "华语听说中级课程"),
        ("课程性质／类别", "语言技能课程；具体性质和类别待学校正式课程资料确认"),
        ("授课对象", "荣市大学华语学习学生；具体年级和行政班待学校确认"),
        ("先修课程", "已完成《博雅汉语听说：初级起步篇》第一、二册"),
        ("教材范围", "《准中级加速篇 I》12课；正文印刷P1–P113；来源盘点72段音频，每课6段"),
        ("课堂总学时", "44课时；每课时50分钟；约2200分钟课堂教学时间"),
        ("课外学习", "学生自行完成阅读、听音频、关键词记录和口语准备，不替代或缩减课堂课时"),
        ("文件状态", "11周整学期教学计划与教案修订版；课程代码、学分、正式课表和审核栏待补"),
        ("编制日期", "2026/08/30"),
    ], widths=[1.7, 5.55], font_size=10.5)
    add_heading(doc, "一、课程定位与教学设计", 1)
    add_para(doc, "本课程面向已完成初级汉语听说学习的学生，使用《博雅汉语听说：准中级加速篇 I》开展听力理解、互动口语和成段表达训练。课堂以教材内容为输入，以真实任务为组织方式；学生先接触材料，再通过问答、信息交换、比较、推荐、建议和报告完成口语产出。", size=11.5)
    add_table(doc, ["项目", "本课程的实施要求"], [
        ("教学重点", "从听力中抓主旨、关键细节、态度和依据；在熟悉主题中完成提问、追问、确认、澄清和协商；根据对象、目的和限制完成成段表达。"),
        ("教学难点", "把零散听力信息组织成可说内容；在理解不一致时及时修补；在第9、10周每周两课的压缩安排中保留听力证据、互动任务和个人口语。"),
        ("教学原则", "语言形式服务于任务完成。教师先观察学生表现，再对影响理解或表达的语言问题进行短时修补；每节直接讲解原则上不超过5分钟。"),
        ("学生学习方式", "课外自主预习，课堂完成听力核对、同伴互动、个人表达、反馈和重做；线上学习不替代实体课时。"),
    ], widths=[1.45, 5.8], font_size=10.3, header_fill=LIGHT)

    add_heading(doc, "二、课程目标与 Can-Do 达成要求", 1)
    add_para(doc, "本课程结束时，学生应能把听到或读到的信息用于交流，而不是只复述词语和句式。以下目标既是授课依据，也是阶段检核和期末整合表现的观察依据。", size=11.5)
    add_table(doc, ["课程目标", "目标描述", "主要考核证据"], [
        ("课程目标1：理解", "听懂教材及同主题新材料的主旨、关键细节、人物关系、时间顺序、原因结果和说话者态度，并记录支持答案的依据。", "听力记录、教材题目回答、新听力检核"),
        ("课程目标2：互动", "在家庭、日常、学习、兴趣、文化、饮食和健康等主题中提问、回答、追问、确认、澄清，并在有限条件下与同伴共同做出决定。", "双人／小组任务、互动观察记录、共同方案"),
        ("课程目标3：表达", "根据对象和目的完成个人说明、经历叙述、比较、推荐、建议或计划，组织信息、给出理由并回应听者追问。", "60–90秒口语、2–3分钟期末口语、任务提纲"),
        ("课程目标4：修正", "根据教师或同伴反馈重做一段口语，说明自己的修改，并整理听力记录、口语样本、反馈和反思。", "反馈表、重做版本、学习档案"),
    ], widths=[1.45, 4.1, 1.7], font_size=9.4, header_fill=MINT)
    add_para(doc, "具体 Can-Do 与最低证据要求：", size=11.5, bold=True, before=7, after=4)
    add_table(doc, ["模式", "具体 Can-Do（学生能够完成的任务）", "课堂证据与最低要求"], [
        ("Interpretive 1", "听一段教材音频或同主题新听力，先看3–5个问题、抓关键词；听1–2遍后写出主旨和至少4条细节信息（人物、时间、原因、结果或态度）。", "完成5个问题，至少4题能在听力记录中指出依据。"),
        ("Interpretive 2", "读／听一段短文后，不逐句翻译，按照时间、人物关系、原因结果或两类比较整理3–5条信息，并用2–3句话说出重点。", "信息记录表加2–3句口头概括；同伴能根据概括复述主要内容。"),
        ("Interpersonal 1", "在家庭、日程、学习、兴趣、饮食、文化或健康主题的双人任务中连续交流4–6分钟；提出至少3个问题，回答同伴问题，并追问至少1次。", "信息差记录表；教师或同伴能听到问题、回答和追问。"),
        ("Interpersonal 2", "面对不同意见或信息不清时，用确认、改述或补充说明修补交流；根据任务限制比较选项，与同伴共同做出一个选择并说出理由。", "互动观察记录；对话中至少出现1次确认／修补和1次共同决定。"),
        ("Presentational 1", "根据个人提纲，完成60–90秒的家庭、日常、学习经历或兴趣说明；交代主题，按先后或因果说出至少3条信息，并回答1个追问。", "现场口语或录音；内容完整、顺序清楚，能回应追问。"),
        ("Presentational 2", "针对一个对象和目的，完成2–3分钟的比较、推荐、建议或计划；比较两个选择，给出至少2个理由，说明一项条件变化下的调整方法，并回答1个追问。", "成段口语与任务提纲；听者能说出你的选择、理由和调整办法。"),
        ("修补与重做", "根据教师或同伴反馈，把一段60–90秒口语重做一次，指出自己修改的1项内容，并把听力记录、口语样本、反馈和反思放入学习档案。", "原始版、反馈表、重做版和一句反思记录。"),
    ], widths=[1.2, 4.35, 1.7], font_size=9.2, header_fill=MINT)

    add_heading(doc, "三、教学方法、课堂组织与时间结构", 1)
    add_table(doc, ["教学方法", "课堂实施"], [
        ("能力导向教学", "以 Interpretive、Interpersonal、Presentational 三类表现组织学习证据；语言知识在任务前后按需要处理。"),
        ("任务型教学", "每次任务交代角色、信息差或限制，以及学生要留下的记录、决定或口语产出。"),
        ("翻转学习", "学生课外阅读教材、接触音频、记录关键词并准备个人资料；课堂把时间用于理解、互动、表达和重做。"),
        ("形成性评价", "每周收集听力记录、互动观察、个人口语、同伴反馈、重做和出口记录；反馈聚焦一项高影响问题。"),
        ("学习支持", "未预习学生使用3–5分钟恢复路线；恢复后回到全班同一任务，不把整节课改成逐项讲解。"),
    ], widths=[1.45, 5.8], font_size=10.1, header_fill=LIGHT)
    add_para(doc, "每次实体课的时间结构：", size=11.5, bold=True, before=7, after=4)
    add_table(doc, ["节次", "时间", "教学环节", "教师执行要点", "学生学习成果"], [
        ("第1节", "0–50分钟", "预习回收与第一次听力", "检查预习证据；先看题、抓关键词；未预习学生走恢复路线后回到同一任务。", "预测、关键词、主旨和一个不懂处"),
        ("第2节", "50–100分钟", "听力核对与教材理解", "先让学生回答和比较，再修补影响理解的问题；短讲原则上不超过5分钟。", "听力记录、教材题目回答、改述"),
        ("第3节", "100–150分钟", "双人／小组信息差任务", "明确角色、限制、共同产出和听者回应；观察提问、追问、确认和协商。", "访谈记录、共同方案或任务决定"),
        ("第4节", "150–200分钟", "个人表现、反馈与重做", "每位学生完成一次成段表达；反馈后保留一次重做和出口记录。", "60秒以上口语、反馈、重做、出口卡"),
    ], widths=[0.65, 1.0, 1.55, 2.65, 1.65], font_size=9.4, header_fill=MINT)

    add_heading(doc, "四、课程内容、周次与学时分配", 1)
    add_para(doc, "本学期按11周、每周一次、每次4课时编排，共44课时。第9周和第10周各安排两课，课前预习承担材料接触，实体课保留听力证据、互动任务和口语表现。", size=11.5)
    schedule_rows = [
        ("第1周", "09/07–09/13", "第1课《丽丽是独生女》", "课程诊断；家庭信息听说；60秒家庭说明；一次重做", "预习第2课"),
        ("第2周", "09/14–09/20", "第2课《王红的一天》", "日程信息交换；共同安排一天；说明取舍", "预习第3课"),
        ("第3周", "09/21–09/27", "第3课《我对学中文越来越有兴趣》", "听力证据；学习兴趣访谈；原因和变化说明", "预习第4课"),
        ("第4周", "09/28–10/04", "第4课《在中国学汉语》", "学习经历理解；同伴咨询；建议说明", "预习第5课"),
        ("第5周", "10/05–10/11", "第5课《我的音乐老师》", "人物信息整理；访谈；60–90秒推荐", "预习第6课"),
        ("第6周", "10/12–10/18", "第6课《大圣参加了学校的合唱团》；阶段检核", "团队活动信息交换；共同方案；阶段听说检核", "预习第7课"),
        ("第7周", "10/19–10/25", "第7课《小张热爱登山》", "经历与条件听辨；路线或准备方案；安全建议", "预习第8课"),
        ("第8周", "10/26–11/01", "第8课《孙子和〈孙子兵法〉》", "文化材料主旨与观点；策略案例说明；回应追问", "预习第9课"),
        ("第9周", "11/02–11/08", "第9课《北方菜和南方菜》＋第10课《中国人喜欢聚餐》", "比较和推荐；聚餐共同计划；处理临时变化", "预习第11、12课"),
        ("第10周", "11/09–11/15", "第11课《原来他们是关心我》＋第12课《散步》", "澄清误解；回应关心；安排健康休闲计划", "整理学习档案"),
        ("第11周", "11/16–11/22", "期末整合表现、学习档案与补做", "新听力；综合互动；2–3分钟个人表达；反思与重做", "完成学期反思"),
    ]
    add_table(doc, ["周次", "日期", "主要教学内容", "课堂产出与评价证据", "课外衔接"], schedule_rows, widths=[0.7, 1.15, 2.2, 2.25, 1.0], font_size=9.3)

    add_heading(doc, "五、每周授课教案", 1)
    add_para(doc, "各周表格中的“学生学习成果”是当周可回收的课堂证据。教师可根据班级进度微调活动顺序，但不得删除本周听力理解、互动口语或个人表达的核心证据。", size=11.5)
    period_times = ["0–50分钟", "50–100分钟", "100–150分钟", "150–200分钟"]
    for index, week in enumerate(WEEKS):
        add_heading(doc, f"{week['week']}｜{week['title']}｜{week['dates']}", 2)
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=4, line=1.15)
        add_text(p, "本周教学目标：", size=11.5, bold=True, color=CORAL)
        add_text(p, week["function"], size=11.5)
        add_callout(doc, "课前学习：", week["prep"], fill=GOLD, accent=CORAL)
        weekly_rows = [
            (row[0], period_times[row_index], row[1], row[2], row[3])
            for row_index, row in enumerate(week["rows"])
        ]
        add_table(doc, ["课时", "时间", "教学内容", "教学活动与教师组织", "学生学习成果／评价证据"], weekly_rows, widths=[0.65, 1.0, 1.65, 2.75, 1.75], font_size=9.0, header_fill=LIGHT)
        if index in (3, 7):
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    add_heading(doc, "六、教材内容覆盖与课内外学习安排", 1)
    add_para(doc, "教材页码和音频编号来自当前教材来源盘点。具体词语、句式、教材题号、答案状态以及音频与板块的对应关系，须在各课 lesson_key 来源包审核后用于正式课次材料。", size=11.5)
    add_table(doc, ["课次", "课名", "教材印刷页", "音频", "计划周次", "主要表现任务"], LESSONS, widths=[0.65, 2.0, 0.9, 0.8, 0.8, 2.0], font_size=9.3)
    add_heading(doc, "课外自主预习", 2)
    for item in [
        "打开指定教材页，先看题目、图片和板块标题，预测人物、地点、问题或观点。",
        "按音频编号接触本课6段音频；第一次抓主旨，之后记录人物、时间、原因、结果、态度或关键动作。",
        "标记最多3个听不清、看不懂或不知道如何使用的词语／句子。",
        "准备个人资料、一个问题或3–5句短回答，带到实体课使用。",
        "线上学习由学生自行安排，不设固定分钟数，不要求计时，也不抵扣实体课时。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "实体课衔接", 2)
    for item in [
        "先回收预习证据，再进入听力理解和口语任务；课堂不因少数学生未预习而改为整节重讲教材。",
        "未预习学生使用3–5分钟恢复路线；恢复后回到全班同一任务。",
        "语言形式只在影响理解或任务完成时进行短修补；每节直接讲解原则上不超过5分钟。",
        "每周保留个人听力证据、互动观察、个人口语、反馈重做和出口记录。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "七、考核、成绩评定与学习档案", 1)
    add_para(doc, "以下为本课程的建议评价结构。若院系已有正式评分规定，应以院系规定替换比例，同时保留听力理解、互动口语、成段表达和反馈重做四类证据。", size=11.5)
    add_table(doc, ["考核项目", "建议比例", "考核内容与提交证据"], [
        ("形成性课堂表现", "30%", "每周听力记录、互动观察、短讲、任务产出和反馈后重做。"),
        ("阶段听说检核", "20%", "第6周新听力、双人互动、60–90秒说明和一次重做。"),
        ("期末整合表现", "30%", "第11周新听力、综合互动、2–3分钟个人口语和追问回应。"),
        ("学习档案与预习证据", "20%", "预习记录、听力证据、口语样本、同伴反馈、反思与修订。"),
    ], widths=[1.55, 0.9, 4.95], font_size=10.0, header_fill=MINT)
    add_heading(doc, "口语观察维度", 2)
    add_table(doc, ["观察维度", "判定问题"], [
        ("任务完成", "是否完成说明、询问、比较、推荐、建议或共同决定；是否留下任务要求的记录或方案。"),
        ("可理解度与必要准确度", "听者能否理解主要意思；学生能否修补影响理解的词语、句子或信息。"),
        ("互动与修补", "是否会回答、追问、确认、澄清和回应不同意见；是否能让交流继续。"),
        ("内容组织与情境适切度", "是否交代对象和目的；是否有清楚顺序、理由、条件和结尾。"),
    ], widths=[1.55, 5.7], font_size=10.2, header_fill=LIGHT)
    add_callout(doc, "答案政策：", "开放题和迁移任务使用“教师示例／评分面向”，不把教师示例写成教材唯一标准答案。", fill=GOLD, accent=CORAL)

    add_heading(doc, "八、教材、教学资源与质量保障", 1)
    add_table(doc, ["资源类别", "使用说明"], [
        ("核心教材", "《博雅汉语听说：准中级加速篇 I》；按印刷页码P1–P113和各课来源包使用。"),
        ("音频资源", "按来源盘点编号使用；正式交付前分别确认来源、解码、语义核对和课堂播放状态。"),
        ("课堂材料", "学生预习卡、听力关键词记录表、访谈表、信息差卡、角色卡、调查表、报告提纲、同伴反馈表和出口卡。"),
        ("教师记录", "每周保存预习证据、听力记录、互动观察、口语样本、反馈、重做版本和阶段改进目标。"),
        ("质量保障", "课前核对页码、音频和材料；课后检查学生证据；每6周进行一次阶段检核；期末依据 Can-Do 证据完成整合评价。"),
    ], widths=[1.55, 5.7], font_size=10.2, header_fill=LIGHT)

    add_heading(doc, "九、教学执行、调课与风险处理", 1)
    add_table(doc, ["情况", "处理原则"], [
        ("正式课表变化", "若学校规定的周次、每周次数或每次课时不同，按正式课表平移日期和每次内部时间，不删除学生口语证据。"),
        ("学生未预习", "使用3–5分钟恢复路线，完成必要的关键词和题目准备后回到同一课堂任务。"),
        ("第9、10周进度不足", "优先保留听力主旨与细节、互动任务、个人口语和反馈重做；教材逐项讲解可移至课外预习。"),
        ("音频或来源待确认", "只使用已登记且可测试的播放方式；来源、解码、语义和实播状态分别记录，不把未核对材料写成已完成。"),
        ("缺席或补做", "依据同一 Can-Do 要求完成听力记录、互动替代任务、个人口语和反思；补做证据进入学习档案。"),
    ], widths=[1.55, 5.7], font_size=10.2, header_fill=LIGHT)

    add_heading(doc, "十、开课前审核与签字记录", 1)
    add_para(doc, "以下栏目用于学校、院系或教研室审阅。当前未获确认的信息留空，不以推测内容代替正式记录。", size=11.5)
    add_table(doc, ["审核项目", "姓名／意见", "日期"], [
        ("编制人／任课教师", "____________________________", "____________"),
        ("教研室或课程负责人审核", "____________________________", "____________"),
        ("院系审核", "____________________________", "____________"),
        ("学校正式课表确认", "每周____次；每次____课时；每课时____分钟", "____________"),
        ("课程代码、学分和课程类别", "____________________________", "____________"),
        ("教材来源、音频和教学材料审核", "____________________________", "____________"),
    ], widths=[2.25, 3.8, 1.2], font_size=10.0, header_fill=MINT)
    add_callout(doc, "归档说明：", "本文件为11周排课修订版。正式归档前，应补齐学校课程信息和正式课表，并完成逐课来源审核、练习 coverage、音频播放和教师执行确认。", fill=GOLD, accent=CORAL)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
