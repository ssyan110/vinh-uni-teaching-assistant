#!/usr/bin/env python3
"""Build the Lesson 1 student performance rubric draft.

The draft is written only under lessons/boya-intermediate-i/lesson-01/10-design/support-draft.
It uses the four 0–3 dimensions already defined in the approved teacher guide
and does not change the authority or release package.
"""

from __future__ import annotations

import os
import sys
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_ROOT = PROJECT_ROOT / "lessons" / "boya-intermediate-i" / "lesson-01" / "10-design" / "support-draft"
os.environ["BOYA_LESSON_DRAFT_ROOT"] = str(OUTPUT_ROOT)
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

from production_gate import assert_ready  # noqa: E402


assert_ready("support", OUTPUT_ROOT)

import build_lesson_01_support_materials as support  # noqa: E402

# Use the project-wide standard CJK face for the editable rubric.
support.CJK_FONT = support.CONFIG.get("font_policy", {}).get("cjk", "KaiTi")

from build_lesson_01_support_materials import (  # noqa: E402
    DARK_ACCENT,
    INK,
    LIGHT_BLUE,
    add_body,
    add_heading,
    add_label_box,
    add_meta,
    add_title,
    new_document,
    set_cell_border,
    set_cell_shading,
    set_cell_margins,
    set_paragraph,
    set_run_font,
    set_table_geometry,
)
from docx.enum.section import WD_ORIENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402


ASSESSMENT_DIR = OUTPUT_ROOT / "assessment"


RUBRIC_ROWS = [
    (
        "诠释理解",
        "没有说出相关信息，也没有听力或文本依据。",
        "说出零散信息，但不能说明从哪里听到或读到。",
        "说出主要信息或一个关键细节，并指出一项听力或文本依据。",
        "说清主要信息和关键细节，能准确指出听力或文本依据。",
    ),
    (
        "人际互动",
        "不能回答，也不能让对话继续。",
        "能回答简单问题，但不能追问、澄清或接住同伴的回答。",
        "能回答并完成一次追问；听不清时能请同伴再说一次。",
        "能根据对方的回答追问、澄清、确认或换一种说法，让对话继续。",
    ),
    (
        "表达呈现",
        "不能完成口语表达。",
        "只能说出片段，信息不完整，听者很难理解。",
        "能按顺序说出主要信息，有例子或理由，听者基本听得懂。",
        "结构清楚，理由或例子具体；能根据听者反应补充或换一种说法。",
    ),
    (
        "任务完成",
        "没有完成角色、条件或产出要求。",
        "完成一部分，但漏掉一项关键要求。",
        "完成任务的主要要求，留下规定的口语或记录证据。",
        "完成全部主要要求，并根据反馈改进或重做。",
    ),
]


def set_landscape(document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(0.55)


def add_rubric_table(document, rubric_rows) -> None:
    headers = ["面向", "0 未完成", "1 部分完成", "2 基本达标", "3 清楚并能调整"]
    table = document.add_table(rows=1 + len(rubric_rows), cols=len(headers))
    set_table_geometry(table, [1800, 3000, 3000, 3000, 3000], indent_dxa=0)

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_border(cell, color="9FB3C8", size=8)
        set_cell_margins(cell, top=120, start=120, bottom=120, end=120)
        paragraph = cell.paragraphs[0]
        paragraph.text = ""
        set_paragraph(paragraph, after=0, line_spacing=1.05, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        run = paragraph.add_run(header)
        set_run_font(run, size=10.5, color=DARK_ACCENT, bold=True)

    for row_index, row_data in enumerate(rubric_rows, start=1):
        row = table.rows[row_index]
        for column_index, value in enumerate(row_data):
            cell = row.cells[column_index]
            set_cell_shading(cell, "F8FAFC" if column_index == 0 else "FFFFFF")
            set_cell_border(cell, color="CBD5E1", size=7)
            set_cell_margins(cell, top=130, start=130, bottom=130, end=130)
            paragraph = cell.paragraphs[0]
            paragraph.text = ""
            set_paragraph(
                paragraph,
                after=0,
                line_spacing=1.08,
                alignment=WD_ALIGN_PARAGRAPH.CENTER if column_index == 0 else WD_ALIGN_PARAGRAPH.LEFT,
            )
            run = paragraph.add_run(value)
            set_run_font(run, size=10.5, color=DARK_ACCENT if column_index == 0 else INK, bold=column_index == 0)

    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build_student_rubric() -> Path:
    document = new_document()
    set_landscape(document)
    add_title(
        document,
        "第一课《中国人的姓名》学生表现评量表",
        "用于听、说、问、答、说明与重做",
    )
    add_meta(
        document,
        [
            ("姓名", "________________________"),
            ("日期", "________________________"),
            ("任务", "________________________"),
            ("观察者", "________________________"),
        ],
    )
    add_label_box(
        document,
        "怎么用",
        "先完成任务，再根据实际听到和说出的表现选择一个等级。一个词读错不等于整项低分；要看听者能不能理解、对话能不能继续、任务有没有完成。",
    )
    add_heading(document, "本次评分面向", 2)
    add_body(
        document,
        "□ 诠释理解　　□ 人际互动　　□ 表达呈现　　□ 任务完成　　\n"
        "本次分数：________／________",
        after=6,
    )
    add_rubric_table(document, RUBRIC_ROWS[:2])
    add_heading(document, "评量标准（继续）", 1, page_break_before=True)
    add_rubric_table(document, RUBRIC_ROWS[2:])
    add_heading(document, "表现证据", 2)
    add_body(document, "我听到／看到的一项具体证据：____________________________________________________________")
    add_body(document, "我做得比较好的一点：________________________________________________________________")
    add_body(document, "下一次我要先改进：__________________________________________________________________")
    add_heading(document, "反馈与重做", 2)
    add_body(document, "同伴或教师给我的一句话：_______________________________________________________________")
    add_body(document, "我根据反馈重做以后，改进了：□ 信息　□ 理由／证据　□ 互动　□ 可理解度　□ 任务完成")
    add_body(document, "重做后的新证据：____________________________________________________________________")

    path = ASSESSMENT_DIR / "第一课学生表现评量表.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def main() -> None:
    path = build_student_rubric()
    print(path)


if __name__ == "__main__":
    main()
