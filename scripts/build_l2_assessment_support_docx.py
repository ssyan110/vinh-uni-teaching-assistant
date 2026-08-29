#!/usr/bin/env python3
"""Render Lesson 2 assessment/check Markdown drafts as editable DOCX files.

The Markdown files remain the design drafts and the only text source.  This
script only writes the two DOCX siblings in the lesson's ``10-design``
support-materials directory; it never touches approved or release paths.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SUPPORT = ROOT / "lessons/boya-quasi-intermediate-i/lesson-02/10-design/support-materials"

INK = RGBColor(0x17, 0x32, 0x4D)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
MUTED = RGBColor(0x64, 0x74, 0x8B)
LIGHT = "E8EEF5"
NOTE = "FFF5D6"


def set_run_font(run, size: float = 11, *, bold: bool = False, italic: bool = False,
                 color: RGBColor | None = None) -> None:
    """Set both East Asian and Latin font slots on every generated run."""

    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), "KaiTi")
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:cs"), "Times New Roman")
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "KaiTi")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    normal_lang = normal._element.rPr.find(qn("w:lang"))
    if normal_lang is None:
        normal_lang = OxmlElement("w:lang")
        normal._element.rPr.append(normal_lang)
    normal_lang.set(qn("w:val"), "zh-CN")
    normal_lang.set(qn("w:eastAsia"), "zh-CN")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before, after in (
        ("Heading 1", 17, 0, 10),
        ("Heading 2", 13, 14, 6),
        ("Heading 3", 12, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE if name != "Heading 3" else INK
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "KaiTi")
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
        style_lang = style._element.rPr.find(qn("w:lang"))
        if style_lang is None:
            style_lang = OxmlElement("w:lang")
            style._element.rPr.append(style_lang)
        style_lang.set(qn("w:val"), "zh-CN")
        style_lang.set(qn("w:eastAsia"), "zh-CN")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def inline_parts(text: str):
    """Yield visible Markdown text and a small formatting hint.

    The drafts currently use inline code for the compound lesson key.  The
    backticks are syntax, not student-facing content, so they are removed
    while the text itself is retained exactly.
    """

    # Keep this intentionally narrow: no rewriting or punctuation conversion.
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|_[^_]+_)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            yield text[pos:match.start()], False, False
        token = match.group(0)
        if token.startswith("`") and token.endswith("`"):
            yield token[1:-1], False, False
        elif token.startswith("**") and token.endswith("**"):
            yield token[2:-2], True, False
        elif token.startswith("_") and token.endswith("_"):
            yield token[1:-1], False, True
        pos = match.end()
    if pos < len(text):
        yield text[pos:], False, False


def add_inline(paragraph, text: str, *, size: float = 11, color=None,
               italic: bool = False) -> None:
    for value, bold, emphasis in inline_parts(text):
        if not value:
            continue
        run = paragraph.add_run(value)
        set_run_font(run, size=size, bold=bold, italic=italic or emphasis, color=color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 110,
                     bottom: int = 90, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_borders(table, color: str = "C7D2DE", size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table_borders(table)
    for row_index, row in enumerate(rows):
        for col_index in range(width):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.05
            value = row[col_index] if col_index < len(row) else ""
            add_inline(paragraph, value, size=9.5, color=INK if row_index == 0 else None)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
                set_cell_shading(cell, LIGHT)
    # Empty paragraph is layout only; it adds no wording from outside source.
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def parse_table(lines: list[str], start: int):
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw = lines[index].strip().strip("|")
        cells = [cell.strip() for cell in raw.split("|")]
        # Markdown separator row; alignment punctuation is formatting syntax.
        if cells and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            index += 1
            continue
        rows.append(cells)
        index += 1
    return rows, index


def render_markdown(source: Path, destination: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_document(doc)

    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            rows, index = parse_table(lines, index)
            add_markdown_table(doc, rows)
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(10)
            add_inline(paragraph, line[2:].strip(), size=19, color=INK)
            for run in paragraph.runs:
                run.bold = True
            index += 1
            continue
        if line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 2")
            add_inline(paragraph, line[3:].strip(), size=13, color=BLUE)
            for run in paragraph.runs:
                run.bold = True
            index += 1
            continue
        if line.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 3")
            add_inline(paragraph, line[4:].strip(), size=12, color=INK)
            for run in paragraph.runs:
                run.bold = True
            index += 1
            continue
        if line.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.2)
            paragraph.paragraph_format.right_indent = Inches(0.1)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.line_spacing = 1.1
            add_inline(paragraph, line[2:].strip(), size=10.2, color=INK, italic=True)
            p_pr = paragraph._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), NOTE)
            p_pr.append(shd)
            index += 1
            continue
        # Keep list and numbered markers as literal source text to avoid Word
        # auto-numbering changing the visible wording.
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.15
        add_inline(paragraph, line.strip(), size=11)
        index += 1

    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)


def main() -> None:
    jobs = (
        (SUPPORT / "lesson-02-综合口语评量表-draft.md",
         SUPPORT / "lesson-02-综合口语评量表-draft.docx"),
        (SUPPORT / "lesson-02-课末检查-draft.md",
         SUPPORT / "lesson-02-课末检查-draft.docx"),
    )
    for source, destination in jobs:
        if not source.is_file():
            raise FileNotFoundError(source)
        render_markdown(source, destination)
        print(destination)


if __name__ == "__main__":
    main()
