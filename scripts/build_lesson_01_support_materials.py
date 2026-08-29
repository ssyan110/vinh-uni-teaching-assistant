#!/usr/bin/env python3
"""Build printable Simplified-Chinese support materials for Boya Lesson 1."""

from __future__ import annotations

import json
import os
from pathlib import Path
from typing import Iterable, Sequence

from production_gate import assert_ready


PROJECT_ROOT = Path(__file__).resolve().parents[1]
CONFIG = json.loads((PROJECT_ROOT / "project.config.json").read_text(encoding="utf-8"))
HISTORICAL_LESSON_KEY = "boya-intermediate-i:lesson-01"
active_lesson_key = CONFIG.get("active_context", {}).get("lesson_key")
expected_lesson_root = (PROJECT_ROOT / "lessons/boya-intermediate-i/lesson-01").resolve()
configured_lesson_root = (PROJECT_ROOT / CONFIG.get("lesson_root", "")).resolve()
if active_lesson_key != HISTORICAL_LESSON_KEY or configured_lesson_root != expected_lesson_root:
    raise RuntimeError(
        "build_lesson_01_support_materials.py is historical and scoped to "
        f"{HISTORICAL_LESSON_KEY}; active context is {active_lesson_key or 'missing'} "
        f"and lesson_root is {configured_lesson_root}. Use a lesson-key-scoped builder for the current offering."
    )
DESIGN = json.loads((PROJECT_ROOT / CONFIG["design_system"]).read_text(encoding="utf-8"))
LESSON_ROOT = PROJECT_ROOT / CONFIG["lesson_root"]
# Generators write drafts only. An approved path must be supplied through the
# explicit approval/migration workflow, never by a default build command.
OUTPUT_ROOT = Path(os.environ.get("BOYA_LESSON_DRAFT_ROOT", str(LESSON_ROOT / "10-design" / "support-draft")))
STUDENT_DIR = OUTPUT_ROOT / "student"
ACTIVITY_DIR = OUTPUT_ROOT / "activities"
ASSESSMENT_DIR = OUTPUT_ROOT / "assessment"
SUPPORT_DIR = OUTPUT_ROOT / "support"

# Run before importing optional document/rendering libraries or creating any
# output directory. A bad path or missing approval must fail without side effects.
assert_ready("support", OUTPUT_ROOT)

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = DESIGN["fonts"]["latin"]
CJK_FONT = DESIGN["fonts"]["cjk"]
COLORS = DESIGN["colors"]
INK = COLORS["ink"]
ACCENT = COLORS["teal"]
DARK_ACCENT = COLORS["purple"]
LIGHT_BLUE = COLORS["blue"]
LIGHT_GRAY = COLORS["paper"]
FORM_LINE = COLORS["line"]
WHITE = COLORS["white"]


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, size: float = 12, color: str = INK, bold: bool | None = None,
                 italic: bool | None = None) -> None:
    run.font.name = FONT
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), FONT)
    r_pr.rFonts.set(qn("w:hAnsi"), FONT)
    r_pr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    r_pr.rFonts.set(qn("w:cs"), FONT)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph(paragraph, before: float = 0, after: float = 6,
                  line_spacing: float = 1.25, alignment=None) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing
    if alignment is not None:
        paragraph.alignment = alignment


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 140,
                     bottom: int = 100, end: int = 140) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color: str = FORM_LINE, size: int = 8,
                    edges: Iterable[str] = ("top", "left", "bottom", "right")) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in edges:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: Sequence[int], indent_dxa: int = 120) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_page(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
    normal.font.size = Pt(12)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Title", 22, INK, 0, 8),
        ("Heading 1", 16, ACCENT, 16, 8),
        ("Heading 2", 14, DARK_ACCENT, 12, 6),
        ("Heading 3", 12, DARK_ACCENT, 9, 4),
    ):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = style_name != "Title"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15

    for section in document.sections:
        header = section.header
        hp = header.paragraphs[0]
        hp.text = ""
        set_paragraph(hp, after=2, line_spacing=1.0)
        run = hp.add_run("博雅汉语听说 · 第一课《中国人的姓名》")
        set_run_font(run, size=9, color="64748B", bold=True)
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = ""
        set_paragraph(fp, after=0, line_spacing=1.0, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
        run = fp.add_run("荣市大学华语听说课程")
        set_run_font(run, size=8.5, color="64748B")


def new_document() -> Document:
    document = Document()
    set_page(document)
    return document


def add_title(document: Document, title: str, subtitle: str | None = None) -> None:
    p = document.add_paragraph(style="Title")
    set_paragraph(p, after=4, line_spacing=1.05)
    run = p.add_run(title)
    set_run_font(run, size=22, color=INK, bold=True)
    if subtitle:
        p = document.add_paragraph()
        set_paragraph(p, after=12, line_spacing=1.1)
        run = p.add_run(subtitle)
        set_run_font(run, size=12, color="526579")


def add_heading(document: Document, text: str, level: int = 1,
                page_break_before: bool = False) -> None:
    p = document.add_paragraph(text, style=f"Heading {level}")
    set_paragraph(p, before={1: 16, 2: 12, 3: 9}.get(level, 8),
                  after={1: 8, 2: 6, 3: 4}.get(level, 4), line_spacing=1.1)
    p.paragraph_format.page_break_before = page_break_before
    for run in p.runs:
        set_run_font(run, size={1: 16, 2: 14, 3: 12}.get(level, 12),
                     color=ACCENT if level == 1 else DARK_ACCENT, bold=True)


def add_body(document: Document, text: str, bold_prefix: str | None = None,
             after: float = 6, align=None) -> None:
    p = document.add_paragraph()
    set_paragraph(p, after=after, alignment=align)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, size=12, color=INK, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, size=12, color=INK)
    else:
        run = p.add_run(text)
        set_run_font(run, size=12, color=INK)


def add_numbered(document: Document, items: Sequence[str]) -> None:
    for item in items:
        p = document.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        set_paragraph(p, after=4, line_spacing=1.2)
        run = p.add_run(item)
        set_run_font(run, size=12, color=INK)


def add_bulleted(document: Document, items: Sequence[str]) -> None:
    for item in items:
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        set_paragraph(p, after=4, line_spacing=1.2)
        run = p.add_run(f"· {item}")
        set_run_font(run, size=12, color=INK)


def add_meta(document: Document, fields: Sequence[tuple[str, str]]) -> None:
    table = document.add_table(rows=len(fields), cols=2)
    set_table_geometry(table, [1900, 7460])
    for row, (label, value) in zip(table.rows, fields):
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for cell in row.cells:
            set_cell_border(cell, color="D2DCE7", size=6)
        p = row.cells[0].paragraphs[0]
        p.text = ""
        set_paragraph(p, after=0, line_spacing=1.1)
        run = p.add_run(label)
        set_run_font(run, size=11, color=DARK_ACCENT, bold=True)
        p = row.cells[1].paragraphs[0]
        p.text = ""
        set_paragraph(p, after=0, line_spacing=1.1)
        run = p.add_run(value)
        set_run_font(run, size=11, color=INK)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_form_table(document: Document, headers: Sequence[str], rows: Sequence[Sequence[str]],
                   widths: Sequence[int], header_fill: str = LIGHT_BLUE,
                   row_height: float | None = None) -> None:
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, header_fill)
        set_cell_border(cell, color="B7C6D6", size=7)
        p = cell.paragraphs[0]
        p.text = ""
        set_paragraph(p, after=0, line_spacing=1.05, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run(header)
        set_run_font(run, size=10.5, color=DARK_ACCENT, bold=True)
    for r_index, row_data in enumerate(rows, start=1):
        row = table.rows[r_index]
        if row_height:
            row.height = Inches(row_height)
        for c_index, value in enumerate(row_data):
            cell = row.cells[c_index]
            set_cell_border(cell, color="CBD5E1", size=6)
            p = cell.paragraphs[0]
            p.text = ""
            set_paragraph(p, after=0, line_spacing=1.1)
            run = p.add_run(value)
            set_run_font(run, size=10.5, color=INK)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_label_box(document: Document, label: str, text: str = "") -> None:
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_border(cell, color="C7D2DF", size=8)
    p = cell.paragraphs[0]
    p.text = ""
    set_paragraph(p, after=2, line_spacing=1.1)
    run = p.add_run(label)
    set_run_font(run, size=11, color=DARK_ACCENT, bold=True)
    if text:
        run = p.add_run(f"  {text}")
        set_run_font(run, size=11, color=INK)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_card_table(document: Document, cards: Sequence[tuple[str, str]], columns: int = 2,
                   card_fill: str = "F8FAFC") -> None:
    rows = (len(cards) + columns - 1) // columns
    table = document.add_table(rows=rows, cols=columns)
    widths = [9360 // columns] * columns
    widths[-1] += 9360 - sum(widths)
    set_table_geometry(table, widths, indent_dxa=120)
    for index, (title, body) in enumerate(cards):
        cell = table.cell(index // columns, index % columns)
        set_cell_shading(cell, card_fill)
        set_cell_border(cell, color="9FB3C8", size=10)
        p = cell.paragraphs[0]
        p.text = ""
        set_paragraph(p, after=4, line_spacing=1.05)
        run = p.add_run(title)
        set_run_font(run, size=12, color=DARK_ACCENT, bold=True)
        for paragraph_text in body.split("\n"):
            p = cell.add_paragraph()
            set_paragraph(p, after=3, line_spacing=1.1)
            run = p.add_run(paragraph_text)
            set_run_font(run, size=11, color=INK)
    for index in range(len(cards), rows * columns):
        set_cell_border(table.cell(index // columns, index % columns), color=WHITE, size=0)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_role_card(document: Document, title: str, task: str,
                  steps: Sequence[str], completion: str) -> None:
    add_heading(document, title, 1, page_break_before=True)
    add_label_box(document, "你的任务", task)
    add_heading(document, "活动中", 2)
    add_bulleted(document, steps)
    add_label_box(document, "完成条件", completion)


def add_page_break(document: Document) -> None:
    p = document.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_student_header(document: Document, title: str, subtitle: str, time_text: str) -> None:
    add_title(document, title, subtitle)
    add_meta(document, [("姓名", "________________"), ("班级", "________________"), ("日期", "________________")])
    add_label_box(document, "任务", f"{time_text}  ·  请先读清楚，再开始活动。")


def build_prep_a() -> Path:
    document = new_document()
    add_student_header(document, "课前预习卡 A：姓名", "第一课《中国人的姓名》", "约 15 分钟")
    add_heading(document, "一、先准备", 1)
    add_numbered(document, [
        "快速阅读教材第 1–9 页，先看懂主题，不要求逐字翻译。",
        "听音频 1-1 至 1-6 至少一次。听不清的地方，用“？”标出来。",
        "准备自己的姓名意思、来历和读音；没有中文名字时，可以准备一个你想使用的中文名字。",
        "访问三位中文使用者，了解他们名字的意思或来历。",
    ])
    add_heading(document, "二、我的姓名信息", 1)
    add_form_table(document, ["项目", "我的记录"], [
        ["中文名字", ""],
        ["怎么读", ""],
        ["有什么意思", ""],
        ["从哪里来", ""],
        ["我喜欢／不太喜欢，因为", ""],
    ], [2200, 7160], row_height=0.42)
    add_heading(document, "三、上课问同学（选一个）", 1)
    add_bulleted(document, [
        "你的名字有什么意思？",
        "你的名字从哪里来？",
        "你喜欢自己的名字吗？为什么？",
    ])
    add_heading(document, "四、小调查", 1)
    add_body(document, "请记录三位中文使用者的资料。课堂上，你要选择一个相同点、不同点或想继续追问的问题。")
    add_form_table(document, ["朋友", "姓", "名", "名字的意思和来历", "一个相同点或不同点"], [
        ["A", "", "", "", ""],
        ["B", "", "", "", ""],
        ["C", "", "", "", ""],
    ], [900, 1200, 1200, 3560, 2500], row_height=0.62)
    add_heading(document, "五、带到课堂", 1)
    add_body(document, "上课前检查：□ 我看过教材 1–9 页　　□ 我听过音频 1-1 至 1-6　　□ 我准备了姓名资料　　□ 我准备了一个问题")
    path = STUDENT_DIR / "lesson-01-prep-card-a.docx"
    document.save(path)
    return path


def build_prep_b() -> Path:
    document = new_document()
    add_student_header(document, "课前预习卡 B：姓氏", "第一课《中国人的姓名》", "约 15 分钟")
    add_heading(document, "一、先准备", 1)
    add_numbered(document, [
        "阅读教材第 10–16 页，先看懂“姓氏”这个主题。",
        "听音频 2-1 至 2-5 至少一次，标出一个听不清或看不懂的地方。",
        "准备本国两个常见的姓或名字，并想一想它们为什么常见。",
        "选择一位历史人物，准备一个 30 秒介绍。",
    ])
    add_heading(document, "二、本国姓氏比较", 1)
    add_body(document, "课堂上，你要和同伴比较姓名顺序、常见姓氏和名字的原因。")
    add_form_table(document, ["项目", "我的资料"], [
        ["常见姓氏 1", ""],
        ["常见姓氏 2", ""],
        ["常见名字", ""],
        ["为什么常见", ""],
        ["姓名顺序", ""],
    ], [3000, 6360], row_height=0.48)
    add_heading(document, "三、上课问同学（选一个）", 1)
    add_bulleted(document, [
        "你们国家的姓名是姓在前，还是名在前？",
        "在你们国家，人的姓名是什么时候产生的？",
        "你知道一个常见的姓吗？",
    ])
    add_heading(document, "四、历史人物研究卡", 1)
    add_form_table(document, ["项目", "我的记录"], [
        ["人物姓名", ""],
        ["他的姓", ""],
        ["生活的时代", ""],
        ["我想介绍的一件事", ""],
        ["我准备怎样开始介绍", ""],
    ], [2300, 7060], row_height=0.45)
    add_heading(document, "五、带到课堂", 1)
    add_body(document, "□ 我看过教材 10–16 页　　□ 我听过音频 2-1 至 2-5　　□ 我准备了两个姓氏　　□ 我准备了一个人物　　□ 我选好了课堂问题")
    path = STUDENT_DIR / "lesson-01-prep-card-b.docx"
    document.save(path)
    return path


def build_activity_cards_legacy() -> Path:
    """Retain the former combined file builder only for historical recovery."""
    document = new_document()
    add_student_header(document, "课堂活动卡", "第一课《中国人的姓名》", "按课堂指示完成")
    add_heading(document, "活动一：姓名访谈", 1)
    add_label_box(document, "小组", "三人一组　·　每轮更换角色　·　每轮约 4 分钟")
    add_card_table(document, [
        ("访谈者", "先问一个问题，再根据回答追问。\n请保持中文互动。"),
        ("回答者", "回答问题，并说出自己的例子。\n听不清时可以说：“请再说一次。”"),
        ("观察者", "记录一个好的追问和一句需要再说的话。\n不要记录所有错误。"),
    ], columns=3)
    add_heading(document, "访谈问题", 2)
    add_form_table(document, ["先问", "再追问"], [
        ["你的名字怎么读？", "这个名字有什么意思？"],
        ["这个名字从哪里来？", "你喜欢这个名字吗？为什么？"],
        ["你的姓和名有什么故事？", "如果重新起名儿，你会怎么选？"],
    ], [4680, 4680], row_height=0.48)
    add_heading(document, "观察记录", 2)
    add_form_table(document, ["我听到的一个好问题", "这句话需要再说一次"], [["", ""]], [4680, 4680], row_height=0.8)

    add_heading(document, "活动二：起名儿公司", 1, page_break_before=True)
    add_label_box(document, "任务", "三至四人一组　·　约 20 分钟　·　先问清楚，再提出姓名")
    add_heading(document, "活动怎么做", 2)
    add_numbered(document, [
        "客户拿一张客户卡，先自己读懂要求，圈出最重要的两个条件。客户卡先不给顾问看。",
        "顾问先问至少两个问题。客户根据客户卡回答；如果顾问没有问到重要条件，客户可以补充。",
        "记录员在命名顾问角色卡的“姓名提案”处写下客户的条件。顾问根据回答提出一个姓名，并说明读音、字义和两个理由。",
        "客户听完提案后，可以接受，也可以提出修改要求；最后问顾问一个问题。",
        "小组使用命名顾问角色卡中的“姓名提案”准备 90 秒小组呈现。",
    ])
    add_label_box(document, "开始前", "先分角色，再抽一张客户卡；每组使用一张客户卡完成一次提案。")

    add_role_card(document, "角色卡：客户",
                  "先拿一张客户卡，读懂要求，圈出两个最重要的条件。",
                  [
                      "客户卡先不给顾问看；让顾问通过提问了解你的条件。",
                      "根据客户卡回答顾问的问题；没有问到重要条件时，可以补充。",
                      "听完姓名提案后，可以接受，也可以提出修改要求。",
                      "最后问顾问一个问题。",
                  ],
                  "顾问知道你的两个重要条件，你提出一个追问。")

    add_role_card(document, "角色卡：命名顾问",
                  "先问清楚客户想要什么，再提出一个合适的姓名。",
                  [
                      "至少问客户两个问题，不要马上猜姓名。",
                      "提出姓名，并说明怎么读、字的意思和两个理由。",
                      "听到客户的追问后，回答问题；必要时修改提案。",
                  ],
                  "你提出一个姓名，说清楚读音、字义和两个理由。")

    add_role_card(document, "角色卡：记录员",
                  "记录客户的要求和小组最后的姓名提案。",
                  [
                      "写下客户最重视的两个条件和顾问问过的问题。",
                      "记录顾问提出的姓名、读音、字义和两个理由。",
                      "记录客户的接受或修改要求，把最后结果写在命名顾问角色卡的“姓名提案”处。",
                  ],
                  "命名顾问角色卡中的“姓名提案”内容完整，小组可以用它准备 90 秒呈现。")

    add_role_card(document, "角色卡：观察员",
                  "观察小组怎样提问、回答和提出姓名。",
                  [
                      "检查顾问有没有问至少两个问题。",
                      "检查顾问有没有说明读音、字义和两个理由。",
                      "检查客户有没有回答追问；选一句话准备重做。",
                  ],
                  "你能说出一个做得好的地方，并选一句话和小组重做。")

    add_heading(document, "客户卡一：明亮的名字", 1, page_break_before=True)
    add_body(document, "我想给女儿起一个两个字的名字。名字要让人想到阳光、春天或希望，读起来要容易。")
    add_label_box(document, "客户先做", "先自己读懂这张卡，圈出两个最重要的条件。不要把卡给顾问看，等顾问提问。")
    add_form_table(document, ["客户最重视的条件", "顾问要继续问什么"], [["", ""]], [4680, 4680], row_height=1.0)
    add_heading(document, "客户卡二：有力量的名字", 1, page_break_before=True)
    add_body(document, "我想给儿子起一个名字。名字可以表示勇敢、努力或责任，但是不要太夸张，也不要太难读。")
    add_label_box(document, "客户先做", "先自己读懂这张卡，圈出两个最重要的条件。不要把卡给顾问看，等顾问提问。")
    add_form_table(document, ["客户最重视的条件", "顾问要继续问什么"], [["", ""]], [4680, 4680], row_height=1.0)
    add_heading(document, "客户卡三：和自然有关的名字", 1, page_break_before=True)
    add_body(document, "我希望孩子的名字和水、山、树或天空有关。名字可以有一个字，也可以有两个字；我更重视字的意思。")
    add_label_box(document, "客户先做", "先自己读懂这张卡，圈出两个最重要的条件。不要把卡给顾问看，等顾问提问。")
    add_form_table(document, ["客户最重视的条件", "顾问要继续问什么"], [["", ""]], [4680, 4680], row_height=1.0)
    add_heading(document, "客户卡四：容易介绍的名字", 1, page_break_before=True)
    add_body(document, "我希望孩子的名字容易读、容易写，也容易向不同国家的人介绍。名字要有友好、快乐或平安的意思。")
    add_label_box(document, "客户先做", "先自己读懂这张卡，圈出两个最重要的条件。不要把卡给顾问看，等顾问提问。")
    add_form_table(document, ["客户最重视的条件", "顾问要继续问什么"], [["", ""]], [4680, 4680], row_height=1.0)
    add_heading(document, "命名顾问角色卡：姓名提案", 1, page_break_before=True)
    add_label_box(document, "记录员填写", "记录最后的姓名、读音、字的意思、两个理由和客户的追问。")
    add_form_table(document, ["姓名", "怎么读", "字的意思", "两个理由", "客户的追问"], [["", "", "", "", ""]], [1700, 1500, 1900, 2600, 1660], row_height=1.05)

    add_heading(document, "活动三：电影演员中文名", 1, page_break_before=True)
    add_label_box(document, "任务", "四人一组　·　约 10 分钟　·　根据演员名字的声音和意思提出中文名")
    add_card_table(document, [
        ("演员甲", "原名或声音：________________\n你想到的意思：________________\n中文名：________________\n理由：________________"),
        ("演员乙", "原名或声音：________________\n你想到的意思：________________\n中文名：________________\n理由：________________"),
        ("演员丙", "原名或声音：________________\n你想到的意思：________________\n中文名：________________\n理由：________________"),
    ], columns=1)
    add_heading(document, "小组呈现", 2, page_break_before=True)
    add_form_table(document, ["我们选择的两个名字", "读音和意思", "同伴的问题", "我们的回答"], [["", "", "", ""]], [2300, 2500, 2200, 2360], row_height=0.85)

    add_heading(document, "活动四：姓氏信息站", 1, page_break_before=True)
    add_label_box(document, "轮换", "三人一组　·　每站约 6 分钟　·　每个人都要提问和回答")
    add_heading(document, "第一站：句式信息交换", 2)
    add_body(document, "每个人拿到不同的姓氏资料。请先问清楚，再把资料告诉同伴。")
    add_form_table(document, ["我的资料", "我问同伴", "我听到的资料"], [
        ["姓：________\n单姓／复姓：________", "这是单姓还是复姓？", ""],
        ["姓名顺序：________\n一个特点：________", "你们国家呢？", ""],
        ["一个想比较的地方：________", "你同意吗？为什么？", ""],
    ], [2900, 3000, 3460], row_height=0.7)
    add_heading(document, "第二站：文化比较", 2, page_break_before=True)
    add_body(document, "先说本国资料，再和同伴比较中国姓名。不要只说“不同”，要说出一个具体例子。")
    add_form_table(document, ["比较项目", "本国", "中国", "我的问题"], [
        ["姓和名的顺序", "", "姓在前，名在后", ""],
        ["常见的姓", "", "张、王、李、刘等", ""],
        ["单姓／复姓", "", "两种都有", ""],
    ], [2200, 2300, 2700, 2160], row_height=0.58)

    add_heading(document, "活动四：姓氏信息站", 1, page_break_before=True)
    add_heading(document, "第三站：阅读拼图", 2)
    add_body(document, "每个人读一张资料卡，找出一个重点，再教给新小组。请使用“因为……所以……”或“例如……”说清楚。")
    add_card_table(document, [
        ("资料卡 A", "中国人的姓大多是单姓。张、王、李、刘是常见的姓。\n我教给小组的重点：________________"),
        ("资料卡 B", "有些人的姓是复姓，例如司马、欧阳、诸葛、上官。\n我教给小组的重点：________________"),
        ("资料卡 C", "有些姓和古代的国名、地名或官名有关系。\n我教给小组的重点：________________"),
        ("资料卡 D", "《百家姓》收集了很多姓氏。不同姓氏的人都可以有自己的故事。\n我教给小组的重点：________________"),
    ], columns=2)
    add_heading(document, "第四站：朗读与人物研究", 2)
    add_body(document, "先朗读下面的姓，再介绍一位历史人物。听者记录两个姓，并提出一个问题。")
    add_form_table(document, ["单姓", "复姓", "历史人物的姓", "同伴的问题"], [["张、王、李、刘\n________________", "司马、欧阳、诸葛、上官\n________________", "", ""]], [2350, 2550, 2300, 2160], row_height=0.85)

    add_heading(document, "活动五：调查与研究表", 1, page_break_before=True)
    add_heading(document, "姓名调查：名字的含义", 2)
    add_body(document, "请把你访问到的名字按意思分类。一个名字可以有多个意思，请写出你听到的理由。")
    add_form_table(document, ["时代意义", "地域特点", "美好愿望"], [
        ["1. __________________", "1. __________________", "1. __________________"],
        ["2. __________________", "2. __________________", "2. __________________"],
        ["3. __________________", "3. __________________", "3. __________________"],
        ["4. __________________", "4. __________________", "4. __________________"],
        ["5. __________________", "5. __________________", "5. __________________"],
    ], [3120, 3120, 3120], row_height=0.48)
    add_heading(document, "小组报告大纲", 2)
    add_form_table(document, ["部分", "我们的内容"], [
        ["开头：我们调查了什么", ""],
        ["发现一：一个名字和理由", ""],
        ["发现二：一个比较或例子", ""],
        ["结尾：我们的观察", ""],
        ["同伴可能问的问题", ""],
    ], [2600, 6760], row_height=0.58)
    add_heading(document, "完成检查", 2)
    add_body(document, "□ 我们有资料　　□ 我们有例子　　□ 我们说明了理由　　□ 每个人都有一句话要说")

    raise RuntimeError(
        "The retired combined activity-card builder is disabled. "
        "Use build_lesson_01_activity_packages.py and its DOCX files instead."
    )


def build_assessment() -> Path:
    document = new_document()
    add_student_header(document, "同伴反馈与出口卡", "第一课《中国人的姓名》", "说完以后填写")
    add_heading(document, "同伴反馈表", 1)
    add_body(document, "请先听懂同伴的意思，再给一条具体反馈。不要把所有错误都写下来。")
    add_form_table(document, ["我听懂了什么", "哪一句需要再说", "下一次可以怎么做"], [["", "", ""]], [3120, 3120, 3120], row_height=1.0)
    add_heading(document, "短讲／提案评量", 2)
    add_form_table(document, ["面向", "0", "1", "2", "3"], [
        ["信息完整", "□", "□", "□", "□"],
        ["理由或证据", "□", "□", "□", "□"],
        ["互动回应", "□", "□", "□", "□"],
        ["容易听懂", "□", "□", "□", "□"],
        ["结构清楚", "□", "□", "□", "□"],
    ], [2800, 1640, 1640, 1640, 1640], row_height=0.5)
    add_body(document, "我给同伴的一句话：____________________________________________________________")
    add_body(document, "同伴重做以后，我听到的变化：________________________________________________")

    add_heading(document, "出口卡怎么用", 2)
    add_body(document, "每节课结束时填写一张对应的出口卡。先完成课堂活动，再填写；教师用卡片安排下一节的短修补。")
    add_form_table(document, ["出口卡", "填写时间", "对应课堂活动"], [
        ["1", "P1 结束", "姓名对话听力与姓名访谈（E01-001–E01-005、E01-034）"],
        ["2", "P2 结束", "句式情境对话（E01-009–E01-012）"],
        ["3", "P3 结束", "第二次听力与谐音表达（E01-006–E01-008、E01-015）"],
        ["4", "P4 结束", "起名儿公司提案（E01-016）"],
        ["5", "P5 结束", "本国姓氏比较与对话重建（E01-019–E01-023）"],
        ["6", "P6 结束", "姓氏信息站与对话反思（E01-027–E01-033）"],
    ], [1100, 1900, 6360], row_height=0.48)

    prompts = [
        ("出口卡 1｜P1 姓名访谈", "姓名对话听力与姓名访谈（E01-001–E01-005、E01-034）", "说出一个姓名理由，再写下一个你在音频中听到的信息。", "我的姓名理由：________________________\n我的听力证据：________________________"),
        ("出口卡 2｜P2 句式对话", "句式情境对话（E01-009–E01-012）", "用一个句式回应姓名选择问题，并说明你的理由。", "我说：________________________________\n我使用的句式：________________________"),
        ("出口卡 3｜P3 听力与表达", "第二次听力与谐音表达（E01-006–E01-008、E01-015）", "写下一个你听懂的内容和一个还想确认的问题。", "我听懂了：____________________________\n我还想确认：__________________________"),
        ("出口卡 4｜P4 起名儿公司", "起名儿公司提案（E01-016）", "写出一个重要的命名条件，并说明理由。", "重要条件：____________________________\n理由：________________________________"),
        ("出口卡 5｜P5 姓氏比较", "本国姓氏比较与对话重建（E01-019–E01-023）", "说出本国和中国姓名的一个相同点或不同点，再写一个追问。", "相同点／不同点：______________________\n我的追问：____________________________"),
        ("出口卡 6｜P6 姓氏信息站", "姓氏信息站与对话反思（E01-027–E01-033）", "完成一句“现在我能……”，再写下下一步目标。", "现在我能：____________________________\n下一步我想改进：______________________"),
    ]
    for index, (title, activity, prompt, lines) in enumerate(prompts):
        add_page_break(document)
        add_heading(document, title, 1)
        add_meta(document, [("姓名", "________________"), ("小组", "________________"), ("日期", "________________")])
        add_label_box(document, "对应课堂活动", activity)
        add_label_box(document, "请完成", prompt)
        add_form_table(document, ["我的回答"], [[lines]], [9360], row_height=1.6)
        add_body(document, "我今天最想记住的一个说法：________________________________________________")
        add_body(document, "□ 我完成了今天的听说任务　　□ 我给同伴提了一个问题　　□ 我根据反馈重做了一次")

    path = ASSESSMENT_DIR / "lesson-01-feedback-exit-tickets.docx"
    document.save(path)
    return path


def write_manifest(paths: Sequence[Path],
                   activity_specs: Sequence[dict] = (),
                   activity_manifest_path: Path | None = None) -> Path:
    SUPPORT_DIR.mkdir(parents=True, exist_ok=True)
    coverage = [
        {"material_id": "PREP-A", "file": "student/lesson-01-prep-card-a.docx", "activities": ["E01-034", "E01-035"], "periods": ["P1", "P2", "P3", "P4"]},
        {"material_id": "PREP-B", "file": "student/lesson-01-prep-card-b.docx", "activities": ["E01-019", "E01-031", "E01-032"], "periods": ["P5", "P6"]},
    ]
    for spec in activity_specs:
        guide = spec.get("teacher_guide") or {}
        coverage.append({
            "material_id": spec["activity_id"],
            "file": guide.get("docx"),
            "folder": spec["folder"],
            "student_materials": [item["docx"] for item in spec.get("student_materials", [])],
            "activities": {
                "ACT-01": ["E01-003", "E01-034", "E01-035"],
                "ACT-02": ["E01-016"],
                "ACT-03": ["E01-017"],
                "ACT-04": ["E01-027", "E01-028", "E01-029", "E01-032", "E01-033"],
                "ACT-05": ["E01-018", "E01-031", "E01-035"],
            }.get(spec["activity_id"], []),
            "periods": spec["periods"],
        })
    output_files = [str(path.relative_to(OUTPUT_ROOT)) for path in paths]
    data = {
        "package": "boya-intermediate-lesson-01-support-materials",
        "status": "draft_for_support_material_review",
        "lesson_completion_status": "lesson_01_in_production",
        "language": "简体中文",
        "vietnamese_explanation_policy": "仅在必要时用于说明",
        "source_of_truth": "lessons/boya-intermediate-i/lesson-01/20-approved/teacher-manual/第一课简易教案.docx",
        "teacher_guide_approval": "checked_by_production_gate_from_authority_manifest",
        "activity_material_layout": "每个活动一个资料夹；合并后的 Word 文件直接放在活动资料夹内，不增加中间资料夹；不生成活动卡 PDF。",
        "activity_package_manifest": str(activity_manifest_path.relative_to(OUTPUT_ROOT)) if activity_manifest_path else None,
        "coverage": coverage + [
            {"material_id": "ASSESS-01", "file": "assessment/lesson-01-feedback-exit-tickets.docx", "activities": ["E01-015", "E01-016", "E01-018", "E01-030", "E01-031"], "periods": ["P3", "P4", "P5", "P6"]},
            {"material_id": "ASSESS-02", "file": "assessment/lesson-01-feedback-exit-tickets.docx", "activities": ["P1", "P2", "P3", "P4", "P5", "P6"], "periods": ["P1", "P2", "P3", "P4", "P5", "P6"]},
        ],
        "output_files": output_files,
        "exit_ticket_mapping": [
            {"card": "出口卡 1", "period": "P1", "activities": ["E01-001", "E01-002", "E01-003", "E01-004", "E01-005", "E01-034"]},
            {"card": "出口卡 2", "period": "P2", "activities": ["E01-009", "E01-010", "E01-011", "E01-012"]},
            {"card": "出口卡 3", "period": "P3", "activities": ["E01-006", "E01-007", "E01-008", "E01-015"]},
            {"card": "出口卡 4", "period": "P4", "activities": ["E01-016"]},
            {"card": "出口卡 5", "period": "P5", "activities": ["E01-019", "E01-020", "E01-021", "E01-022", "E01-023"]},
            {"card": "出口卡 6", "period": "P6", "activities": ["E01-027", "E01-028", "E01-029", "E01-030", "E01-031", "E01-032", "E01-033"]}
        ],
        "next_gate": "support_material_review_and_pre_delivery_reconciliation_then_recheck_internal_ppt_storyboard",
        "unlock_rule": "lesson_02_remains_locked_until_lesson_01_delivery",
    }
    manifest_path = SUPPORT_DIR / "manifest.json"
    manifest_path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest_path


def main() -> None:
    for directory in (STUDENT_DIR, ACTIVITY_DIR, ASSESSMENT_DIR, SUPPORT_DIR):
        directory.mkdir(parents=True, exist_ok=True)
    from build_lesson_01_activity_packages import build_all

    prep_a = build_prep_a()
    prep_b = build_prep_b()
    assessment = build_assessment()
    activity_docx, activity_specs, activity_manifest = build_all()
    paths = [prep_a, prep_b, *activity_docx, assessment]
    manifest = write_manifest(
        paths,
        activity_specs=activity_specs,
        activity_manifest_path=activity_manifest,
    )
    print(json.dumps({
        "outputs": [str(path) for path in paths],
        "activity_pdf_count": 0,
        "manifest": str(manifest),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
