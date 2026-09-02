from copy import deepcopy
from hashlib import sha256
from pathlib import Path
import json

from pptx import Presentation

ROOT = Path('/Users/ssyan110/Development/vinh-uni-teaching-assistant')
REFERENCE = ROOT / 'lessons/boya-quasi-intermediate-i/lesson-01/20-approved/pptx/lesson-01-在线预习.pptx'

def nonempty(slide):
    return [s for s in slide.shapes if hasattr(s, 'text') and s.text.strip()]

def replace_text_preserving_run_style(shape, value):
    runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
    if runs:
        runs[0].text = value
        for run in runs[1:]:
            run.text = ''
    else:
        shape.text = value

ref = Presentation(REFERENCE).slides[2]
ref_count = len(nonempty(ref))

for n in range(5, 11):
    path = ROOT / f'lessons/boya-quasi-intermediate-i/lesson-{n:02d}/10-design/pptx-draft/online/lesson-{n:02d}-在线预习.pptx'
    if not path.exists():
        print(f'lesson-{n:02d}: skipped (PPTX missing)')
        continue
    prs = Presentation(path)
    target = prs.slides[2]
    current_text = [s.text for s in nonempty(target)]
    if len(current_text) != ref_count:
        raise ValueError(f'{path}: current goals page has {len(current_text)} text shapes; reference has {ref_count}')
    for shape in list(target.shapes):
        target.shapes._spTree.remove(shape._element)
    for shape in ref.shapes:
        target.shapes._spTree.append(deepcopy(shape._element))
    copied = nonempty(target)
    for shape, value in zip(copied, current_text):
        replace_text_preserving_run_style(shape, value)
    tmp = path.with_suffix('.goals-style.tmp.pptx')
    prs.save(tmp)
    tmp.replace(path)
    manifest = path.parent / 'manifest.json'
    if manifest.exists():
        data = json.loads(manifest.read_text())
        data['sha256'] = sha256(path.read_bytes()).hexdigest()
        data['bytes'] = path.stat().st_size
        data['status'] = 'draft_goals_style_refined_2026-08-31'
        manifest.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')
    print(f'lesson-{n:02d}: goals_style_updated; text_preserved={current_text[5:]}')
