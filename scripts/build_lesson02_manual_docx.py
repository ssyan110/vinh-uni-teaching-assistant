#!/usr/bin/env python3
"""Render the Lesson 2 teacher-manual markdown draft as a formatted DOCX.

This builder is deliberately draft-only: it reads the current markdown file and
does not touch 20-approved, 30-qa, or 40-release.
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "lessons/boya-quasi-intermediate-i/lesson-02/10-design/teacher-manual/lesson-02-教师手册-draft.md"
OUT = MD.with_suffix(".docx")

INK = RGBColor(0x17, 0x32, 0x4D)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
MUTED = RGBColor(0x64, 0x74, 0x8B)
LIGHT = "E8EEF5"
CAUTION = "FFF5D6"

def set_run_font(run, cjk="KaiTi", latin="Times New Roman", size=11, bold=None, color=None, italic=None):
    run.font.name = cjk
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), cjk)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), latin)
    r_pr = run._element.get_or_add_rPr()
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_table_borders(table, color="C7D2DE", size="6"):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)

def add_text_runs(paragraph, text, size=11, color=None):
    # Minimal inline handling for bold and code spans.
    parts = re.split(r"(\*\*.*?\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        value = part[2:-2] if bold else part[1:-1] if code else part
        run = paragraph.add_run(value)
        set_run_font(run, size=size, bold=bold, color=color or (MUTED if code else None))

def add_para(doc, text="", style=None, size=11, color=None, bold=False, italic=False, after=6, before=0):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    if text:
        add_text_runs(p, text, size=size, color=color)
        for r in p.runs:
            if bold:
                r.bold = True
            if italic:
                r.italic = True
    return p

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    add_text_runs(p, text, size=10.5)
    return p

def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    add_text_runs(p, text, size=10.5)
    return p

def add_table(doc, rows):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    set_table_borders(table)
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if i == 0:
                shade(cell, LIGHT)
            value = row[j] if j < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            add_text_runs(p, value, size=9.2, color=INK if i == 0 else None)
            if i == 0:
                for r in p.runs:
                    r.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        cells = [c.strip() for c in raw.split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            i += 1
            continue
        rows.append(cells)
        i += 1
    return rows, i

def build():
    lines = MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "KaiTi")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    normal_lang = normal._element.rPr.find(qn("w:lang"))
    if normal_lang is None:
        normal_lang = OxmlElement("w:lang")
        normal._element.rPr.append(normal_lang)
    normal_lang.set(qn("w:val"), "zh-CN")
    normal_lang.set(qn("w:eastAsia"), "zh-CN")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, INK, 10, 5),
    ]:
        st = styles[name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "KaiTi")
        st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
        st_lang = st._element.rPr.find(qn("w:lang"))
        if st_lang is None:
            st_lang = OxmlElement("w:lang")
            st._element.rPr.append(st_lang)
        st_lang.set(qn("w:val"), "zh-CN")
        st_lang.set(qn("w:eastAsia"), "zh-CN")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    # Header/footer are intentionally quiet and operational.
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("准中级加速篇 I · 第2课教师手册草案")
    set_run_font(hr, size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("Draft only · lesson_key: boya-quasi-intermediate-i:lesson-02")
    set_run_font(fr, size=8, color=MUTED)

    i = 0
    first = True
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(line[2:].strip())
            set_run_font(r, size=22, color=INK, bold=True)
            sub = doc.add_paragraph()
            sub.paragraph_format.space_after = Pt(14)
            sr = sub.add_run("荣市大学华语听说课程 · 仅供来源批准后的教师手册定稿")
            set_run_font(sr, size=11, color=MUTED)
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(line[3:].strip(), style="Heading 1")
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(line[4:].strip(), style="Heading 2")
            i += 1
            continue
        if line.startswith("#### "):
            p = doc.add_paragraph(line[5:].strip(), style="Heading 3")
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.10)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.12
            add_text_runs(p, line[2:].strip(), size=10.2, color=INK)
            for r in p.runs:
                r.italic = True
            # Add a pale background to the paragraph.
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), CAUTION)
            pPr.append(shd)
            i += 1
            continue
        if line.startswith("| ") or line.startswith("|"):
            rows, i = parse_table(lines, i)
            if rows:
                add_table(doc, rows)
            continue
        if re.match(r"^\d+\.\s+", line):
            add_number(doc, re.sub(r"^\d+\.\s+", "", line))
            i += 1
            continue
        if line.startswith("- "):
            add_bullet(doc, line[2:].strip())
            i += 1
            continue
        # Standalone bold line (section cue) becomes a compact callout heading.
        if line.startswith("**") and line.endswith("**"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(line[2:-2])
            set_run_font(r, size=11.5, color=INK, bold=True)
            i += 1
            continue
        add_para(doc, line.strip(), size=10.5)
        i += 1

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)

if __name__ == "__main__":
    build()
