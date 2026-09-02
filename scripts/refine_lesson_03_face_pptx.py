from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib
import os

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
FACE = ROOT / "lessons/boya-quasi-intermediate-i/lesson-03/10-design/pptx-draft/face-to-face/lesson-03-实体课.pptx"
ONLINE3 = ROOT / "lessons/boya-quasi-intermediate-i/lesson-03/20-approved/pptx/lesson-03-在线预习.pptx"
ONLINE2 = ROOT / "lessons/boya-quasi-intermediate-i/lesson-02/20-approved/pptx/lesson-02-在线预习.pptx"
TMP = FACE.with_suffix(".refined.tmp.pptx")


def replace_shape_text(slide, old, new):
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text == old:
            shape.text_frame.paragraphs[0].runs[0].text = new
            return
    raise ValueError(f"Missing text: {old}")


prs = Presentation(str(FACE))

# The user explicitly asked for the online goals page. Copy its complete body
# styling and shapes while keeping the face-to-face header and page label.
online3 = Presentation(str(ONLINE3))
face_goals = prs.slides[2]
for shape in list(face_goals.shapes)[3:]:
    face_goals.shapes._spTree.remove(shape._element)
for shape in list(online3.slides[2].shapes)[3:]:
    face_goals.shapes._spTree.append(deepcopy(shape._element))

# Keep all current page content except the three explicitly requested task
# requirement lines.
for slide_no in (27, 41, 52):
    replace_shape_text(prs.slides[slide_no - 1], "必须使用10个课本中的词语和5个句式。", "必须使用5个课本中的词语和3个句式。")

# Delete requested slides by current thumbnail index, highest first.
for slide_no in (51, 49):
    slide_id = prs.slides._sldIdLst[slide_no - 1]
    prs.part.drop_rel(slide_id.rId)
    prs.slides._sldIdLst.remove(slide_id)

prs.save(str(TMP))
os.replace(TMP, FACE)

# Replace only the embedded route image on slide 2. The existing slide
# relationship already targets ppt/media/image2.png in this deck.
tmp_zip = FACE.with_suffix(".route.tmp.pptx")
with ZipFile(FACE) as zin, ZipFile(tmp_zip, "w", ZIP_DEFLATED) as zout, ZipFile(ONLINE2) as route_zip:
    route = route_zip.read("ppt/media/image2.png")
    for info in zin.infolist():
        data = route if info.filename == "ppt/media/image2.png" else zin.read(info.filename)
        zout.writestr(info, data)
os.replace(tmp_zip, FACE)

with ZipFile(FACE) as z:
    assert z.testzip() is None
    assert z.read("ppt/media/image2.png") == route
print("saved", FACE)
print("slides", len(Presentation(str(FACE)).slides))
print("route_sha256", hashlib.sha256(route).hexdigest())
print("deck_sha256", hashlib.sha256(FACE.read_bytes()).hexdigest())
print("deck_bytes", FACE.stat().st_size)
